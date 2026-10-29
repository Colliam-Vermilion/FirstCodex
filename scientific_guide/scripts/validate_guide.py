#!/usr/bin/env python3
"""Read-only checks of scientific homepage, citations and preserved old results."""
import argparse,ast,csv,datetime,hashlib,json,re,subprocess,sys
from pathlib import Path
from urllib.parse import unquote,urlsplit
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[2]
G=ROOT/'scientific_guide'
M=ROOT/'results/ad-oxygen-literature-20261006/oxygen_response_review'
BASE='3966c96f48f3194a4e27fc9ada0837f009345f35'
DISPATCH='results/ad-oxygen-literature-20261006/branch_migration/oxygen_response_mechanisms_20261029/validate_migration.py'
def git(*args):return subprocess.check_output(['git',*args],cwd=ROOT)
def load(p):return json.loads(p.read_text(encoding='utf-8'))
def anchors(path):
 s=path.read_text();found=set(re.findall(r'<a\s+id="([^"]+)"',s));seen={}
 for line in s.splitlines():
  if not re.match(r'^#{1,6}\s',line):continue
  t=re.sub(r'^#{1,6}\s+','',line).strip().lower();t=re.sub(r'[`*\[\]]','',t)
  t=''.join(c for c in t if c.isalnum() or c in [' ','-','_']);t=t.replace(' ','-')
  n=seen.get(t,0);seen[t]=n+1;found.add(t if n==0 else t+'-'+str(n))
 return found
def run():
 parser=argparse.ArgumentParser();parser.add_argument('--report',type=Path);args=parser.parse_args()
 summary=load(G/'summary.json');articles=load(G/'discussed_articles.json');fullmap=load(G/'literature_proposition_map.json');props=load(G/'propositions.json');natural=load(G/'natural_communities.json')
 cat=load(M/'catalog.json');C={c['id']:c for c in cat};assert len(C)==1536
 assert {r['id'] for r in fullmap}==set(C)
 assert len(articles)==42 and len({a['id'] for a in articles})==42
 assert len(fullmap)==1587 and len(natural)==25 and len(props)==10
 mapped={a['id'] for a in articles};evidence_rows=[r for r in fullmap if r.get('proposition_id')]
 assert len(evidence_rows)==93
 assert all(not r.get('proposition_id') for r in fullmap if r['id'] not in mapped)
 assert len(C)-len(mapped)==1494
 coverage={r['term']:r for r in load(G/'query_coverage.json')}
 for term,ids in {'sediment':['C03','P06'],'peat':['P07'],'wetland':['P07'],'flooded soil':['P07'],'reoxygenation':['P06','P07'],'redox fluctuation':['P07']}.items():assert coverage[term]['executed_query_ids']==ids,(term,coverage[term]['executed_query_ids'])
 for a in articles:
  c=C[a['id']];assert a['doi']==c['doi'] and a['title']==c['title'] and a['authors']==c.get('authors',[])
  assert (ROOT/a['evidence_card']).exists()
  assert set(a['propositions'])<={p['proposition_id'] for p in props}
  if a['reading_depth'].startswith('已核实'):
   old=next(e for e in load(M/'mechanism_evidence.json') if e['id']==a['id'])
   assert old['screening_stage']=='全文核实'
  else:assert '仅依据完整缓存摘要' in a['reading_depth']
 # Archive is the old root README, byte for byte. The human-readable technical
 # copy retains all original paragraphs with only relative-link rebasing.
 original=git('show',BASE+':README.md');archive=(ROOT/'technical/root_README_original_20261029.txt').read_bytes()
 assert original==archive
 assert hashlib.sha256(archive).hexdigest()==summary['original_root_README_sha256']
 technical=(ROOT/'technical/README.md').read_text().split('\n---\n\n',1)[1]
 restored=re.sub(r'(\]\()\.\./(results/[^)]+)(\))',lambda m:m[1]+m[2]+m[3],technical)
 assert restored==original.decode()
 # All old study artifacts must stay exactly the same Git objects. This detects
 # changes without reading or hashing shared PDFs, or re-downloading anything.
 changed=git('diff','--name-only',BASE).decode().splitlines()
 assert all(p=='README.md' or p==DISPATCH or p.startswith('scientific_guide/') or p.startswith('technical/') for p in changed),changed
 prior=git('ls-tree','-r','--name-only',BASE).decode().splitlines()
 preserved=[p for p in prior if p not in ['README.md',DISPATCH]]
 assert not set(preserved)&set(changed)
 # Complete material links and article DOI values are checked against existing
 # records. We do not request publisher URLs or claim newly verified availability.
 docs=[ROOT/'README.md',ROOT/'technical/README.md',*G.rglob('*.md')]
 links=0;external=0;errors=[];tracked_only=0
 tracked=set(git('ls-files').decode().splitlines())
 for doc in docs:
  for target in re.findall(r'(?<!!)\[[^\]\n]+\]\(([^)\n]+)\)',doc.read_text()):
   u=urlsplit(target.strip('<>'))
   if u.scheme:external+=1;continue
   path=(doc.parent/unquote(u.path)).resolve() if u.path else doc
   if not path.exists():
    try:rel=path.relative_to(ROOT).as_posix()
    except ValueError:errors.append((str(doc.relative_to(ROOT)),target));continue
    if rel not in tracked and not any(x.startswith(rel.rstrip('/')+'/') for x in tracked):errors.append((str(doc.relative_to(ROOT)),target));continue
    tracked_only+=1
   if u.fragment and path.suffix.lower()=='.md' and path.exists() and unquote(u.fragment) not in anchors(path):errors.append((str(doc.relative_to(ROOT)),target+' [anchor missing]'));continue
   links+=1
 assert not errors,errors
 assert all(x in anchors(ROOT/'README.md') for x in ['part-1','part-2','part-3','part-4','natural-communities','natural-search'])
 assert '<details' not in (ROOT/'README.md').read_text().lower()
 for name,expected in [('literature_proposition_map_utf8.csv',1587),('discussed_evidence_links_utf8.csv',93),('natural_communities_utf8.csv',25),('reading_log_20261029_utf8.csv',42),('query_coverage_utf8.csv',7)]:
  with (G/name).open(encoding='utf-8-sig',newline='') as f:assert len(list(csv.DictReader(f)))==expected
 wb=load_workbook(G/'scientific_reading_guide.xlsx',read_only=True,data_only=False)
 expected={'全部文献—命题':1587,'本页逐项证据':93,'天然群落与边界':25,'科学命题':10,'七类至四部分':7,'实际查询覆盖':7,'本轮补读':42}
 assert set(wb.sheetnames)==set(expected)
 for name,n in expected.items():assert sum(1 for _ in wb[name].iter_rows(values_only=True))==n+1
 wb.close()
 for p in (G/'scripts').glob('*.py'):ast.parse(p.read_text())
 result={'checked_Asia_Shanghai':datetime.datetime.now(datetime.timezone(datetime.timedelta(hours=8))).isoformat(),'commit':git('rev-parse','HEAD').decode().strip(),'branch':git('branch','--show-current').decode().strip(),'passed':True,'editorial_baseline':BASE,'old_README_preserved_byte_for_byte':True,'technical_copy_all_old_content_preserved_links_rebased':True,'existing_tracked_files_preserved_except_root_and_validation_dispatch':len(preserved),'old_review_catalog_screening_evidence_priority_workbook_RIS_and_progress_unchanged':True,'articles_discussed':42,'proposition_links':93,'all_existing_ids_represented':1536,'other_records_without_new_proposition_claim':1494,'natural_and_boundary_records':25,'workbook_sheets_rows_verified':expected,'local_file_and_anchor_links_resolved':links,'links_checked_by_Git_tree_without_retrieving_shared_payload':tracked_only,'external_links_values_preserved_not_publisher_HTTP_retested':external,'new_searches':0,'new_scope_screening':0,'new_paper_downloads':0,'shared_original_PDF_rehashes':0}
 if args.report:args.report.write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n')
 print(json.dumps(result,ensure_ascii=False))
if __name__=='__main__':run()
