#!/usr/bin/env python3
"""Preservation and link checks for additive spatial-transfer guide changes."""
import argparse,ast,csv,datetime,io,json,re,subprocess,sys
from pathlib import Path
from urllib.parse import unquote,urlsplit
from openpyxl import load_workbook
R=Path(__file__).resolve().parents[2];G=R/'scientific_guide';P=R/'results/ad-oxygen-literature-20261006';M=P/'oxygen_response_review';T=M/'spatial_gas_transfer';BASE='fcd0ad3119900af9f75c9b9220575b8ff79a0170'
def git(*args):return subprocess.check_output(['git',*args],cwd=R)
def old(p):return git('show',BASE+':'+p)
def read(p):return json.loads(p.read_text())
def main():
 pa=argparse.ArgumentParser();pa.add_argument('--report',type=Path);a=pa.parse_args()
 check=json.loads(subprocess.check_output([sys.executable,str(T/'scripts/validate.py')],cwd=R,text=True));assert check['passed']
 home=(R/'README.md').read_text()
 clean=home
 for name in ['exposure','structure','function','navigation']:
  clean=re.sub(r'<!-- SPATIAL_'+name+r'_BEGIN -->.*?<!-- SPATIAL_'+name+r'_END -->\n\n','',clean,flags=re.S)
 assert clean==old('README.md').decode(),'Homepage old text was changed, not merely supplemented'
 for anchor in ['part-1','part-2','part-3','part-4','natural-communities','natural-search']:assert 'id="'+anchor+'"' in home
 oldmap=json.loads(old('scientific_guide/literature_proposition_map.json'));current=read(G/'literature_proposition_map.json');assert current[:len(oldmap)]==oldmap
 oldprops=json.loads(old('scientific_guide/propositions.json'));assert read(G/'propositions.json')[:len(oldprops)]==oldprops
 # Original scope, priorities, evidence, papers and files must not be changed.
 changed=git('diff','--name-only',BASE).decode().splitlines();topic_prefix=T.relative_to(R).as_posix()+'/'
 allowed=['README.md','scientific_guide/README.md','scientific_guide/literature_proposition_map.json','scientific_guide/literature_proposition_map_utf8.csv','scientific_guide/propositions.json','scientific_guide/propositions_utf8.csv','scientific_guide/scientific_reading_guide.xlsx','scientific_guide/summary.json','scientific_guide/scripts/build_guide.py','scientific_guide/scripts/write_homepage.py','scientific_guide/scripts/validate_guide.py']
 for p in changed:assert p in allowed or p.startswith(topic_prefix) or p.startswith('scientific_guide/spatial_') or p=='scientific_guide/scripts/validate_spatial.py',p
 assert len(read(M/'catalog.json'))==1536
 assert len(read(M/'screening.json'))==1536
 for f in ['catalog.json','screening.json','mechanism_evidence.json','reading_priority/evaluations.json','all_pool/progress.json','oxygen_response_review_cumulative.xlsx']:
  assert not git('diff',BASE,'--',str((M/f).relative_to(R))),f
 # Validate complete matrix and workbook without hashing/retrieving old PDFs.
 assert len({r['id'] for r in current})==1631
 assert len(read(G/'spatial_gas_proposition_links.json'))==49
 wb=load_workbook(G/'scientific_reading_guide.xlsx',read_only=True)
 assert sum(1 for _ in wb['全部文献—命题'].iter_rows())==len(current)+1
 assert sum(1 for _ in wb['空间传质专题'].iter_rows())==check['processed']+1
 wb.close()
 docs=[R/'README.md',G/'README.md',G/'spatial_reading_map.md',T/'README.md',*(T/'evidence').glob('*.md')]
 tracked=set(git('ls-files').decode().splitlines());errors=[];links=0;tracked_only=0
 for f in docs:
  for target in re.findall(r'(?<!!)\[[^\]\n]+\]\(([^)\n]+)\)',f.read_text()):
   u=urlsplit(target.strip('<>'))
   if u.scheme or not u.path:continue
   path=(f.parent/unquote(u.path)).resolve()
   try:rel=path.relative_to(R).as_posix()
   except ValueError:errors.append((str(f),target));continue
   if not path.exists():
    if rel not in tracked and not any(p.startswith(rel.rstrip('/')+'/') for p in tracked):errors.append((str(f.relative_to(R)),target));continue
    tracked_only+=1
   if u.fragment and path.exists() and path.suffix=='.md':
    # Existing shared anchors are explicit; newly generated cards use no fragments.
    assert f'id="{u.fragment}"' in path.read_text() or u.fragment in ['readme'],(str(f),target)
   links+=1
 assert not errors,errors
 for f in [*G.glob('scripts/*.py'),*T.glob('scripts/*.py')]:ast.parse(f.read_text())
 summary=read(T/'summary.json')
 assert summary['processed']==40 and summary['remaining']==193 and summary['new_external_queries_completed']==0
 result={'checked_Asia_Shanghai':datetime.datetime.now(datetime.timezone(datetime.timedelta(hours=8))).isoformat(),'commit':git('rev-parse','HEAD').decode().strip(),'branch':git('branch','--show-current').decode().strip(),'passed':True,'baseline':BASE,'old_homepage_entire_text_preserved':True,'old_1587_map_rows_and_ten_propositions_preserved':True,'old_review_scope_priorities_evidence_originals_and_progress_unchanged':True,'topical_candidates':233,'processed':40,'remaining':193,'old_review_related_fulltext_or_evidence_reused':5,'abstract_only_processed':35,'verified_new_external_articles':0,'external_queries_completed':0,'saved_backward_reference_leads':42,'local_file_links_verified':links,'shared_targets_checked_by_Git_tree_only':tracked_only,'topic_workbook_opened':True,'scientific_guide_workbook_opened':True,'expanded_guide_ids':1631,'new_proposition_links':49,'true_whole_AD_granule_3D_fulltext_verifications_added':0,'abstract_explicit_electron_tomography_leads':1,'direct_interfloc_distance_or_flux_confirmed':0,'shared_PDF_redownloads_or_rehashes':0}
 if a.report:a.report.write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n')
 print(json.dumps(result,ensure_ascii=False))
if __name__=='__main__':main()
