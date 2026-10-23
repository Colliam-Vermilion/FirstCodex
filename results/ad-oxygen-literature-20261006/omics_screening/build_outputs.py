"""Rebuild cumulative CSV/XLSX from human-reviewed JSON; never classify papers automatically."""
import csv,json,collections,hashlib,re
from pathlib import Path
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.utils import get_column_letter
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

ROOT=Path(__file__).resolve().parent
def read(name):return json.loads((ROOT/name).read_text(encoding='utf-8'))
def dump(name,obj):(ROOT/name).write_text(json.dumps(obj,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
def cell(value):
    if isinstance(value,(list,dict)):return json.dumps(value,ensure_ascii=False)
    if isinstance(value,str):return ILLEGAL_CHARACTERS_RE.sub(lambda m:f'[U+{ord(m.group()):04X}]',value)
    return value
def csvout(name,rows,fields):
    with (ROOT/name).open('w',encoding='utf-8-sig',newline='') as f:
        w=csv.DictWriter(f,fieldnames=fields,extrasaction='ignore');w.writeheader()
        w.writerows({k:cell(v) for k,v in r.items()} for r in rows)

RECORD_FIELDS={'record_id':'文献ID','title':'题名','year':'年份','journal':'期刊或载体','doi':'DOI','original_category':'原研究分类','document_type':'原核实载体类型','document_group':'原统计文献组','fulltext_status':'正文状态','progress':'进度','conclusion':'筛查结论','sequencing_type':'测序类型','data_access_status':'数据获取状态','batch_id':'批次','pending_reason':'待确认原因','substrate':'底物','oxygen_supply':'供氧方式','treatment_control':'处理组与对照组','platform':'测序平台','sequencing_sample_count':'测序样本数量','biological_replicates':'生物重复','method_summary':'方法判断说明','reviewed_locations':'已阅读位置','fulltext_paths':'原件路径','fulltext_sha256':'原件SHA256','supplementary_status':'补充材料状态','supplementary_paths':'补充材料路径','supplementary_identity_pending_count':'附件归属待核实件数','supplementary_review_notes':'附件核查备注','evidence_ids':'证据ID','completed_at':'完成时间UTC'}
PROJECT_FIELDS={'record_id':'文献ID','dataset_key':'独立项目去重键','raw_accession':'文中原样编号','normalized_accession':'规范编号','accession_level':'编号层级','database':'数据库','sequencing_type':'测序类型','data_origin':'本研究产生或既有复用','target_eligible':'是否计入目标数据','official_url':'官方链接','public_status':'官方公开状态','data_access_status':'数据获取状态','official_checked_at':'官方核验时间','official_title':'官方题名','library_strategy':'官方library strategy','sample_accession':'样本编号','experiment_accession':'实验编号','run_accession':'Run编号','sample_treatment_mapping':'处理与样本对应关系','sample_count':'样本数','mapping_certainty':'对应确定程度','evidence_ids':'文章或附件证据ID','verification_notes':'核验备注'}
EVIDENCE_FIELDS={'evidence_id':'证据ID','record_id':'文献ID','evidence_kind':'证据类型','excerpt':'简短原文证据','interpretation':'上下文判断','source_path':'原件路径','source_sha256':'源SHA256','pdf_pages':'PDF页码（从1起）','section':'章节或印刷页码','worksheet':'工作表','cells':'单元格','batch_id':'批次'}
PENDING_FIELDS={'record_id':'文献ID','issue_type':'问题或线索类型','path':'文件路径','sha256':'源SHA256','pdf_pages':'PDF页码','raw_accession':'原样编号','normalized_accession':'规范编号','accession_level':'编号层级','database':'数据库','official_url':'官方链接','reason':'问题说明','action':'后续处理','resolution_status':'解决状态','resolved_in_batch':'解决批次','resolution':'解决依据','resolution_evidence_id':'解决证据ID'}

def main():
    rows=read('screening_records.json');evidence=read('evidence.json');projects=read('data_projects_samples.json');pending=read('pending_related.json')
    assert len(rows)==len({r['record_id'] for r in rows})
    done=[r for r in rows if r['progress']=='已完成'];targets=[r for r in done if r['conclusion'] in ['确认宏基因组','确认宏转录组','确认二者兼有']]
    remaining=[r for r in rows if r['progress']=='未开始'];states=collections.Counter(r['progress'] for r in rows)
    target_ids={r['record_id'] for r in targets};eligible_projects=[p for p in projects if p.get('target_eligible') is True and p['record_id'] in target_ids]
    project_keys={p['dataset_key'] for p in eligible_projects if p.get('dataset_key')}
    public_ids={p['record_id'] for p in eligible_projects if p.get('data_access_status')=='1 已定位编号，官方确认公开'}
    request_ids={r['record_id'] for r in targets if r['data_access_status']=='4 明确声明相关测序数据可向作者索取'}
    processed=[r for r in rows if r['batch_id'] and r['completed_at']]
    newly_identified_reviews=[r for r in processed if r['progress']=='综述不适用']
    summary={'scope_records':len(rows),'eligible_research_queue_initial':read('baseline.json')['initial_eligible_research_records'],'eligible_research_queue_after_type_verification':read('baseline.json')['initial_eligible_research_records']-len(newly_identified_reviews),'progress_counts':dict(states),'completed_articles':len(done),'processed_records_including_new_reviews':len(processed),'new_reviews_identified':len(newly_identified_reviews),'target_articles':len(targets),'metagenomics_only_articles':sum(r['conclusion']=='确认宏基因组' for r in targets),'metatranscriptomics_only_articles':sum(r['conclusion']=='确认宏转录组' for r in targets),'both_articles':sum(r['conclusion']=='确认二者兼有' for r in targets),'public_target_data_articles':len(public_ids),'explicit_sequencing_data_on_request_articles':len(request_ids),'independent_target_projects':len(project_keys),'remaining_unstarted':len(remaining),'next_batch_start':remaining[0]['record_id'] if remaining else None,'next_batch_default_size':20,'batch_ids':sorted({r['batch_id'] for r in processed}),'last_completion_at_utc':max((r['completed_at'] for r in processed),default=''),'attachment_identity_pending_objects':sum(p['issue_type']=='附件唯一归属待核实' and p.get('resolution_status')!='已解决' for p in pending),'screening_uncertainty_articles':sum(r['progress']=='待确认' for r in rows),'note':'未完成、缺正文及综述不计阴性；本批零命中不代表剩余队列没有目标测序。'}
    dump('summary.json',summary)
    batches=[]
    for bid in summary['batch_ids']:
        br=[r for r in done if r['batch_id']==bid];pr=[r for r in processed if r['batch_id']==bid];reviews=[r for r in pr if r['progress']=='综述不适用']
        batches.append({'batch_id':bid,'completed_ids':[r['record_id'] for r in br],'completed_count':len(br),'processed_record_ids':[r['record_id'] for r in pr],'processed_record_count':len(pr),'review_not_applicable_ids':[r['record_id'] for r in reviews],'review_not_applicable_count':len(reviews),'target_articles':sum(r['record_id'] in target_ids for r in br),'metagenomics_only':sum(r['conclusion']=='确认宏基因组' for r in br),'metatranscriptomics_only':sum(r['conclusion']=='确认宏转录组' for r in br),'both':sum(r['conclusion']=='确认二者兼有' for r in br),'public_data_articles':sum(r['record_id'] in public_ids for r in br),'explicit_request_articles':sum(r['record_id'] in request_ids for r in br),'status':'已完成','note':'按稳定ID顺序，逐篇阅读方法；实际综述单列后顺延补足研究篇数。' if reviews else '按稳定ID顺序，逐篇阅读方法；扫描件按需OCR/原页核查。'})
    dump('batches.json',batches)
    for batch in batches:
        csvout(f"batch_{batch['batch_id']}_results_utf8.csv",[r for r in rows if r['batch_id']==batch['batch_id']],list(RECORD_FIELDS))
    for name,data,fields in [('screening_records_utf8.csv',rows,list(rows[0])),('target_studies_utf8.csv',targets,list(RECORD_FIELDS)),('data_projects_samples_utf8.csv',projects,list(PROJECT_FIELDS)),('evidence_utf8.csv',evidence,list(EVIDENCE_FIELDS)),('pending_related_utf8.csv',pending,list(PENDING_FIELDS))]:csvout(name,data,fields)
    wb=Workbook();wb.remove(wb.active)
    def sheet(name,data,fields):
        ws=wb.create_sheet(name);ws.append(list(fields.values()))
        for r in data:ws.append([cell(r.get(k,'')) for k in fields])
        ws.freeze_panes='C2';ws.auto_filter.ref=ws.dimensions;ws.row_dimensions[1].height=30
        for c in ws[1]:c.fill=PatternFill('solid',fgColor='17365D');c.font=Font(color='FFFFFF',bold=True);c.alignment=Alignment(vertical='center',wrap_text=True)
        for i,(key,label) in enumerate(fields.items(),1):
            ws.column_dimensions[get_column_letter(i)].width=100 if key=='rule' else (65 if key in ['title','excerpt','interpretation','method_summary','reason','source_path','path','fulltext_paths','reviewed_locations','value'] else 25)
        for row in ws.iter_rows(min_row=2):
            ws.row_dimensions[row[0].row].height=90 if name=='证据记录' else 60
            for c in row:c.alignment=Alignment(vertical='top',wrap_text=True)
        return ws
    sheet('筛查总表',rows,RECORD_FIELDS);sheet('目标测序研究',targets,RECORD_FIELDS);sheet('数据项目与样本',projects,PROJECT_FIELDS);sheet('证据记录',evidence,EVIDENCE_FIELDS);sheet('待确认与相关线索',pending,PENDING_FIELDS)
    stat=[]
    for k,v in [('范围内全部任务',len(rows)),('可筛查研究初始队列',summary['eligible_research_queue_initial']),('实际类型核验后研究队列',summary['eligible_research_queue_after_type_verification']),('累计已核查记录（含新识别综述）',len(processed)),('累计已完成研究',len(done)),('累计目标测序文章（兼有只计一次）',len(targets)),('仅宏基因组文章',summary['metagenomics_only_articles']),('仅宏转录组文章',summary['metatranscriptomics_only_articles']),('二者兼有文章',summary['both_articles']),('官方确认公开目标数据文章',len(public_ids)),('明确可索取相关测序数据文章',len(request_ids)),('独立目标项目数（按dataset_key去重）',len(project_keys)),('待确认研究',summary['screening_uncertainty_articles']),('未开始研究',len(remaining)),('缺正文跳过（研究）',states['缺正文跳过']),('综述不适用（含4条缺正文综述）',states['综述不适用']),('附件唯一归属待核实件数',summary['attachment_identity_pending_objects']),('下一批起点',summary['next_batch_start']),('下一批默认篇数',20)]:stat.append({'metric':k,'value':v,'batch':'累计','note':summary['note'] if k=='累计目标测序文章（兼有只计一次）' else ''})
    for b in batches:
        stat.append({'metric':'本批完成ID','value':', '.join(b['completed_ids']),'batch':b['batch_id'],'note':b['note']})
        for key in ['completed_count','processed_record_count','review_not_applicable_ids','target_articles','metagenomics_only','metatranscriptomics_only','both','public_data_articles','explicit_request_articles']:stat.append({'metric':key,'value':b[key],'batch':b['batch_id'],'note':''})
    sheet('批次进度与统计',stat,{'metric':'指标','value':'数值或内容','batch':'口径或批次','note':'说明'})
    notes=[{'field':'筛查口径','rule':'范围按现行完整分类名；不限制期刊/JCR；学位和其他研究如有核验正文也进入队列。综述引用不当作自身测序。'}, {'field':'零命中','rule':'目标测序研究/项目表仅列真实阳性；零记录表表示累计已完成研究没有确认阳性；未完成记录不作阴性。'}, {'field':'正文覆盖','rule':'全范围426条含388有核验正文、38无核验正文；基线综述50条（4条缺正文），B002按原件新增识别R00915为综述，当前51条；34条研究缺正文跳过。'}, {'field':'目标测序','rule':'群落DNA shotgun宏基因组或明确群落RNA宏转录组；既有公开数据复用单列；仅16S/DGGE/功能预测/单菌/引用不纳入；仅病毒组作线索。'}, {'field':'数据状态1–3','rule':'1 已定位编号，官方确认公开；2 文中给出编号，但访问或公开状态未核实；3 已定位项目，但数据未公开、受限或处于保留状态。'}, {'field':'数据状态4–7','rule':'4 明确声明相关测序数据可向作者索取；5 仅泛指数据可索取，是否含原始测序数据不明确；6 确认使用目标测序但未找到编号或相关声明；7 证据不足，待确认。'}, {'field':'阴性数据状态','rule':'阴性用不适用，不冒用状态6；缺正文或身份冲突不能标未使用目标测序。'}, {'field':'官方核验','rule':'仅对目标或疑似目标项目读取官方数据库元数据，检查题名、样本、library strategy、公开状态和分组；不下载原始reads。不成功保留文中编号且标未核实。本批无目标/疑似编号，未发起目标库请求。'}, {'field':'项目层级','rule':'项目/样本/实验/run/单条核酸记录分开；独立项目以数据库+规范项目号dataset_key跨文章去重，无项目号不推定为独立项目。'}, {'field':'页码和OCR','rule':'PDF页从1起；原印刷页另记section。R00608全5页OCR与当前PDF渲染图逐像素一致，p2–3方法原页视觉复核；科学判断经上下文阅读。'}, {'field':'附件归属','rule':'本轮R01727已核实为同文接受稿，剩60件暂不唯一归属；待题名/DOI/正文引用核实，不据目录分配编号。可检查现有附件，不追下载缺件。'}, {'field':'当前附件限制','rule':'首批5篇均无独立补充材料原件；R00273正文内附录I–V已读，R00313正文提及Supplement Tables 1–2但未取得，方法证据明确为16S。'}, {'field':'续跑','rule':'读取已提交screening_records.json，取未开始项按稳定ID顺序20篇；待确认独立复核；逐篇写入evidence/进度后再生成本表。勿运行初始化覆盖累计结果。'}]
    sheet('口径与字段说明',notes,{'field':'字段或事项','rule':'规则与限制'})
    outfile=ROOT/'omics_screening_cumulative.xlsx';wb.save(outfile)
    reopened=load_workbook(outfile,read_only=True,data_only=True)
    assert reopened['筛查总表'].max_row==len(rows)+1 and reopened['证据记录'].max_row==len(evidence)+1
    reopened.close();print(json.dumps(summary,ensure_ascii=False,indent=2))

if __name__=='__main__':main()
