"""Cumulative projects/runs/papers export. No reads or originals are downloaded."""
import collections,csv,hashlib,json,re,xml.etree.ElementTree as ET
from pathlib import Path
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Font,PatternFill
from common import ROOT,dump,now
from collect_projects import parse_sra

def read(name,default):
 p=ROOT/name;return json.loads(p.read_text()) if p.exists() else default
def csvout(name,rows):
 fields=list(dict.fromkeys(k for r in rows for k in r)) or ['id','status']
 with (ROOT/name).open('w',encoding='utf-8-sig',newline='') as f:
  w=csv.DictWriter(f,fieldnames=fields);w.writeheader()
  for r in rows:w.writerow({k:json.dumps(v,ensure_ascii=False) if isinstance(v,(list,dict)) else v for k,v in r.items()})
def cls(x,known,settings):
 if settings.get('assay_exclusion'):return settings['assay_exclusion']
 strategy=(x.get('library_strategy') or '').upper();source=(x.get('library_source') or '').upper();sci=(x.get('scientific_name') or '').lower()
 if strategy=='AMPLICON':return '非目标：扩增子'
 if x['run_accession'] in known:return '群落宏基因组：既有全文方法及官方run联合确认'
 if source=='METATRANSCRIPTOMIC' and strategy in ['RNA-SEQ','OTHER','WGS']:return '群落宏转录组候选'
 if source=='METAGENOMIC' and strategy in ['WGS','METAGENOMIC','OTHER','WGA']:return '群落宏基因组候选'
 if source=='GENOMIC' and strategy=='WGS' and 'metagenom' in sci:return '群落宏基因组候选：GENOMIC标签，项目/物种需联合核对'
 if source=='TRANSCRIPTOMIC' and 'metagenom' in sci and strategy=='RNA-SEQ':return '群落宏转录组候选：source需复核'
 if 'VIRAL' in source or 'virome' in sci:return '特殊：病毒富集/病毒组'
 if source in ['GENOMIC','TRANSCRIPTOMIC'] and 'metagenom' not in sci:return '特殊：纯培养/非群落待核'
 return '类型待核'
def mapping(p,x,old):
 t=x.get('sample_alias') or '';attrs={v['tag']:v['value'] for v in x.get('sample_attributes',[])};r={'treatment_label':None,'timepoint':None,'time_unit':None,'biological_replicate':None,'experimental_unit_id':None,'oxygen_exposure':None,'mapping_status':'未核实处理/样本映射','core_comparison_eligible':False,'sample_attributes':x.get('sample_attributes',[]),'prior_mapping_evidence':sorted({a.get('sample_treatment_mapping','') for a in old if a.get('run_accession')==x['run_accession']})}
 if p=='PRJNA1013451':
  label=attrs.get('source_material_id') or x.get('experiment_title','');r['treatment_label']=label
  n=re.search(r'JL\.(\d+)',t);core=n and int(n[1])<=12;tm=re.search(r'week\s*(\d+)',label,re.I)
  if tm:r.update(timepoint=int(tm[1]),time_unit='week')
  if core:
   g='MST' if 'air' in label.lower() and 'sound' in label.lower() else ('MAT' if 'air' in label.lower() else ('SND' if 'sound' in label.lower() else 'CON'));r.update(treatment_label=g,experimental_unit_id=g,biological_replicate='每处理1反应器，非独立重复',mapping_status='官方source_material_id+既有SI键明确',core_comparison_eligible=True,oxygen_exposure=(g in ['MAT','MST'] and r['timepoint']!=6))
  else:r['mapping_status']='项目额外样本，非主比较；保留进料/初始/擦拭信息'
 elif p=='PRJNA1287080' and t in ['R1-M','R4-M']:r.update(treatment_label=t,timepoint=40,time_unit='day',oxygen_exposure=t=='R4-M',mapping_status='官方alias与论文R1/R4处理明确对应',core_comparison_eligible=True)
 elif p=='PRJNA1447957':
  m=re.fullmatch(r'([A-Z]+)(\d+)',t)
  if m:r.update(treatment_label=m[1]+'（原始代号，氧/对照键待核）',timepoint=int(m[2]),time_unit='day',mapping_status='温度/时间可核，ZN/ZZ/CN/CZ与MN/MD/PN/PD对照键未核')
 elif p=='PRJNA565588':r.update(treatment_label=attrs.get('isolation_source'),oxygen_exposure=('microaerat' in attrs.get('isolation_source','').lower()),mapping_status='Rep1/Rep2明确，但论文采样日与官方alias日标签不同，缺厌氧MG对照')
 elif p=='PRJNA357876':r.update(treatment_label=t,mapping_status='SM/MP/TP与既有论文对应；无供氧—厌氧匹配比较')
 elif p=='PRJNA1086167':r.update(treatment_label=t,mapping_status='官方alias与论文给定SRR处理顺序冲突，不能自动纠正')
 elif p=='PRJNA1128434':r.update(treatment_label=t,timepoint=30,time_unit='day',mapping_status='A/B/C代号可核，DW/CO2NW/O2NW处理键未核')
 override=read('sample_annotations.json',{}).get(x['run_accession']) or read('sample_annotations.json',{}).get(p+'|'+x.get('sample_accession',''))
 if override:r.update(override)
 return r
def render():
 papers=read('existing_articles_availability_audit.json',[])+read('new_associated_papers.json',[]);links=read('existing_project_links.json',{});ass=read('project_assessments.json',{});old=read('existing_dataset_rows.json',[]);known={x['run_accession'] for x in old if x.get('run_accession')};discovered=read('discovered_projects.json',[]);allacc=set(links)|{x['accession'] for x in discovered}|set(ass);runs=[];projects=[];aliases=[]
 for acc in sorted(allacc):
  p=ROOT/'project_cache'/acc;status=read(str(p.relative_to(ROOT)/'collection_status.json'),{});ena=read(str(p.relative_to(ROOT)/'ena_runs.json'),[]);ncbi=read(str(p.relative_to(ROOT)/'ncbi_sra_packages.json'),[]);byrun={x['run_accession']:dict(x) for x in ncbi}
  for e in ena:
   key=e['run_accession']
   if key not in byrun:byrun[key]=dict(e)
   row=byrun[key];row['ENA_fields']=e
   for k,v in e.items():
    if k not in row or not row[k]:row[k]=v
  prs=[]
  for x in byrun.values():
   en=x.get('ENA_fields',{});typ=cls(x,known,ass.get(acc,{}));target=typ.startswith('群落');entry={'project_id':acc,'study_id':ass.get(acc,{}).get('study_id') or 'OST-'+acc,'linked_paper_ids':sorted({d['record_id'] for d in links.get(acc,[])})+sorted({a['paper_id'] for a in papers if acc in a.get('project_accessions',[]) and a.get('newly_discovered_paper')}),'run_accession':x['run_accession'],'experiment_accession':x.get('experiment_accession'),'BioSample_accession':x.get('sample_accession'),'secondary_sample_accession':x.get('secondary_sample_accession'),'sample_alias':x.get('sample_alias'),'sample_title':x.get('sample_title'),'experiment_title':x.get('experiment_title'),'scientific_name':x.get('scientific_name'),'library_strategy':x.get('library_strategy'),'library_source':x.get('library_source'),'library_selection':x.get('library_selection'),'library_construction_protocol':x.get('library_construction_protocol'),'library_layout':x.get('library_layout'),'platform':x.get('instrument_platform'),'instrument_model':x.get('instrument_model'),'sequencing_class':typ,'target_community_reads':target,'ENA_library_strategy':en.get('library_strategy'),'ENA_library_source':en.get('library_source'),'ENA_library_selection':en.get('library_selection'),'ENA_read_count_raw':en.get('read_count') or None,'NCBI_total_spots_raw':x.get('read_count') or None,'base_count':en.get('base_count') or x.get('base_count'),'count_unit_warning':'ENA read_count/SRA total_spots保留原字段；PAIRED时不能直接当单条read总数','fastq_urls':['https://'+u for u in en.get('fastq_ftp','').split(';') if u],'fastq_bytes':[int(v) for v in en.get('fastq_bytes','').split(';') if v.isdigit()],'fastq_md5':[v for v in en.get('fastq_md5','').split(';') if v],'submitted_urls':['https://'+u for u in en.get('submitted_ftp','').split(';') if u],'submitted_formats':en.get('submitted_format'),'SRA_file_records':x.get('sra_files',[]),'file_entry_status':'官方清单给出；未下载reads，未校验实际文件MD5','first_public':en.get('first_public'),'metadata_source_paths':['project_cache/'+acc+'/ena_runs.json','project_cache/'+acc+'/ncbi_sra_packages.json'],'mapping':mapping(acc,x,old),'raw_reads_downloaded':False}
   entry.update(entry.pop('mapping'));prs.append(entry)
  prs.sort(key=lambda x:x['run_accession']);cnt=collections.Counter(x['BioSample_accession'] for x in prs)
  for x in prs:x['runs_same_BioSample']=cnt[x['BioSample_accession']];x['replicate_warning']='同一BioSample多run/实验不算独立生物重复；分lane/技术重复具体关系待核' if x['runs_same_BioSample']>1 else '一run不自动证明一个独立生物重复'
  runs+=prs;a=ass.get(acc,{});target=[x for x in prs if x['target_community_reads']];projectpapers=sorted({d['record_id'] for d in links.get(acc,[])})+sorted({x['paper_id'] for x in papers if x.get('newly_discovered_paper') and acc in x.get('project_accessions',[])});design=[x for x in papers if x['paper_id'] in projectpapers];types=sorted(set(x['sequencing_class'] for x in prs));row={'project_id':acc,'study_id':ass.get(acc,{}).get('study_id') or 'OST-'+acc,'independent_study_status':'项目/run/BioSample/MD5去重；项目号本身不证明实验独立；跨项目关系见去重表','linked_paper_ids':projectpapers,'new_project_relative_to_existing_omics':acc not in links,'tier':a.get('tier'),'tier_rationale':a.get('rationale'),'review_status':'已审计' if a else '未开始项目核实','availability_code':a.get('availability_code'),'availability':a.get('availability') or '官方检索候选，未完成处理/文件/类型核实','official_title':status.get('project_title') or next((x.get('study_title') for x in ncbi+ena if x.get('study_title')),None),'official_description':status.get('project_description') or next((x.get('study_description') for x in ncbi if x.get('study_description')),None),'data_types':types,'all_runs_count':len(prs) if status else None,'target_runs_count':len(target) if status else None,'target_BioSamples_count':len({x['BioSample_accession'] for x in target}) if status else None,'BioSample_is_not_biological_replicate':'归档样本数独立列；混样、技术重复和重复归档未核实者不算生物重复','confirmed_biological_specimen_count':a.get('confirmed_biological_specimen_count'),'confirmed_experimental_units':a.get('confirmed_experimental_units'),'biological_replicates':a.get('biological_replicates'),'mapping_status':a.get('mapping_status'),'oxygen_comparisons_confirmed':a.get('oxygen_comparisons_confirmed'),'comparison_definition':a.get('comparison_definition'),'target_fastq_files_listed':sum(len(x['fastq_urls']) for x in target),'target_fastq_bytes_listed':sum(sum(x['fastq_bytes']) for x in target) if target and all(x['fastq_bytes'] for x in target) else None,'base_count_listed':sum(int(x['base_count']) for x in target) if target and all(str(x.get('base_count','')).isdigit() for x in target) else None,'target_runs_missing_base_count':sum(not str(x.get('base_count','')).isdigit() for x in target),'paired_omics_group':a.get('paired_omics_group'),'project_aliases_official':a.get('project_aliases_official',[]),'matched_DNA_RNA_same_BioSample_count':len({x['BioSample_accession'] for x in target if '宏基因组' in x['sequencing_class']}&{x['BioSample_accession'] for x in target if '宏转录组' in x['sequencing_class']}),'oxygen_designs_reused':[{'paper_id':x['paper_id'],'oxygen':x.get('oxygen_supply'),'substrate':x.get('substrate'),'treatment_control':x.get('treatment_control')} for x in design],'ARG_MGE_feasibility':'群落总DNA候选可重注释；核建库偏向、宿主DNA、深度及分组。短reads共现不是ARG宿主/MGE传播因果' if any('宏基因组' in x['sequencing_class'] for x in target) else '目标DNA未核实；RNA不能替代DNA丰度/基因存在性','virus_host_feasibility':'总DNA可探索DNA病毒/前噬菌体，但非专门病毒组；深度/组装完整性/宿主匹配须实测，不保证低丰度病毒或RNA病毒覆盖' if any('宏基因组' in x['sequencing_class'] for x in target) else '依类型另评：宏转录组不等于完整病毒组，纯培养/富集需分列','limitations':a.get('limitations'),'priority':a.get('priority') or '待筛，指标不替代设计核查','public_files_listed_not_download_verified':bool(target and all(x['fastq_urls'] for x in target)),'oxygen_delivery_and_dose_as_recorded':a.get('oxygen_delivery_and_dose_as_recorded'),'oxygen_time_and_frequency':a.get('oxygen_time_and_frequency'),'temperature':a.get('temperature'),'reactor_or_environment':a.get('reactor_or_environment'),'substrate':a.get('substrate'),'official_metadata_checked_utc':status.get('checked_utc'),'NCBI_SRA_search_count':status.get('ncbi_sra_count'),'ENA_run_count':status.get('ena_run_count'),'official_metadata_cache':'project_cache/'+acc,'raw_reads_downloaded':False}
  projects.append(row)
  for alias in sorted({x.get('secondary_study_accession') for x in ncbi+ena if x.get('secondary_study_accession')}):aliases.append({'alias':alias,'canonical_project':acc,'relation':'同一项目SRA/ENA secondary study，不计第二项目'})
 # File entry manifest, not a success/download manifest.
 files=[]
 for x in runs:
  for i,u in enumerate(x['fastq_urls']):files.append({'project_id':x['project_id'],'run_accession':x['run_accession'],'BioSample_accession':x['BioSample_accession'],'target_community_reads':x['target_community_reads'],'core_comparison_eligible':x['core_comparison_eligible'],'url':u,'bytes_as_listed':x['fastq_bytes'][i] if i<len(x['fastq_bytes']) else None,'md5_as_listed':x['fastq_md5'][i] if i<len(x['fastq_md5']) else None,'download_status':'未下载','checksum_verified':False})
 duplicates=[];md5=collections.defaultdict(list)
 for f in files:
  if f['md5_as_listed']:md5[f['md5_as_listed']].append(f)
 for h,fs in md5.items():
  if len({f['run_accession'] for f in fs})>1:duplicates.append({'type':'相同官方文件MD5','md5':h,'records':fs,'interpretation':'重复文件线索，不作独立生物重复；未实际下载核验'})
 aliases += [{'alias':'R09236 + R12985','canonical_project':'PRJNA1013451','relation':'两论文使用同21 BioSample/同21run；同研究/数据复用，不能重复累计'}]
 pending=[]
 for a in papers:
  if not a.get('project_accessions') or a.get('availability_code') in [2,3,4,5]:pending.append({'paper_id':a['paper_id'],'project_accessions':a.get('project_accessions'),'availability_code':a.get('availability_code'),'gap':a.get('availability'),'verbatim_statement_evidence':a.get('statements_with_locations'),'next_action':'核官方编号/完整样本键表或明确reads索取范围；不代联系作者','author_contacted':False})
 study_dedupe=aliases+duplicates;comparisons=read('confirmed_comparisons.json',[])
 stats={'generated_utc':now(),'existing_target_articles':22,'existing_MG_articles':21,'existing_MT_articles':1,'existing_project_accessions':9,'existing_projects_currently_public_target_reads':sum(p['project_id'] in links and (p['target_runs_count'] or 0)>0 for p in projects),'existing_article_linked_target_runs':len({x['run_accession'] for x in runs if x['run_accession'] in known}),'newly_found_community_runs_within_existing_projects':len({x['run_accession'] for x in runs if x['project_id'] in links and x['target_community_reads'] and x['run_accession'] not in known}),'existing_project_all_target_runs':len({x['run_accession'] for x in runs if x['project_id'] in links and x['target_community_reads']}),'existing_target_BioSamples':len({x['BioSample_accession'] for x in runs if x['project_id'] in links and x['target_community_reads']}),'reviewed_projects':sum(p['review_status']=='已审计' for p in projects),'unreviewed_project_candidates':sum(p['review_status']!='已审计' for p in projects),'new_project_candidates':sum(p['new_project_relative_to_existing_omics'] for p in projects),'new_reviewed_projects':sum(p['new_project_relative_to_existing_omics'] and p['review_status']=='已审计' for p in projects),'new_associated_papers':sum(p.get('newly_discovered_paper',False) for p in papers),'confirmed_oxygen_control_comparisons':sum(p['oxygen_comparisons_confirmed'] or 0 for p in projects if p['review_status']=='已审计'),'comparison_count_warning':'多个比较可共享同一研究/反应器/时间；不当独立生物重复','metadata_indicated_community_candidate_runs':len({x['run_accession'] for x in runs if x['target_community_reads']}),'metadata_indicated_target_BioSamples':len({x['BioSample_accession'] for x in runs if x['target_community_reads']}),'reviewed_public_community_projects':sum(p['review_status']=='已审计' and (p['target_runs_count'] or 0)>0 and p['availability_code'] not in [6,7] for p in projects),'reviewed_new_public_community_projects':sum(p['project_id'] not in links and p['review_status']=='已审计' and (p['target_runs_count'] or 0)>0 and p['availability_code'] not in [6,7] for p in projects),'reviewed_project_community_runs':len({x['run_accession'] for x in runs if x['target_community_reads'] and x['project_id'] in ass}),'comparisons_with_reused_fulltext_methods':sum(x.get('design_evidence_level')=='既有全文方法+官方样本映射' for x in comparisons),'comparisons_metadata_only_pending_methods':sum(x.get('design_evidence_level')!='既有全文方法+官方样本映射' for x in comparisons),'confirmed_biological_specimens_known_subset':sum(p['confirmed_biological_specimen_count'] or 0 for p in projects),'projects_biological_specimen_count_unknown':sum(p['confirmed_biological_specimen_count'] is None for p in projects),'independent_study_count':'未作所有项目实验独立性的强制总和；已确认PRJNA1013451的两篇论文同研究，官方别名/文件重复另表','raw_FASTQ_downloads':0,'original_paper_redownloads':0,'mechanism_review_paused':True}
 for name,rows in [('data_projects_master',projects),('samples_runs_mapping',runs),('associated_papers',papers),('request_or_pending',pending),('dedup_relationships',study_dedupe),('raw_file_entry_manifest',files)]:dump(ROOT/(name+'.json'),rows);csvout(name+'_utf8.csv',rows)
 dump(ROOT/'summary.json',stats)
 wb=Workbook();wb.remove(wb.active)
 for name,rows in [('数据项目总表',projects),('样本与run映射',runs),('关联论文表',papers),('可索取与待确认',pending),('去重与别名',study_dedupe),('检索记录',read('search_log.json',[])),('分析可行性汇总',[stats]),('文件入口清单',files),('明确氧对照比较',comparisons)]:
  ws=wb.create_sheet(name);fields=list(dict.fromkeys(k for r in rows for k in r)) or ['id','status'];ws.append(fields)
  for row in rows:
   vals=[]
   for k in fields:
    v=row.get(k);v=json.dumps(v,ensure_ascii=False) if isinstance(v,(dict,list)) else v
    if isinstance(v,str):v=v[:32700];v="'"+v if v.startswith(('=','+','@')) else v
    vals.append(v)
   ws.append(vals)
  ws.freeze_panes='C2';ws.auto_filter.ref=ws.dimensions
  for c in ws[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='244766')
  for col in ws.columns:ws.column_dimensions[col[0].column_letter].width=38
 wb.save(ROOT/'omics_dataset_discovery_cumulative.xlsx');wb.close()
 ris=[]
 for p in papers:
  ris+=['TY  - JOUR','ID  - '+p['paper_id'],'TI  - '+p['title']]
  for tag,key in [('PY','year'),('JO','journal'),('DO','doi')]:
   if p.get(key):ris.append(tag+'  - '+str(p[key]))
  ris+=['N1  - '+';'.join(p.get('project_accessions',[]))+'；项目可用性见累计表，不等于已下载reads','ER  -','']
 (ROOT/'associated_papers.ris').write_text('\n'.join(ris),encoding='utf-8')
 check=load_workbook(ROOT/'omics_dataset_discovery_cumulative.xlsx',read_only=True);assert check['数据项目总表'].max_row-1==len(projects);assert check['样本与run映射'].max_row-1==len(runs);check.close();assert len({x['project_id'] for x in projects})==len(projects);assert len({x['run_accession'] for x in runs})==len(runs)
 print(json.dumps(stats,ensure_ascii=False))
if __name__=='__main__':render()
