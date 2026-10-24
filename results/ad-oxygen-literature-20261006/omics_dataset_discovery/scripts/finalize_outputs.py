"""Refresh exports, separated analysis subsets, provenance, queues and validation."""
import json,collections,re,hashlib
from pathlib import Path
from openpyxl import load_workbook
from common import ROOT,dump,now
from render import render,csvout
def read(n):return json.loads((ROOT/n).read_text())
def export(n,x):dump(ROOT/(n+'.json'),x);csvout(n+'_utf8.csv',x)
def finalize():
 # Register paper-origin accessions without pretending these are screened projects.
 registry={x['accession']:x for x in read('discovered_projects.json')}
 for p in read('new_associated_papers.json'):
  for acc in p.get('project_accessions',[]):
   if not re.fullmatch(r'PRJ(?:NA|EB|DB)\d+',acc):continue
   v=registry.setdefault(acc,{'accession':acc,'source_queries':[],'study_title':None,'run_accessions_discovered':[],'raw_source_paths':[]})
   tag='paper_accession:'+p['paper_id']
   if tag not in v['source_queries']:v['source_queries'].append(tag)
 dump(ROOT/'discovered_projects.json',list(registry.values()))
 ass=read('project_assessments.json')
 ass['PRJNA565588']['mapping_status']='4WGS中新增2个厌氧样本；microaeration/anaerobic标签可核，但A150/B79与A192/B147取样日不同，尚无匹配时间对照'
 ass['PRJNA565588']['limitations']='官方B79别名与B97原文件名冲突、SINGLE/PAIRED及取样日差异待论文键核；不忽略新增厌氧WGS，也不直接认定氧效应'
 for chk in read('assembly_status_checks.json'):
  a=ass[chk['project_id']];a['derived_assemblies_official_count']=chk['reported_count'];a['availability']='仅定位公开组装/派生数据；当前项目未查到目标raw run，其他关联项目仍待追踪';a['derived_data_evidence']='assembly_status_checks.json';a['limitations']=(a.get('limitations') or '').split('；官方Assembly已查到')[0]+'；官方Assembly已查到'+str(chk['reported_count'])+'条（含可能的GenBank/RefSeq同组装别名，不当独立样本）。'
 dump(ROOT/'project_assessments.json',ass)
 logs=read('search_log.json')
 for e in logs:
  if e.get('query_id')=='P06':e['interpretation_warning']='NCBI未识别引号短语，实际扩展为methane×oxidation×community×function；2125条完整保存，宽泛扩展记录不当范围筛选完成';e['screening_role']='低优先背景扩展，未完成筛查'
  if e.get('query_id')=='E05':e.pop('failure',None);e['page_limit']=0
 if not any(x.get('query_id')=='GSA-availability' for x in logs):logs.append({'query_id':'GSA-availability','database':'GSA','query':None,'reported_count':None,'complete':False,'status':'仅入口连通性测试；代理403，未执行GSA数据库查询','raw_source':'raw/GSA/c718d2bbdeae064bd679e23dac7323bf82630547c44b43f3585e1925035d17e7.body'})
 dump(ROOT/'search_log.json',logs);render()
 projects=read('data_projects_master.json');pm={p['project_id']:p for p in projects};runs=read('samples_runs_mapping.json');papers=read('associated_papers.json');files=read('raw_file_entry_manifest.json');stat=read('summary.json');comparisons=read('confirmed_comparisons.json')
 reviewed=[p for p in projects if p['review_status']=='已审计'];public=[p for p in reviewed if p['target_runs_count'] and p['availability_code'] not in [6,7]];pubids={p['project_id'] for p in public};main=[p for p in public if p['tier'] in ['A','B']];background=[p for p in public if p['tier']=='C'];special=[p for p in reviewed if p['availability_code']==7];pending_projects=[p for p in reviewed if p not in public and p not in special]
 methods_count=sum('全文方法' in c['design_evidence_level'] for c in comparisons)
 for p in projects:
  p['potential_oxygen_control_comparisons']=p['oxygen_comparisons_confirmed'];p['comparisons_fulltext_method_and_mapping']=sum(c['project_id']==p['project_id'] and '全文方法' in c['design_evidence_level'] for c in comparisons);p['oxygen_comparisons_confirmed']=p['comparisons_fulltext_method_and_mapping']
 export('data_projects_master',projects)
 for p in papers:
  rr=[pm[a] for a in p.get('project_accessions',[]) if a in pm];ids=[a for a in p.get('project_accessions',[]) if a in pubids]
  p['official_page']='https://pubmed.ncbi.nlm.nih.gov/'+str(p['pmid'])+'/' if p.get('pmid') else 'https://doi.org/'+str(p.get('doi',''))
  p['public_community_project_links']=ids
  p['dataset_scope_status']='关联已审计公开群落项目；论文具体方法仍按阅读深度' if ids else ('既有目标文章：数据状态见原22篇审计' if p['paper_id'] in {x['paper_id'] for x in read('existing_articles_availability_audit.json')} else '仅关联非目标/特殊数据，不纳入群落主集合' if rr and all(x['availability_code']==7 for x in rr) else '关联项目/文章范围尚待核实，非纳入结论')
 export('associated_papers',papers)
 # Preserve all links while providing a separate RIS without host/pure-culture papers.
 subset=[p for p in papers if p.get('public_community_project_links') or p['paper_id'] in {x['paper_id'] for x in read('existing_articles_availability_audit.json')}]
 def ris(ps):
  out=[]
  for p in ps:
   out+=['TY  - JOUR','ID  - '+p['paper_id'],'TI  - '+p['title']]
   for tag,k in [('PY','year'),('JO','journal'),('DO','doi'),('UR','official_page')]:
    if p.get(k):out.append(tag+'  - '+str(p[k]))
   for a in p.get('authors',[]):out.append('AU  - '+str(a))
   out+=['N1  - '+p['dataset_scope_status']+'；'+','.join(p.get('project_accessions',[])),'ER  -','']
  return '\n'.join(out)
 (ROOT/'associated_papers.ris').write_text(ris(papers),encoding='utf-8');(ROOT/'community_dataset_papers.ris').write_text(ris(subset),encoding='utf-8')
 # Official aliases and same-experiment cohorts remain distinct from mere shared titles.
 dedup=read('dedup_relationships.json')
 for p in reviewed:
  for alias in p.get('project_aliases_official',[]):dedup.append({'type':'官方项目别名','alias':alias,'canonical_project':p['project_id'],'relation':'同一项目；不是新项目/新数据，官方sample属性保留于缓存'})
 groups=collections.defaultdict(list)
 for p in public:groups[p['study_id']].append(p['project_id'])
 for sid,accs in groups.items():
  if len(accs)>1:dedup.append({'type':'同实验/队列跨项目线索','study_id':sid,'projects':accs,'relation':'项目不同但不自动计独立研究；处理/物理样本一一对应仍需证据'})
 sruns=collections.defaultdict(list)
 for x in runs:
  if x['target_community_reads']:sruns[x['BioSample_accession']].append(x)
 for b,rs in sruns.items():
  if len({r['project_id'] for r in rs})>1:dedup.append({'type':'跨项目共享BioSample','BioSample':b,'projects':sorted({r['project_id'] for r in rs}),'runs':[r['run_accession'] for r in rs],'relation':'共享归档样本；是否同一物理采样及DNA/RNA配对另核'})
 export('dedup_relationships',dedup)
 # Metadata-ready means a manifest can be downloaded, not download/quality validation.
 ready=[]
 for p in main:
  if p['availability_code']==1:
   cs=[c for c in comparisons if c['project_id']==p['project_id']];rs=[r for r in runs if r['project_id']==p['project_id'] and r['target_community_reads']];core={v for c in cs for v in c['treatment_runs']+c['control_runs']}
   ready.append({'project_id':p['project_id'],'tier':p['tier'],'priority':p['priority'],'all_target_runs':p['target_runs_count'],'selected_comparison_runs':len(core),'comparison_ids':[c['comparison_id'] for c in cs],'design_evidence':'已读全文方法+官方映射' if any('全文方法' in c['design_evidence_level'] for c in cs) else '官方元数据暂定，须补方法','download_manifest_entries':sum(f['project_id']==p['project_id'] and f['run_accession'] in core for f in files),'current_environment_transfer_status':'ENA FASTQ域名HEAD被代理403；未下载、未校验reads','caveat':p['limitations'],'raw_download_authorized_this_round':False})
 export('download_preparation_priority',ready)
 selected_runs={r for c in comparisons for r in c['treatment_runs']+c['control_runs']};download_files=[dict(f,selection_status='下一阶段候选清单；本轮未下载FASTQ') for f in files if f['run_accession'] in selected_runs];export('download_preparation_file_manifest',download_files)
 # Missing information is tracked at project as well as article level.
 pending=read('request_or_pending.json')
 for p in reviewed:
  if p['availability_code'] in [2,3,6] or p in main:
   pending.append({'project_id':p['project_id'],'availability_code':p['availability_code'],'gap':p['mapping_status'] or p['availability'],'next_action':p['limitations'] or '核建库/样本键/独立重复/氧剂量原单位','official_url':'https://www.ncbi.nlm.nih.gov/bioproject/'+p['project_id'],'author_contacted':False,'evidence_paths':[p['official_metadata_cache']+'/collection_status.json',p['official_metadata_cache']+'/ncbi_sra_packages.json']})
 export('request_or_pending',pending)
 # Stable remaining queue, no quota; score is routing only, not scientific inclusion.
 queue=[]
 for p in projects:
  if p['review_status']=='已审计':continue
  d=registry.get(p['project_id'],{});t=' '.join(str(d.get(k) or '') for k in ['study_title','project_description']).lower();route='LakeWashington多子项目/合成群落，先按研究去重' if ('lake washington' in t or 'methane oxidation as a community' in t or 'synthetic' in t and 'methan' in t) else '优先供氧/对照设计线索' if any(k in t for k in ['anaerobic digestion','microaer','reoxygen','oxygen exposure','redox','oxygen_injection']) else '项目关联/数据类型及范围待筛'
  queue.append({'project_id':p['project_id'],'status':'未开始','route_hint_not_inclusion':route,'study_title':d.get('study_title'),'source_queries':d.get('source_queries'),'matched_runs_in_queries':len(d.get('run_accessions_discovered',[])),'matched_run_count_is_not_project_total':True,'raw_source_paths':d.get('raw_source_paths'),'next_action':'复用现有响应；补全官方项目run/实验/样本与处理键，再人工判断A/B/C或特殊/非目标'})
 queue.sort(key=lambda q:({'优先供氧/对照设计线索':0,'LakeWashington多子项目/合成群落，先按研究去重':1}.get(q['route_hint_not_inclusion'],2),q['project_id']))
 for i,q in enumerate(queue,1):q['stable_order']=i
 export('remaining_project_queue',queue)
 pcandidates=read('paper_discovery_candidates_enriched.json');linkedpmids={str(p.get('pmid')) for p in papers};paperqueue=[dict(p,processing_status='未开始项目关联/范围筛查；完整摘要已缓存，非已读') for p in pcandidates if str(p.get('pmid')) not in linkedpmids];export('remaining_paper_queue',paperqueue)
 evidence=[]
 for p in reviewed:
  for field in ['official_title','official_description']:
   if p.get(field):evidence.append({'project_id':p['project_id'],'source_type':'官方项目元数据，不是论文全文','location':field,'verbatim':p[field],'source_path':p['official_metadata_cache']+'/collection_status.json','official_url':'https://www.ncbi.nlm.nih.gov/bioproject/'+p['project_id']})
 for c in comparisons:evidence.append({'project_id':c['project_id'],'source_type':c['design_evidence_level'],'location':'比较记录及精确run键','verbatim':None,'comparison':c,'source_path':'confirmed_comparisons.json'})
 for ev in read('reused_cached_fulltext_methods.json'):
  for sec in ev['sections_read']:evidence.append({'project_id':ev['project_id'],'paper_id':ev['paper_id'],'source_type':ev['source_type'],'location':sec['section'],'xml_section_id':sec['xml_section_id'],'verbatim':sec['verbatim'],'source_path':ev['existing_cached_source'],'official_url':ev['official_page']})
 export('dataset_evidence',evidence)
 source=[{'database':'NCBI BioProject/SRA/PubMed/Assembly','status':'官方API实际成功；完整查询、ID分页和响应已缓存','access_test':'request_log.jsonl','raw_reads_downloaded':False},{'database':'ENA','status':'官方search/filereport/project XML实际成功；E05官方limit=0全量1282并核count','access_test':'search_log.json','raw_reads_downloaded':False},{'database':'GSA','status':'入口代理403；未执行平台检索，返回数未知','access_test':'search_log.json:GSA-availability','raw_reads_downloaded':False},{'database':'ENA FASTQ文件服务器','status':'7项目抽检HEAD被环境代理403；仅核官方URL/bytes/MD5列表，文件未下载验证','access_test':'file_entry_HEAD_probes.json','raw_reads_downloaded':False}];export('source_availability',source)
 pruns=[r for r in runs if r['project_id'] in pubids and r['target_community_reads']]
 oxygen_runs=[r for r in pruns if pm[r['project_id']]['tier'] in ['A','B']];background_runs=[r for r in pruns if pm[r['project_id']]['tier']=='C'];non_target=[{k:r[k] for k in ['project_id','run_accession','BioSample_accession','experiment_accession','sequencing_class','library_strategy','library_source','sample_alias']} for r in runs if not r['target_community_reads']]
 export('oxygen_candidate_community_runs',oxygen_runs);export('background_community_runs',background_runs);export('non_target_assay_run_index',non_target);export('oxygen_comparison_records',comparisons);csvout('search_log_utf8.csv',logs)
 stat.update({'total_project_candidates':len(projects),'reviewed_new_public_community_projects':sum(p['new_project_relative_to_existing_omics'] for p in public),'new_public_A_projects':sum(p['tier']=='A' and p['new_project_relative_to_existing_omics'] for p in public),'new_public_B_projects':sum(p['tier']=='B' and p['new_project_relative_to_existing_omics'] for p in public),'new_public_C_projects':sum(p['tier']=='C' and p['new_project_relative_to_existing_omics'] for p in public),'public_A_projects':sum(p['tier']=='A' for p in public),'public_B_projects':sum(p['tier']=='B' for p in public),'public_C_projects':len(background),'public_A_B_projects':len(main),'A_B_public_runs':sum(p['target_runs_count'] for p in main),'C_public_runs':sum(p['target_runs_count'] for p in background),'public_A_B_archive_BioSamples':len({r['BioSample_accession'] for r in pruns if pm[r['project_id']]['tier'] in ['A','B']}),'public_C_archive_BioSamples':len({r['BioSample_accession'] for r in pruns if pm[r['project_id']]['tier']=='C'}),'provisional_archive_study_groups_public':len(groups),'provisional_archive_groups_not_verified_independent_studies':True,'independent_studies_total':None,'biological_samples_total':None,'associated_papers_total_including_pending_special':len(papers),'new_associated_paper_links_relative_to_old22':len(papers)-22,'papers_new_to_existing_literature_library':sum(p.get('newly_discovered_paper',False) for p in papers),'papers_linked_reviewed_public_community_projects':sum(bool(p['public_community_project_links']) for p in papers),'new_papers_linked_reviewed_public_community_projects':sum(bool(p['public_community_project_links']) and p.get('newly_discovered_paper',False) for p in papers),'community_paper_RIS_records':len(subset),'PubMed_cached_unique_paper_metadata':len(pcandidates),'remaining_paper_association_candidates':len(paperqueue),'potential_oxygen_control_comparison_records':len(comparisons),'confirmed_oxygen_control_comparisons':methods_count,'comparisons_with_reused_fulltext_methods':methods_count,'comparisons_metadata_only_pending_methods':len(comparisons)-methods_count,'comparison_count_warning':'12项由已读全文方法+官方样本键支持，另4项元数据暂定；多项共享研究/反应器，绝非16独立研究或有重复统计比较','download_preparation_projects':len(ready),'remaining_project_candidates':len(queue),'confirmed_biological_specimens_known_subset':12,'raw_FASTQ_downloads':0})
 export('analysis_feasibility_by_project',[{k:p[k] for k in ['project_id','study_id','tier','review_status','availability_code','linked_paper_ids','target_runs_count','target_BioSamples_count','confirmed_biological_specimen_count','confirmed_experimental_units','biological_replicates','oxygen_comparisons_confirmed','potential_oxygen_control_comparisons','ARG_MGE_feasibility','virus_host_feasibility','limitations']} for p in reviewed]);dump(ROOT/'summary.json',stat)
 export('public_community_projects',public);export('special_or_non_target_projects',special)
 # Append separated sheets. JSON/CSV retain long full fields; Excel cells capped explicitly.
 wb=load_workbook(ROOT/'omics_dataset_discovery_cumulative.xlsx')
 if '明确氧对照比较' in wb:del wb['明确氧对照比较']
 sheets={'数据项目总表':projects,'供氧比较及证据层级':comparisons,'关联论文表':papers,'可索取与待确认':pending,'去重与别名':dedup,'分析可行性汇总':[stat],'公开群落数据子集':public,'AB候选群落run':oxygen_runs,'C背景群落run':background_runs,'非目标run索引':non_target,'特殊及非目标数据':special,'下载准备优先队列':ready,'剩余项目队列':queue,'逐项目分析可行性':read('analysis_feasibility_by_project.json'),'官方证据与位置':evidence,'来源可用性':source}
 for name,rs in sheets.items():
  if name in wb:del wb[name]
  w=wb.create_sheet(name);keys=list(dict.fromkeys(k for r in rs for k in r)) or ['status'];w.append(keys)
  for r in rs:
   vals=[]
   for k in keys:
    v=r.get(k);v=json.dumps(v,ensure_ascii=False) if isinstance(v,(dict,list)) else v
    if isinstance(v,str):v=v[:32700];v="'"+v if v.startswith(('=','+','@')) else v
    vals.append(v)
   w.append(vals)
  w.freeze_panes='B2';w.auto_filter.ref=w.dimensions
  for col in w.columns:w.column_dimensions[col[0].column_letter].width=34
 wb.save(ROOT/'omics_dataset_discovery_cumulative.xlsx');wb.close()
 dump(ROOT/'progress.json',{'updated_utc':now(),'baseline_commit':read('baseline.json')['source_commit'],'active_task':'omics_dataset_discovery','mechanism_review':'已暂存，不改分类/原判断/原件','completed_batches':['D001','D002','D003','D004','D005'],'reviewed_existing_projects':9,'reviewed_new_projects':60,'remaining_projects':len(queue),'remaining_paper_association_candidates':len(paperqueue),'resume_project_order':'remaining_project_queue.json:stable_order','resume_scripts':'先collect_projects.py按项目补元数据、人工更新project_assessments/sample_annotations，再finalize_outputs.py；不重复全量检索','raw_reads_downloaded':0,'next_actions':['优先解读7项目方法/处理键，形成下载矩阵；仍不在本轮下载FASTQ','解决PRJNA1086167别名/SRR处理键冲突；1447957/1128434/1426923/EB65123分组键','追踪LakeWashington跨子项目同研究/合成群落去重','GSA入口受阻及PRJCA039809官方别名在ENA/DDBJ已识别','其余项目与关联论文从稳定队列续跑']})
 print(json.dumps(stat,ensure_ascii=False))
if __name__=='__main__':finalize()
