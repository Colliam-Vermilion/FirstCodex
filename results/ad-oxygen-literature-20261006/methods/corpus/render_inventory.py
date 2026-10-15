"""Render the unified current inventory; old review workbooks remain snapshots."""
from pathlib import Path
import json,argparse
from openpyxl import Workbook
from audit_library import read_csv


def render(project):
    p=Path(project);wb=Workbook();wb.remove(wb.active)
    current_papers=read_csv(p/'inventory/papers_utf8.csv');current_assets=read_csv(p/'inventory/assets_utf8.csv');current_body=sum(r['main_status']=='downloaded_verified' for r in current_papers);explicit_none=sum(r['supplementary_status']=='explicit_none' for r in current_papers)
    for name,relative in [('文献554','inventory/papers_utf8.csv'),('文件及来源证据','inventory/assets_utf8.csv'),('解析文件','inventory/extracted_files_utf8.csv'),('当前待正文','windows/main_records_current_utf8.csv'),('SI核查426','windows/supplementary_review_current_utf8.csv')]:
        rows=read_csv(p/relative);ws=wb.create_sheet(name);ws.append(list(rows[0]))
        for r in rows:ws.append(list(r.values()))
        ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
    notes=wb.create_sheet('状态及限制',0)
    for row in [('项','说明'),('当前与历史',f'当前正文ID {current_body}；历史91成功清单不覆盖。2026-10-15审计基线同SHA80、另版2、无正文9；原字节缺11。'),('文件分组','107主选核心文件＋1个R11596另存正文PDF＋17 JATS XML＋10传输包＝135对象；89正文ID不变。'),('原件映射','现有original_files真实原件保留；papers是统一逻辑映射，不另建重复原件目录。'),('SI状态',f'未检查/访问失败/仓储未声明附件不等于明确无附件；当前explicit_none={explicit_none}。'),('表格','保留工作表、行列/单元格、公式和缓存值；PDF表格/OCR未经人工逐值校对。'),('旧表','final_results原110列/人工判断不改，本轮当前权威视图是inventory和papers。'),('持久性','已有54 Git原件已远端读回；其余文件以storage_status及实际读回回执为准，私有包准备不等于上传完成。')]:notes.append(row)
    target=p/'inventory/unified_library_inventory.xlsx';wb.save(target);print(target)


if __name__=='__main__':
    ap=argparse.ArgumentParser(description=__doc__);ap.add_argument('--project',type=Path,default=Path(__file__).resolve().parents[2]);render(ap.parse_args().project)
