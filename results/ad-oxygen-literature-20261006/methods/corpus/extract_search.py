"""Extract immutable-source text/tables with precise locators; build/query SQLite FTS5.

No original is edited. OCR and vector-grid PDF tables are derived, unreviewed text.
Archives are listed, never executed. Existing SHA-bound archive derivatives are parsed
as separate assets, retaining the official parent ZIP relationship.
"""
from pathlib import Path
from collections import Counter
from concurrent.futures import ProcessPoolExecutor,as_completed
import argparse,csv,hashlib,json,os,re,shutil,sqlite3,subprocess,tempfile,time,zipfile,io
from datetime import date,datetime
from audit_library import dump,sha,safe_name,write_csv
PARSER_VERSION='1.0.0'


def value(v):
    if isinstance(v,(date,datetime)):return v.isoformat()
    return v


def extract_asset(task):
    a,root,ocr=task;source=Path(a['resolved_source_path']);folder=Path(root)/'papers'/a['record_id']/'extracted'/a['asset_id'];folder.mkdir(parents=True,exist_ok=True)
    report={'record_id':a['record_id'],'asset_id':a['asset_id'],'source_sha256':a['sha256'],'source_filename':a['original_filename'],'canonical_source':a['canonical_path'],'parser_version':PARSER_VERSION,'output_directory':str(folder),'status':'pending','ocr_pages':0,'ocr_pending_pages':[],'table_count':0,'errors':[],'limitations':[]}
    previous=folder/'report.json'
    if previous.is_file():
        r=json.loads(previous.read_text())
        if r.get('source_sha256')==a['sha256'] and r.get('parser_version')==PARSER_VERSION and r.get('ocr_requested')==ocr and r.get('status') not in ['error','pending']:
            if all((folder/f['relative_path']).is_file() and sha(folder/f['relative_path'])==f['sha256'] for f in r.get('output_files',[])):return r
    report['ocr_requested']=ocr
    if not source.is_file() or sha(source)!=a['sha256']:
        report['status']='source_missing_or_changed';dump(previous,report);return report
    chunks=[];tables=[];pages=[];cellfiles=[]
    def chunk(text,location,**extra):
        if str(text).strip():chunks.append({'record_id':a['record_id'],'asset_id':a['asset_id'],'source_filename':a['original_filename'],'source_sha256':a['sha256'],'canonical_source':a['canonical_path'],'role':a['role'],'location':location,'text':str(text),**extra})
    def table(rows,name,location):
        dest=folder/'tables'/name;dest.parent.mkdir(exist_ok=True)
        with dest.open('w',encoding='utf-8-sig',newline='') as f:csv.writer(f,lineterminator='\n').writerows(rows)
        tables.append({'path':str(dest.relative_to(folder)),'location':location,'rows':len(rows),'columns':max(map(len,rows),default=0),'quality':'automatic_extraction_not_manually_reconciled'})
    try:
        fmt=a['format']
        if fmt=='pdf':
            import fitz
            doc=fitz.open(source)
            for n,page in enumerate(doc,1):
                native=page.get_text(sort=True);text=native;method='native';confidence=None;blocks=page.get_text('blocks',sort=True)
                needs_ocr=len(native.strip())<80 and bool(page.get_images())
                if needs_ocr and ocr and shutil.which('tesseract'):
                    try:
                        with tempfile.TemporaryDirectory(prefix='paper-ocr-') as tmp:
                            image=Path(tmp)/'page.png';page.get_pixmap(matrix=fitz.Matrix(2.5,2.5),alpha=False).save(str(image))
                            env=os.environ.copy();env['OMP_THREAD_LIMIT']='1'
                            proc=subprocess.run(['tesseract',str(image),'stdout','-l','eng','--psm','3','tsv'],capture_output=True,text=True,encoding='utf-8',timeout=100,env=env)
                            if proc.returncode:raise RuntimeError('tesseract failed')
                            words=list(csv.DictReader(io.StringIO(proc.stdout),delimiter='\t'));lines={};confs=[]
                            for w in words:
                                if not w.get('text','').strip():continue
                                key=(w['block_num'],w['par_num'],w['line_num']);lines.setdefault(key,[]).append(w['text']);confs.append(float(w['conf']))
                            ocrtext='\n'.join(' '.join(v) for v in lines.values())
                            if len(ocrtext.strip())>len(native.strip()):text=ocrtext;method='ocr_eng';confidence=sum(confs)/len(confs) if confs else None;report['ocr_pages']+=1
                            else:method='ocr_no_additional_text'
                    except Exception as exc:method='ocr_failed';report['errors'].append({'pdf_page':n,'error':type(exc).__name__});report['ocr_pending_pages'].append(n)
                elif needs_ocr:method='ocr_needed_not_run';report['ocr_pending_pages'].append(n)
                pages.append({'pdf_page':n,'pdf_page_label':page.get_label(),'native_characters':len(native),'selected_characters':len(text),'text_method':method,'ocr_mean_word_confidence':confidence,'rotation':page.rotation,'width_points':page.rect.width,'height_points':page.rect.height,'text':text})
                if method=='native':
                    for b,block in enumerate(blocks,1):
                        if len(block)>6 and block[6]!=0:continue
                        chunk(block[4],{'type':'pdf_block','pdf_page':n,'pdf_page_label':page.get_label(),'block':b,'bbox_points':list(block[:4])},text_method='native')
                else:chunk(text,{'type':'pdf_page','pdf_page':n,'pdf_page_label':page.get_label()},text_method=method,ocr_unreviewed=True)
                # Conservative vector-grid extraction. Borderless/image tables remain text/OCR,
                # never falsely reported as structurally recovered tables.
                if len(native)>80:
                    drawings=page.get_drawings()
                    straight=sum(item[0] in ('l','re') for d in drawings for item in d.get('items',[]))
                    if 8<=straight<=2000:
                        try:
                            found=page.find_tables(strategy='lines_strict',paths=drawings)
                            for tn,t in enumerate(found.tables,1):
                                rows=t.extract()
                                if not rows or max(map(len,rows),default=0)<2:continue
                                loc={'type':'pdf_table','pdf_page':n,'table':tn,'bbox_points':list(t.bbox)}
                                table(rows,f'page_{n:04d}_table_{tn:02d}.csv',loc)
                                for rn,row in enumerate(rows,1):chunk(' | '.join('' if x is None else str(x) for x in row),dict(loc,row=rn),table_columns=list(range(1,len(row)+1)))
                        except Exception as exc:report['errors'].append({'pdf_page':n,'table_error':type(exc).__name__})
            report.update(pdf_pages=len(doc),native_pages=sum(p['text_method']=='native' for p in pages));report['limitations']+=['PDF logical reading order and inferred vector-grid cells may be imperfect; borderless/scan tables are not guaranteed structured.','OCR language=eng; OCR values, symbols and units require checking against source page.']
            with (folder/'pages.jsonl').open('w',encoding='utf-8') as f:
                for r in pages:f.write(json.dumps(r,ensure_ascii=False)+'\n')
        elif fmt=='docx':
            from docx import Document
            from docx.text.paragraph import Paragraph
            from docx.table import Table
            doc=Document(source);pn=0;tn=0;section=None
            for part in doc.iter_inner_content():
                if isinstance(part,Paragraph):
                    pn+=1
                    if part.style and part.style.name.startswith('Heading'):section=part.text
                    chunk(part.text,{'type':'word_paragraph','paragraph':pn,'section_heading':section,'style':part.style.name if part.style else None})
                elif isinstance(part,Table):
                    tn+=1;rows=[]
                    for rn,row in enumerate(part.rows,1):
                        values=[]
                        for cn,c in enumerate(row.cells,1):
                            values.append(c.text);chunk(c.text,{'type':'word_table_cell','table':tn,'row':rn,'column':cn,'section_heading':section})
                        rows.append(values)
                    table(rows,f'table_{tn:03d}.csv',{'type':'word_table','table':tn})
            for index,section in enumerate(doc.sections,1):
                for name in ['header','footer']:
                    for n,par in enumerate(getattr(section,name).paragraphs,1):chunk(par.text,{'type':'word_'+name,'document_section':index,'paragraph':n})
            report['limitations'].append('Word OOXML body paragraphs/tables and section headers/footers; page numbers depend on renderer; nested tables, text boxes, footnotes and embedded objects are not guaranteed parsed.')
        elif fmt in ['xlsx','xlsm']:
            from openpyxl import load_workbook
            cached=load_workbook(source,data_only=True,read_only=False,keep_links=False);formulas=load_workbook(source,data_only=False,read_only=False,keep_links=False);sheets=[];cell_count=0;formula_missing=0
            with (folder/'cells.jsonl').open('w',encoding='utf-8') as cells:
                for sn,ws in enumerate(formulas,1):
                    cv=cached[ws.title];rows=[];merged=[str(r) for r in ws.merged_cells.ranges];first_rows=[]
                    for rn,row in enumerate(ws.iter_rows(),1):
                        vals=[];rendered=[];addresses=[]
                        for c in row:
                            v=value(c.value);formula=v if c.data_type=='f' else None;cached_v=value(cv[c.coordinate].value)
                            if v is not None:
                                record={'sheet':ws.title,'cell':c.coordinate,'row':rn,'column':c.column,'value':None if formula else v,'formula':formula,'cached_value':cached_v if formula else None,'number_format':c.number_format,'data_type':c.data_type,'hidden_row':bool(ws.row_dimensions[rn].hidden)}
                                cells.write(json.dumps(record,ensure_ascii=False)+'\n');cell_count+=1
                                if formula and cached_v is None:formula_missing+=1
                                show=cached_v if formula and cached_v is not None else v;rendered.append(c.coordinate+'='+str(show));addresses.append(c.coordinate)
                            vals.append(cached_v if formula and cached_v is not None else v)
                        rows.append(vals)
                        if rn<=10:first_rows.append({'row':rn,'values':vals})
                        chunk(' | '.join(rendered),{'type':'excel_row','sheet':ws.title,'row':rn,'cells':addresses})
                    table(rows,f'sheet_{sn:03d}_{safe_name(ws.title)}.csv',{'type':'excel_sheet','sheet':ws.title,'start_cell':'A1','end_cell':ws.cell(ws.max_row,ws.max_column).coordinate})
                    sheets.append({'sheet':ws.title,'rows':ws.max_row,'columns':ws.max_column,'merged_ranges':merged,'sheet_state':ws.sheet_state,'first_10_rows_header_candidates_not_assumed_headers':first_rows})
            dump(folder/'workbook.json',{'sheets':sheets,'cell_values':cell_count,'formulas_without_cached_values':formula_missing,'formulas_recalculated':False});cached.close();formulas.close();report['limitations'].append('Formulas are preserved, never executed/recalculated; unavailable cached values are explicit. Charts and embedded images are not content-indexed.')
        elif fmt=='xml' and a['role']=='evidence_xml':
            import xml.etree.ElementTree as ET
            xml=ET.parse(source).getroot();tn=[0];graphics=[]
            def content(e):return ' '.join(''.join(e.itertext()).split())
            def walk(e,path,section=None):
                tag=e.tag.rsplit('}',1)[-1]
                if tag=='sec':section=content(e.find('title')) if e.find('title') is not None else section
                if tag in ['p','article-title','ref']:
                    chunk(content(e),{'type':'jats_reference' if tag=='ref' else 'jats_paragraph','xpath':path,'xml_id':e.attrib.get('id'),'section_heading':section});return
                if tag=='table':
                    tn[0]+=1;rows=[]
                    for rn,row in enumerate(e.findall('.//tr'),1):
                        values=[]
                        for cn,cell in enumerate(list(row),1):
                            text=content(cell);values.append(text);chunk(text,{'type':'jats_table_cell','xpath':path,'table':tn[0],'row':rn,'column':cn,'colspan':cell.attrib.get('colspan','1'),'rowspan':cell.attrib.get('rowspan','1'),'section_heading':section})
                        rows.append(values)
                    table(rows,f'jats_table_{tn[0]:03d}.csv',{'type':'jats_table','xpath':path,'table':tn[0],'spans_not_expanded':True});return
                if tag=='graphic':graphics.append({'xpath':path,'href':e.attrib.get('{http://www.w3.org/1999/xlink}href'),'status':'referenced_media_not_content_indexed; may be retained in transport ZIP'})
                counters={}
                for child in e:
                    name=child.tag.rsplit('}',1)[-1];counters[name]=counters.get(name,0)+1;walk(child,path+'/'+name+'['+str(counters[name])+']',section)
            walk(xml,'/article');dump(folder/'referenced_media.json',graphics);report['limitations'].append('Repository JATS XML text/structured tables; linked figures/media not all fetched. It is not a recovered historical PDF. Table rowspan/colspan retained as locators, not expanded into a rectangular semantic grid.')
        elif fmt in ['csv','tsv','txt','md','json','xml','r','py','fasta','fa']:
            data=source.read_bytes();encoding='utf-8-sig'
            try:text=data.decode(encoding)
            except UnicodeDecodeError:encoding='gb18030';text=data.decode(encoding)
            report['text_encoding']=encoding
            if fmt in ['csv','tsv']:
                delimiter='\t' if fmt=='tsv' else csv.Sniffer().sniff(text[:20000],delimiters=',;\t').delimiter
                rows=list(csv.reader(io.StringIO(text),delimiter=delimiter));table(rows,'table.csv',{'type':'delimited_table','delimiter':delimiter})
                for rn,row in enumerate(rows,1):chunk(' | '.join(row),{'type':'delimited_row','row':rn,'columns':list(range(1,len(row)+1))},field_names=rows[0] if rows else [])
            else:
                for n,line in enumerate(text.splitlines(),1):chunk(line,{'type':'text_line','line':n})
        elif fmt=='zip':
            members=[]
            with zipfile.ZipFile(source) as z:
                for item in z.infolist():
                    if item.is_dir():continue
                    member={'path':item.filename,'size_bytes':item.file_size,'compressed_bytes':item.compress_size,'unsafe_path':Path(item.filename).is_absolute() or '..' in Path(item.filename).parts or bool(re.match(r'^[A-Za-z]:',item.filename))}
                    if item.file_size<=100*1024*1024 and sum(i.file_size for i in z.infolist())<=500*1024*1024:
                        member['sha256']=hashlib.sha256(z.read(item)).hexdigest();member['crc_checked']=True
                    members.append(member)
            dump(folder/'archive_members.json',members)
            report['status']='archive_catalogued';report['limitations'].append('ZIP member names/checksums indexed; member content indexed through separately SHA-bound derived assets. No executable/code is run.')
        else:
            report['status']='retained_not_content_indexed';report['limitations'].append('Original retained; no supported content parser for this format. Images/video/models are not claimed searchable.')
            dump(folder/'unparsed_asset.json',{'format':fmt,'reason':'format-specific parser unavailable','original_retained':True})
        with (folder/'chunks.jsonl').open('w',encoding='utf-8') as f:
            for index,c in enumerate(chunks,1):c['chunk_id']=a['asset_id']+':'+str(index);f.write(json.dumps(c,ensure_ascii=False)+'\n')
        with (folder/'text.txt').open('w',encoding='utf-8') as f:
            for c in chunks:f.write('['+json.dumps(c['location'],ensure_ascii=False)+']\n'+c['text']+'\n\n')
        dump(folder/'tables.json',tables);report['table_count']=len(tables);report['search_chunks']=len(chunks)
        if report['status']=='pending':report['status']='parsed_with_limits' if report['errors'] or report['ocr_pending_pages'] else 'parsed'
    except Exception as exc:report['status']='error';report['errors'].append({'type':type(exc).__name__,'message':str(exc)[:300]})
    report['output_files']=[{'relative_path':str(f.relative_to(folder)).replace('\\','/'),'size_bytes':f.stat().st_size,'sha256':sha(f)} for f in sorted(folder.rglob('*')) if f.is_file() and f.name!='report.json']
    dump(previous,report);return report


def build_index(root):
    root=Path(root);target=root/'search.sqlite';tmp=root/'search.sqlite.tmp'
    if tmp.exists():tmp.unlink()
    db=sqlite3.connect(tmp);db.execute('CREATE VIRTUAL TABLE chunks USING fts5(record_id UNINDEXED,asset_id UNINDEXED,source_filename UNINDEXED,source_sha256 UNINDEXED,canonical_source UNINDEXED,location UNINDEXED,text, tokenize="unicode61")')
    count=0
    for path in sorted((root/'papers').glob('*/extracted/*/chunks.jsonl')):
        with path.open(encoding='utf-8') as f:
            for line in f:
                c=json.loads(line);db.execute('INSERT INTO chunks VALUES(?,?,?,?,?,?,?)',(c['record_id'],c['asset_id'],c['source_filename'],c['source_sha256'],c['canonical_source'],json.dumps(c['location'],ensure_ascii=False),c['text']));count+=1
    db.commit();db.close();tmp.replace(target);return count


def extract(project,root,ocr,workers,assets_path=None):
    assets=json.loads((assets_path or project/'inventory/assets.json').read_text());
    for a in assets:
        candidates=[root/a['canonical_path']]
        if a.get('source_repo_path'):candidates.append(project/a['source_repo_path'])
        candidates += [Path(a[k]) for k in ['resolved_source_path','source_local_path'] if a.get(k)]
        a['resolved_source_path']=str(next((f for f in candidates if f.is_file() and sha(f)==a['sha256']),candidates[0]))
    root.mkdir(parents=True,exist_ok=True);reports=[]
    with ProcessPoolExecutor(max_workers=max(1,min(workers,3))) as pool:
        jobs={pool.submit(extract_asset,(a,str(root),ocr)):a for a in assets if a['availability']=='present_sha_verified'}
        for future in as_completed(jobs):
            report=future.result();reports.append(report);print(json.dumps({'completed':len(reports),'asset_id':report['asset_id'],'status':report['status'],'ocr_pages':report['ocr_pages'],'tables':report['table_count']}),flush=True)
    chunks=build_index(root);files=[f for f in (root/'papers').glob('*/extracted/**/*') if f.is_file()]+[root/'search.sqlite']
    summary={'parser_version':PARSER_VERSION,'original_assets_seen':len(assets),'reports':len(reports),'status_counts':dict(Counter(r['status'] for r in reports)),'search_chunks':chunks,'structured_tables':sum(r['table_count'] for r in reports),'ocr_pages':sum(r['ocr_pages'] for r in reports),'ocr_pending_pages':sum(len(r['ocr_pending_pages']) for r in reports),'pdf_pages':sum(r.get('pdf_pages',0) for r in reports),'extraction_files_excluding_summary':len(files),'extraction_bytes_excluding_summary':sum(f.stat().st_size for f in files),'originals_modified':False,'full_outputs_root_is_temporary_not_backup':str(root),'all_reports':reports}
    summary['scope']='intake_assets' if assets_path else 'full_inventory'
    dump(root/('extraction_summary_intake.json' if assets_path else 'extraction_summary.json'),summary);dump(project/'inventory'/('intake_extraction_summary.json' if assets_path else 'extraction_summary.json'),summary)
    print(json.dumps({k:v for k,v in summary.items() if k!='all_reports'},ensure_ascii=False),flush=True)


def main():
    ap=argparse.ArgumentParser(description=__doc__);sub=ap.add_subparsers(dest='action',required=True)
    e=sub.add_parser('extract');e.add_argument('--project',type=Path,default=Path(__file__).resolve().parents[2]);e.add_argument('--root',type=Path,required=True);e.add_argument('--ocr',action='store_true');e.add_argument('--workers',type=int,default=2);e.add_argument('--assets',type=Path,help='Use intake/assets.json for newly reviewed Windows files')
    b=sub.add_parser('index');b.add_argument('--root',type=Path,required=True)
    s=sub.add_parser('search');s.add_argument('--root',type=Path,required=True);s.add_argument('--query',required=True);s.add_argument('--limit',type=int,default=20);s.add_argument('--record-id');s.add_argument('--asset-id');s.add_argument('--contains',action='store_true',help='Literal substring search, including Chinese; no FTS syntax')
    args=ap.parse_args()
    if args.action=='extract':extract(args.project.resolve(),args.root.resolve(),args.ocr,args.workers,args.assets)
    elif args.action=='index':print(json.dumps({'chunks':build_index(args.root)}))
    else:
        db=sqlite3.connect('file:'+str(args.root/'search.sqlite')+'?mode=ro',uri=True);db.row_factory=sqlite3.Row
        where="text LIKE ? ESCAPE '\\'" if args.contains else 'chunks MATCH ?';params=['%'+args.query.replace('\\','\\\\').replace('%','\\%').replace('_','\\_')+'%' if args.contains else args.query]
        if args.record_id:where+=' AND record_id=?';params.append(args.record_id)
        if args.asset_id:where+=' AND asset_id=?';params.append(args.asset_id)
        params.append(args.limit)
        selection="text AS excerpt" if args.contains else 'snippet(chunks,6,"[", "]"," … ",28) AS excerpt'
        for row in db.execute('SELECT record_id,asset_id,source_filename,source_sha256,canonical_source,location,'+selection+' FROM chunks WHERE '+where+(' ORDER BY rank' if not args.contains else '')+' LIMIT ?',params):
            r=dict(row);r['location']=json.loads(r['location'])
            if args.contains:
                pos=r['excerpt'].lower().find(args.query.lower());r['excerpt']=r['excerpt'][max(0,pos-90):max(0,pos)+180]
            print(json.dumps(r,ensure_ascii=False))
        db.close()


if __name__=='__main__':main()
