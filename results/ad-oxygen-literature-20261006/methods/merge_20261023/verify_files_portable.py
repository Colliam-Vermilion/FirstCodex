import argparse,json,hashlib,pathlib,zipfile
from openpyxl import load_workbook
p=argparse.ArgumentParser();p.add_argument("result_root",type=pathlib.Path);args=p.parse_args();root=args.result_root
assets=json.load(open(root/"inventory/merge_20261023/unified_files.json",encoding="utf-8"))
for a in assets:
 f=root/a["canonical_path"];h=hashlib.sha256();size=0
 with f.open("rb") as r:
  for data in iter(lambda:r.read(1024*1024),b""):h.update(data);size+=len(data)
 assert h.hexdigest()==a["sha256"] and size==a["size_bytes"],str(f)
print(len(assets),"real original objects: SHA256 and size match")
for name in ["final_results/journal_jcr_statistics_20261023.xlsx","final_results/ad_oxygen_download_review_20261006.xlsx","inventory/unified_library_inventory.xlsx"]:
 f=root/name
 with zipfile.ZipFile(f) as z:assert z.testzip() is None
 w=load_workbook(f,read_only=True,data_only=True);print(name,"readable",len(w.sheetnames),"worksheets");w.close()
