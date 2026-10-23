import pathlib,json,csv,shutil,collections,hashlib,zipfile
from openpyxl import load_workbook,Workbook
from openpyxl.styles import Font,PatternFill,Alignment
W=pathlib.Path('/workspace/work');SRC=W/'snapshot/results/ad-oxygen-literature-20261006';B=W/'review/results/ad-oxygen-literature-20261006';I=B/'inventory/merge_20261023'
for name in ['inventory','final_results','manual_review']:shutil.copytree(SRC/name,B/name,dirs_exist_ok=True)
def readcsv(p):return list(csv.DictReader(open(p,encoding='utf-8-sig')))
def writecsv(p,rows):
 fields=list(dict.fromkeys(k for r in rows for k in r))
 with open(p,'w',encoding='utf-8-sig',newline='') as f:
  w=csv.DictWriter(f,fieldnames=fields,extrasaction='ignore',lineterminator='\n');w.writeheader()
  for r in rows:w.writerow({k:json.dumps(v,ensure_ascii=False) if isinstance(v,(dict,list)) else v for k,v in r.items()})
A=json.load(open(I/'unified_files.json'));S=json.load(open(I/'paper_status.json'));SM={x['record_id']:x for x in S};R=readcsv(B/'final_results/ad_oxygen_download_review_utf8.csv');RM={r['稳定记录ID']:r for r in R};M=readcsv(I/'source_path_mapping_utf8.csv');meta=readcsv(I/'journal_metadata_verification_utf8.csv');summary=json.load(open(I/'merge_summary.json'));J=json.load(open(I/'jcr_summary.json'))
P=json.load(open(I/'main_preservation_before_cleanup.json'));bysha={a['sha256']:a for a in A};branch_sources=[]
for x in P['files']:
 a=bysha[x['sha256']];a.setdefault('main_misupload_source_paths',[]).append(x['main_path']);a.setdefault('source_branches',['results/ad-oxygen-literature-20261006']);a['source_branches'].append('main@ab8353c');a['source_branches']=list(dict.fromkeys(a['source_branches']))
 branch_sources.append({'record_id':a['record_id'],'source_record_id':x['main_path'].split('/')[0],'title':a['title'],'doi':a['doi'],'role':a['role'],'version':a['version'],'original_filename':pathlib.Path(x['main_path']).name,'source_path':'main@ab8353c:'+x['main_path'],'source':'main_misupload_copy','canonical_path':a['canonical_path'],'size_bytes':a['size_bytes'],'sha256':a['sha256'],'git_blob_oid':a['git_blob_oid'],'verification_status':'preserved_and_read_back_before_main_cleanup','association_status':a['association_status'],'verification_notes':'正常清理main之前，保全提交b739c0e的真实原件已独立读回SHA256/大小通过；规范对象复用','github_url':a['github_url']})
M=[m for m in M if m['source']!='main_misupload_copy']+branch_sources
writecsv(I/'source_path_mapping_utf8.csv',M);writecsv(I/'unified_files_utf8.csv',A);json.dump(A,open(I/'unified_files.json','w'),ensure_ascii=False,indent=2);writecsv(B/'inventory/assets_utf8.csv',[dict(a,role='main' if a['role']=='fulltext' else a['role'],source_repo_path=a['canonical_path'],availability='present_sha_verified') for a in A]);json.dump(A,open(B/'inventory/assets.json','w'),ensure_ascii=False,indent=2)
summary.update(cross_branch_main_source_copies=30,source_mapping_rows=len(M),original_input_occurrences=673);json.dump(summary,open(I/'merge_summary.json','w'),ensure_ascii=False,indent=2)
for m in meta:
 r=RM[m['record_id']];r.update({'本轮规范期刊':m['normalized_journal'],'本轮ISSN':m['issn'],'本轮eISSN':m['eissn'],'本轮核实载体类型':m['publication_kind'],'本轮统计文献组':m['subtype'],'本轮期刊身份核验备注':m['notes']})
writecsv(B/'final_results/ad_oxygen_download_review_utf8.csv',R)
# Reconcile existing per-ID metadata without discarding historical evidence.
for rid,r in RM.items():
 p=B/'papers'/rid/'metadata.json';p.parent.mkdir(parents=True,exist_ok=True);d=json.load(open(p)) if p.exists() else {'record_id':rid,'title':r['完整题名'],'doi':r['DOI'],'category_unchanged':r['最终分类']};s=SM[rid]
 if 'acquisition_before_20261023' not in d:d['acquisition_before_20261023']={k:d.get(k) for k in ['main_status','supplementary_status','supplementary_coverage_complete','supplementary_official_count','preferred_main_asset_id','current_verified_main_asset_ids','assets']}
 fs=[a for a in A if rid in a['associated_record_ids']];mains=[a for a in fs if a['role']=='fulltext'];sup=[a for a in fs if a['role']=='supplementary'];d['main_status']='downloaded_verified' if mains else 'not_downloaded';d['acquisition_20261023']=s;d['assets']=[dict(a,role='main' if a['role']=='fulltext' else a['role'],source_repo_path=a['canonical_path']) for a in fs];d['current_verified_main_asset_ids']=[a['asset_id'] for a in mains];d['preferred_main_asset_id']=mains[0]['asset_id'] if mains else '';d['supplementary_status']=s['supplementary_status'];d['supplementary_original_count_20261023']=len(sup);d['supplementary_official_count']=sum(any(p.get('official_attachment')=='True' for p in a['prior_asset_evidence']) for a in sup);d['source_mapping_standard']='Original files: original_files/<ID>/{fulltext,supplementary,abstract_poster,evidence,uncertain,derived}; source_path_mapping_utf8.csv retains all historical paths and branch copies.';d['checked_at_utc_20261023']='2026-10-23';d['journal_metadata_20261023']=next(x for x in meta if x['record_id']==rid)
 json.dump(d,open(p,'w'),ensure_ascii=False,indent=2)
# ZIP members remain in their original containers; verify the five old derived files without re-extracting.
lineage=[]
for a in A:
 if a['role']=='derived':
  prior=a['prior_asset_evidence'][0];parent=next(x for x in A if prior['archive_parent_asset_id'] in x['legacy_asset_ids']);member=prior['archive_member_path']
  with zipfile.ZipFile(B/parent['canonical_path']) as z:data=z.read(member)
  assert hashlib.sha256(data).hexdigest()==a['sha256']
  lineage.append({'derived_path':a['canonical_path'],'parent_zip_path':parent['canonical_path'],'member_path':member,'sha256':a['sha256'],'match':True,'not_counted_again_as_original_si':True})
json.dump(lineage,open(I/'zip_member_lineage_verified.json','w'),ensure_ascii=False,indent=2)
def style(ws):
 ws.freeze_panes='B2';ws.auto_filter.ref=ws.dimensions;ws.sheet_view.showGridLines=False
 for c in ws[1]:c.font=Font(name='Microsoft YaHei',bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='16324F');c.alignment=Alignment(wrap_text=True)
 for col in ws.columns:ws.column_dimensions[col[0].column_letter].width=min(60,max(15,len(str(col[0].value))*2))
# Existing master workbook: all historic sheets survive, current explanation becomes the first sheet.
master=B/'final_results/ad_oxygen_download_review_20261006.xlsx';mw=load_workbook(master);ws=mw['文献库'];headers=[c.value for c in ws[1]];idcol=headers.index('稳定记录ID')+1
for h in R[0]:
 if h not in headers:headers.append(h);ws.cell(1,len(headers),h)
for row in range(2,ws.max_row+1):
 r=RM[ws.cell(row,idcol).value]
 for c,h in enumerate(headers,1):
  if h in r:ws.cell(row,c,r[h])
if '统计与说明' in mw.sheetnames:mw['统计与说明'].title='历史统计与说明_20261017'
if '20261023当前说明' in mw.sheetnames:del mw['20261023当前说明']
ws=mw.create_sheet('20261023当前说明',0);ws.append(['项目','当前说明']);ws.append(['范围与原件','554条记录、426纳入；388目标全文、38未取目标全文。仅摘要3、仅海报1。']);ws.append(['统一原件','665个不同字节对象，673条本地/既有云端输入，另30条main误传副本来源，共703条映射。']);ws.append(['JCR统计','见同目录journal_jcr_statistics_20261023.xlsx；2024指标年/2025发布版；已核实272/386目标期刊论文，Q1至少一学科213篇、28刊；全部已核实学科Q1为201篇、25刊。']);ws.append(['历史工作表','历史工作表保留作证据；旧获取数量及旧路径不作当前状态。文献库和当前环境总览已更新。']);ws.append(['主索引','inventory/merge_20261023/unified_files_utf8.csv；source_path_mapping_utf8.csv；paper_status.json']);style(ws);mw.save(master)
# Refresh the pre-existing inventory workbook with the same current object and record sets.
iw=Workbook();iw.remove(iw.active)
for name,rows in [('论文状态554',S),('原件对象665',A),('来源路径映射703',M)]:
 sh=iw.create_sheet(name);fields=list(dict.fromkeys(k for r in rows for k in r if k not in ['prior_asset_evidence','archive_members']))
 sh.append(fields)
 for r in rows:sh.append([json.dumps(r.get(k,''),ensure_ascii=False) if isinstance(r.get(k,''),(dict,list)) else r.get(k,'') for k in fields])
 style(sh)
iw.save(B/'inventory/unified_library_inventory.xlsx')
# Add explicit all-scope/nonjournal totals to the statistics sheet; preserve formula caches in the JCR file by editing XML directly.
jpath=B/'final_results/journal_jcr_statistics_20261023.xlsx'
import xml.etree.ElementTree as ET
ns='http://schemas.openxmlformats.org/spreadsheetml/2006/main';q=lambda n:'{'+ns+'}'+n
with zipfile.ZipFile(jpath) as z:files={n:z.read(n) for n in z.namelist()}
root=ET.fromstring(files['xl/worksheets/sheet4.xml']);sheetdata=root.find(q('sheetData'));allrows=[['全范围与非期刊分列','纳入记录数','已取得核验全文数','未取得目标全文数']]
allrows.append(['全部纳入文献（含非期刊，不作JCR分母）',426,388,38])
for group in ['研究类','综述','其他期刊论文','学位/学生报告','会议/非期刊','类型待核实']:
 mids=[m['record_id'] for m in meta if m['subtype']==group and SM[m['record_id']]['included_for_analysis']];full=sum(SM[r]['verified_fulltext'] for r in mids);allrows.append([group,len(mids),full,len(mids)-full])
for ridx,vals in enumerate(allrows,31):
 row=ET.SubElement(sheetdata,q('row'),{'r':str(ridx)})
 for col,v in enumerate(vals,1):
  from openpyxl.utils import get_column_letter
  c=ET.SubElement(row,q('c'),{'r':f'{get_column_letter(col)}{ridx}'})
  if isinstance(v,str):c.set('t','inlineStr');ET.SubElement(ET.SubElement(c,q('is')),q('t')).text=v
  else:ET.SubElement(c,q('v')).text=str(v)
root.find(q('dimension')).set('ref','A1:N38');files['xl/worksheets/sheet4.xml']=ET.tostring(root,encoding='utf-8',xml_declaration=True)
with zipfile.ZipFile(jpath,'w',zipfile.ZIP_DEFLATED) as z:
 for n,data in files.items():z.writestr(n,data)
# Reproducibility scripts contain no credentials. They are audit tools, not new download queues.
scripts=B/'methods/merge_20261023';scripts.mkdir(exist_ok=True)
for name in ['audit_files.py','prepare_merge.py','finalize_merge.py','parse_jcr.py','build_journal_excel.py','finish_delivery.py']:shutil.copy2(W/name,scripts/name)
doc='''# 2026-10-23 本地上传整理、身份核验及JCR统计方法

## 基准与范围

远端结果基准022237df5e83e99711e7fa358ed2b4611d42bab1。读取当前README、18/20轮方法、554行总表、assets清单和20261017原件备份清单。按完整现行分类名称识别纳入范围：直接向AD反应器/内部循环供氧、厌氧水解/产酸阶段供氧、独立好氧/微好氧预处理、综述；共426条。128排除仅留审计。既有ID、分类、人工53条判断及证据字段保持，核验期刊/文献载体不重写范围。未进行新主题检索、缺失全文追下载或全文分析。

## 输入、物理原件和去重

localdownload在原结果分支结果目录下，439文件/295个ID；ID全部匹配，无大小写归一后的重复ID目录或未匹配文献ID。此前original_files除目录说明和归属表外234文件。仓库其余位置没有另一个PDF/Word原件库；papers为历史提取/元数据，不把解析文件作原件。673个文件输入按SHA256归为665个不同字节规范对象，8条额外同字节输入保留映射。另将main的30个同字节误传副本列作附加来源，共703条映射，不重复计原件。

沿用original_files，按ID建立fulltext/supplementary/abstract_poster/evidence/uncertain/administrative/derived子目录。每种真实字节只留一个规范对象；不同版本、14个分章学位文件及不同字节副本分别保留。402个目标正文文件组合为388篇：本地287篇、云端101篇；两个7部分学位组各计1篇。来源文件名、旧路径、原资产ID、来源URL和历史人工版本判断保留于source_path_mapping及prior_asset_evidence/legacy_asset_ids。源路径主仓库误传副本回到ab8353c，保全快照b739c0e。没有提交含重复原件的大ZIP。

## 内容与格式核验

PDF全部页面解析并检查页面对象，首页/前两页与题名、作者、目标DOI核对，再查正文结构、末部结论/参考文献及页序；前两页规则仅是初筛，题名/DOI相近或无DOI的16个本地正文均逐项复核。R00608五页中文扫描件逐页OCR；R02239 pp.3–7为目标正文，封面/目录不使其误判缺全文。R06611为48页接受稿，封面题名/作者/DOI匹配，图像正文跨题名、方法、结果、参考文献和图表页OCR抽查。R12483泰英题名/作者及20页正文复核。R13946为整本207页会议论文集，目标第5章在PDF pp.46–53/印刷页35–42完整存在，不计期刊论文，不把整本章节分别增记录。分章学位沿用既有人工连续性证据并核对实际7部分齐全、SHA、页对象可读。

识别4个本地身份冲突：R04605实际为Bioresource Technology 2019期刊文章10.1016/j.biortech.2018.10.077，非目标博士论文；R03965实际为Waste Management 2019期刊文章10.1016/j.wasman.2019.05.018，目标为2018会议论文，不因同作者和相近主题视为同一文献全文；R13947为Maryland Extension事实表草稿，非目标学位论文；R13949为NaOH碱预处理论文，非目标好氧预处理论文。全部放uncertain，保留实际观察题名及DOI，不增目标全文，不改变目标范围分类。

仅摘要R07247/R11845/R13961三篇、仅海报R05473一篇；另34条没有目标全文，其中包括4条上传身份冲突。合计38条目标缺全文。13种不同字节历史抓取候选（20旧路径）为目录、索引、广告或搜索页等，留uncertain；1个Authorship change request作者变更表为行政附件，非科研SI。未发现损坏或需要口令才能读取的对象；文字层乱码与图像扫描单独注明，不等同无内容。没有据文件名将poster/abstract/目录/错误页算全文。

DOCX/ZIP检查结构及CRC，Word读取正文XML；9个DOC用antiword实际解码；XLSX用openpyxl逐工作簿打开并记工作表维度；TXT解码、JPEG实际解码。图片/Word图形/模型不声称已完成科学内容或全部数值索引。3个ZIP原格式保留，只按成员读取核验，不重新落一套解包文件；5个旧派生Excel与父ZIP成员的SHA逐一核对，另记lineage，不计原件SI数量。

科研补充材料154个不同字节原件：DOCX109、PDF16、DOC9、XLSX9、JPG4、TXT4、ZIP3。标题/DOI、既有官方关系或内容关系分层记录；61件无唯一身份凭据，在待核实项单列，目录及内容相关不等于已完成身份核验。R13964和R13834两目录同字节Word为微污染物HPLC及GC色谱图，实验对象与R13834及其Appendix A相符，规范归R13834；R13964错误来源保留，不计其SI。既有官方附件证据/外部数据链接不删除，未找到附件不能推断没有附件或已收齐。

## main清理门槛

通过main提交ab8353c的路径及差异确认误传仅20个ID/30个文件，其他main项目文件（README）排除。30文件都已存在结果分支localdownload，main独有0。保全提交b739c0e后，按该新提交独立从GitHub重新读取30真实对象，大小和SHA256全部吻合，随后创建以ab8353c为父的新正常提交0bdf0b9清理main。没有强制推送、历史重写或其他main改动。旧副本仍在Git历史。规范化原件检查点e4f3298已另行推送。

## 期刊身份与JCR

仅按既有DOI调用Crossref核对题名、期刊、ISSN与载体类型，不检索新文献。509个DOI请求中500返回元数据、9未返回可用内容；非Crossref DOI不因此视为无效。中文/无DOI期刊另据原件页脚核对；原期刊字段PubMed/DOAJ/仓储平台不是期刊名，修正内容在新增规范字段，原书目字段与基准CSV保留。R02370原件页脚为2011年，年份差异单列，不自动重分类。规范期刊按统一名称及印刷/电子ISSN合并，同刊不同ISSN不重复计算期刊种数。

JCR统一为2024指标年/2025发布版。Clarivate Release Notes明确2025发布提供2024数据：https://journalcitationreports.zendesk.com/hc/en-gb/articles/36310651247889-2025 。JIF Quartile按每个JCR学科保留，https://journalcitationreports.zendesk.com/hc/en-gb/articles/28351430970001-Journal-Profile 。不采用中科院分区、SJR/CiteScore或IF高低推断。

目前可核验资料主要为武汉大学图书馆JCR栏/2024列表的搜索索引片段；保留原URL和查询快照、指标年及分区。直连页面超时，机构完整JCR PDF因网络403未取得；未取得Clarivate账户全量导出。因此这是可核验的部分覆盖，不是完整JCR结论。只解析明确标2024指标及JIF分区的机构信息；不把2023/2025指标或MDPI当前2025指标网页混入。多学科显示项完整保留；只核实列表显示学科的局限逐刊注明。全体已核实学科Q1不等于承诺官方全部学科没有遗漏。

期刊论文目标386条（研究339、综述46、其他技术说明1）；非期刊15、学位/学生18、载体类型待核实7另列。会议系列载体或索引平台元数据不充分者不直接放入已核实期刊分母；不影响原研究范围。JCR已核实272条（70.5%），114条待核实。至少一个已核实学科Q1为213篇/28种期刊，其中研究187、综述26；全部已核实学科均Q1为201篇/25种期刊。Q1占全部目标213/386=55.2%，占已核实分区213/272=78.3%。已取得核验全文的目标期刊论文369条，已核实JCR271条，Q1为213/369=57.7%、213/271=78.6%。详细全部口径、研究/综述及多学科差异在Excel统计汇总，不将缺失信息作非Q1。

请提供2025发布版/2024指标年的JCR导出Excel或CSV，含Journal name、ISSN/eISSN、全部Category、JIF Quartile、JIF Rank、JCR Year及导出日期，以补114条期刊分区及所有期刊完整学科覆盖；具体期刊与类型疑点见待核实项。Excel附公式、正确缓存、过滤和独立核对的静态JSON/CSV。原总Excel保留全部历史证据工作表，另加当前说明和更新总表状态；旧统计明确标历史。历史全文索引未重建，旧提取路径由SHA和映射回到新原件。

## 持久化与最终验收

分批推送真实原件。最大规范文件27,324,975字节，当前没有.gitattributes/LFS规则或LFS指针；全部真实Git blob，小于GitHub普通Git单文件100 MB门槛，不需要引入LFS。最终从独立新克隆读取当前提交，取回全部665原件，检查文件数量、大小、SHA256、原件非LFS指针和Excel ZIP CRC/可打开性；回执另存remote_readback.json。新克隆以GitHub真实对象为源，不使用原仓库alternate/shared对象目录。下载来自用户仓库自己的已上传对象，不是重新尝试缺失论文全文。
'''
(B/'methods/21_local_merge_and_jcr_20261023_zh.md').write_text(doc)
old=(B/'README.md').read_text();old=old[old.find('# 新环境恢复与全部原件备份'):]
(B/'README.md').write_text('''# 当前交付：本地文献合并与JCR统计（2026-10-23）

554条记录和人工分类证据保留；按当前分类名称纳入426条。localdownload的439文件/295个ID，与既有云端234文件合并，665个不同SHA256规范对象；673条输入来源，另含30条main副本来源，共703条路径映射。已取得目标全文388篇（本地287、云端101），仅摘要3、仅海报1，其余缺全文34。不继续追下载或分析缺全文文献。

- **[期刊与JCR统计Excel](final_results/journal_jcr_statistics_20261023.xlsx)**：文献明细、期刊汇总、统计汇总、待核实项、逐学科分区、统一文件清单及非期刊分列。
- [更新文献总Excel](final_results/ad_oxygen_download_review_20261006.xlsx)、[总CSV](final_results/ad_oxygen_download_review_utf8.csv)、[当前文件/论文库存Excel](inventory/unified_library_inventory.xlsx)
- **[真实规范原件目录](original_files/)**、[目录与文件角色](original_files/README.md)、[665原件清单](inventory/merge_20261023/unified_files_utf8.csv)、[703来源路径映射](inventory/merge_20261023/source_path_mapping_utf8.csv)、[逐文件核验](inventory/merge_20261023/file_content_verification.json)
- [38条缺目标全文](manual_review/current_environment_need_fulltext_utf8.csv)、[JCR与类型/附件待核实](inventory/merge_20261023/jcr_pending_utf8.csv)、[完整方法与身份异常说明](methods/21_local_merge_and_jcr_20261023_zh.md)、[ZIP派生核验](inventory/merge_20261023/zip_member_lineage_verified.json)
- [最终独立新克隆/原件和Excel读回回执](inventory/merge_20261023/remote_readback.json)、[合并统计](inventory/merge_20261023/merge_summary.json)、[JCR机器可读统计](inventory/merge_20261023/jcr_summary.json)

JCR统一为**2024指标年、2025发布版**。386目标期刊论文中核实272条（70.5%）：至少一个已核实学科Q1为213篇、28刊；全部已核实学科均Q1为201篇、25刊。已取得核验全文的期刊论文369篇中Q1为213篇。分区缺失114条，未知不作非Q1；请提供同年份JCR全部学科导出补足。本结果为部分核实下限，不是完整JCR结论；来源和多学科差异逐项保留。

科研补充材料154个不同字节原格式文件：DOCX109、PDF16、DOC9、XLSX9、JPG4、TXT4、ZIP3；5个旧ZIP派生成员不重复计。61件附件唯一关联仍待核实。4个上传正文身份冲突、13种历史非目标候选和1份作者变更表均保留，不计目标全文。R13964误放同字节SI已归R13834，两条来源保留。

main误传30文件先在保全提交`b739c0e`独立读回大小/SHA全部通过，再以正常提交`0bdf0b9`清理；README原字节不变，不重写历史，旧副本仍留Git历史。原件规范化检查点`e4f3298`已推送。当前以本节、方法21及本轮清单为准；以下全部为历史快照，旧数量和路径不代表当前状态。

---

'''+old)
json.dump({'main_cleanup_commit':summary['main_cleanup_commit'],'preservation_commit':summary['preservation_commit'],'originals_checkpoint_commit':'e4f3298905eb20ac44224cd386c10aece6243be2','main_files_removed':30,'main_other_files_unchanged':True,'history_rewritten':False,'main_only_originals':0},open(I/'github_commits.json','w'),ensure_ascii=False,indent=2)
print('Prepared final metadata, per-ID records, workbooks and methods',len(A),'original objects',len(M),'mappings')
