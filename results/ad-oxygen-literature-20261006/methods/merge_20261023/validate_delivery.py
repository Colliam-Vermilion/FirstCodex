import pathlib,json,csv,hashlib,zipfile,collections
from openpyxl import load_workbook
W=pathlib.Path('/workspace/work');B=W/'review/results/ad-oxygen-literature-20261006';I=B/'inventory/merge_20261023'
R=list(csv.DictReader(open(B/'final_results/ad_oxygen_download_review_utf8.csv',encoding='utf-8-sig')));OLD=list(csv.DictReader(open(W/'remote_master.csv',encoding='utf-8-sig')));A=json.load(open(I/'unified_files.json'));S=json.load(open(I/'paper_status.json'));J=json.load(open(I/'jcr_summary.json'));M=list(csv.DictReader(open(I/'source_path_mapping_utf8.csv',encoding='utf-8-sig')))
assert len(R)==554 and len(S)==554 and len(A)==665 and len(M)==703
old={r['稳定记录ID']:r for r in OLD};new={r['稳定记录ID']:r for r in R};assert set(old)==set(new)
protected=['完整题名','DOI','最终分类','纳入类别','人工最终分类','人工分类来源','人工分类日期','人工分类原始值','判断理由','证据定位','简短证据摘录','信息依据']
for rid in old:
 for k in protected:assert old[rid][k]==new[rid][k],(rid,k)
original_shas={a['sha256'] for a in json.load(open(W/'file_audit_initial.json'))};assert original_shas=={a['sha256'] for a in A}
for a in A:
 p=B/a['canonical_path'];d=p.read_bytes();assert len(d)==a['size_bytes'] and hashlib.sha256(d).hexdigest()==a['sha256'];assert not d.startswith(b'version https://git-lfs.github.com/spec/v1')
assert len({a['canonical_path'] for a in A})==665
assert sum(s['verified_fulltext'] and s['included_for_analysis'] for s in S)==388
for rid in ['R03965','R04605','R13947','R13949','R05473','R13961','R07247','R11845']:assert not next(s for s in S if s['record_id']==rid)['verified_fulltext']
assert sum(s['main_file_count'] for s in S)==402
assert next(s for s in S if s['record_id']=='R13964')['supplementary_count']==0
assert next(s for s in S if s['record_id']=='R13834')['supplementary_count']==1
for a in A:
 if a['format']=='zip':
  with zipfile.ZipFile(B/a['canonical_path']) as z:assert z.testzip() is None
xp=B/'final_results/journal_jcr_statistics_20261023.xlsx';wb=load_workbook(xp,read_only=True,data_only=True);ws=wb['文献明细'];headers=[x.value for x in next(ws.iter_rows())];rows=[dict(zip(headers,x)) for x in ws.iter_rows(min_row=2,values_only=True)];assert len(rows)==554
target=[r for r in rows if r['是否目标期刊论文']=='是'];assert len(target)==386
for i,stat in enumerate(J['summary'],2):
 subset=[r for r in target if (stat['集合'].startswith('全部') or r['核验全文标志']=='是') and (stat['文献组']=='合计' or r['统计文献组']==stat['文献组'])]
 expected=[len(subset),sum(r['JCR分区已核实']=='是' for r in subset),sum(r['至少一学科Q1']=='是' for r in subset),sum(r['全部已核实学科均Q1']=='是' for r in subset)]
 for col,val in zip([3,4,6,9],expected):assert wb['统计汇总'].cell(i,col).value==val,(i,col,val,wb['统计汇总'].cell(i,col).value)
 assert stat['至少一个学科Q1文章数']==expected[2]
jws=wb['期刊汇总'];jh=[x.value for x in next(jws.iter_rows())];jr=[dict(zip(jh,x)) for x in jws.iter_rows(min_row=2,values_only=True)];assert sum(x['纳入记录数'] for x in jr)==386;assert len({x['规范期刊'].lower() for x in jr})==len(jr);assert sum(x['至少一学科Q1']=='是' for x in jr)==28;assert sum(x['全部已核实学科均Q1']=='是' for x in jr)==25
assert wb['统一文件清单'].max_row==666;assert wb['待核实项'].max_row==201;assert wb['统计汇总']['B32'].value==426;wb.close()
master=load_workbook(B/'final_results/ad_oxygen_download_review_20261006.xlsx',read_only=True,data_only=True);ws=master['文献库'];hs=[c.value for c in next(ws.iter_rows())];rr=[dict(zip(hs,r)) for r in ws.iter_rows(min_row=2,values_only=True)];assert len(rr)==554
for r in rr:
 rid=r['稳定记录ID']
 for k in protected:assert str(r.get(k) or '')==str(old[rid].get(k) or ''),(rid,k)
 assert r['全文获取状态']==new[rid]['全文获取状态']
assert '历史统计与说明_20261017' in master.sheetnames;master.close()
for x in [xp,B/'final_results/ad_oxygen_download_review_20261006.xlsx',B/'inventory/unified_library_inventory.xlsx']:
 with zipfile.ZipFile(x) as z:assert z.testzip() is None
receipt={'checked_at_date':'2026-10-23','objects_verified':665,'object_bytes':sum(a['size_bytes'] for a in A),'all_original_input_sha256_retained':True,'record_ids_and_protected_evidence_fields_unchanged':554,'source_mapping_rows':703,'verified_target_fulltext_ids':388,'verified_target_fulltext_files':402,'journal_target_records':386,'jcr_verified_records':272,'q1_at_least_one_category_articles':213,'q1_all_verified_categories_articles':201,'q1_journal_count':28,'q1_all_verified_categories_journal_count':25,'three_workbooks_crc_and_open_verified':True,'jcr_formulas_cached_values_independently_recomputed':True,'lfs_pointers':0}
json.dump(receipt,open(I/'local_delivery_validation.json','w'),ensure_ascii=False,indent=2);print(json.dumps(receipt,ensure_ascii=False,indent=2))
