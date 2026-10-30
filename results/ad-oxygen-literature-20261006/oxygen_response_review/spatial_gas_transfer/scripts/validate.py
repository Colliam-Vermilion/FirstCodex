"""Read-only topical delivery checks, no paper downloads or result writes."""
import ast,csv,json,re,subprocess
from pathlib import Path
from openpyxl import load_workbook
from prepare import T,M,P,R,read
def main():
 rows=read(T/'literature_master.json');ids=[r['id'] for r in rows];assert len(ids)==len(set(ids))
 progress=read(T/'progress.json');done=[r for r in rows if r['status']!='未开始'];assert len(done)==progress['processed']
 assert sum(progress.get(k,0) for k in ['related_fulltext_or_existing_fulltext_evidence_reused','abstract_only_processed','fragment_only_processed','title_only_processed'])==len(done)
 if (T/'followup_screening.json').exists():
  follow=read(T/'followup_screening.json');base=read(T/'local_followup_baseline.json');assert len({r['id'] for r in follow})==len(follow)
  assert {r['id'] for r in follow}<={r['id'] for r in base}
  candidates={r['id']:r for r in read(T/'candidate_records.json')}
  for r in follow:
   assert not r['old_scope_judgment_changed'] and not r['new_external_fulltext_downloaded']
   if r.get('evidence_quote'):assert r['evidence_quote'] in candidates[r['id']]['abstract']
  prior=json.loads(subprocess.check_output(['git','show','4692a8906cbc9ec028f21bc39b57104710baccf0:'+str((T/'literature_master.json').relative_to(R))],cwd=R))
  current={r['id']:r for r in rows}
  for r in prior:
   if r['status']!='未开始':
    for k in ['reading_depth','decision_reason','actual_finding','author_interpretation_and_alternatives']:assert r[k]==current[r['id']][k],(r['id'],k)
 assert len(rows)==progress['candidates'] and progress['processed']+progress['remaining']==len(rows)
 old={r['id']:r for r in read(M/'catalog.json')}
 for r in rows:
  if r['id'] in old:assert r['existing_scope_decision_preserved']==old[r['id']].get('screening_decision')
  for s in r['source_links']:assert (R/s['path']).exists()
 for q in read(T/'search_log.json'):
  assert q['execution_state']=='completed_local'
  assert len(read(T/q['raw_path']))==q['returned_records']
 for q in read(T/'online_query_queue.json'):assert q['state']=='not_executed_network_policy_block'
 assert read(T/'citation_trace_log.json')['backward_returned']==42
 assert len(read(T/'backward_citation_trace.json'))==42
 assert sum(bool(r['assigned_existing_id']) for r in read(T/'seed_audit.json'))==1
 for name,n in [('spatial_gas_transfer_candidates.ris',len(rows)),('spatial_gas_transfer_reviewed.ris',len(done))]:
  text=(T/name).read_text();assert len(re.findall('^TY  - ',text,re.M))==len(re.findall('^ER  -',text,re.M))==n
 wb=load_workbook(T/'spatial_gas_transfer_cumulative.xlsx',read_only=True)
 assert len(wb.sheetnames)==11 and sum(1 for _ in wb['文献与专题筛选'].iter_rows(values_only=True))==len(rows)+1
 wb.close()
 for file in T.rglob('*.py'):ast.parse(file.read_text())
 # Original AD/omics files and all old review decisions remain untouched.
 paths=subprocess.check_output(['git','diff','--name-only','fcd0ad3119900af9f75c9b9220575b8ff79a0170'],cwd=R,text=True).splitlines()
 protected=['final_results/','manual_review/','original_files/','omics_screening/','omics_dataset_discovery/']
 assert not any(x.startswith(str(P.relative_to(R))+'/'+p) for x in paths for p in protected),paths
 print(json.dumps({'passed':True,'topical_candidates':len(rows),'processed':len(done),'remaining':progress['remaining'],'online_queries_completed':0,'backward_saved_references':42,'old_scope_preserved':True,'workbook_opened':True,'RIS_readable':True},ensure_ascii=False))
if __name__=='__main__':main()
