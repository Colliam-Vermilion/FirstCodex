"""Increment existing outputs without rerunning or altering the 233 initial screens."""
import json,hashlib,csv,copy
from pathlib import Path
from collections import Counter
from openpyxl import load_workbook
from openpyxl.styles import Alignment
from prepare import T,read,dump,now,title,doi
from render import csvout,ris

def sha(r):return hashlib.sha256(json.dumps(r,ensure_ascii=False,sort_keys=True).encode()).hexdigest()
def main():
 snap=T/'checkpoints/pre_online_233_master.json'
 if not snap.exists():dump(snap,read(T/'literature_master.json'))
 base=read(snap);assert len(base)==233;B={r['id']:r for r in base};R=copy.deepcopy(B)
 online=read(T/'online_records.json');O={r['id']:r for r in online};S={r['id']:r for r in read(T/'online_screening.json')};F={r['id']:r for r in read(T/'online_fulltext_evidence.json')}
 prefix=str(T.relative_to(Path('/workspace/FirstCodex')))
 for rid,r in O.items():
  if rid not in R:
   R[rid]={**r,'status':'未开始在线结果筛选','scope_decision':'尚未判断，数据库命中不等于相关/纳入','reading_depth':'未读；仅取回元数据/摘要','source_role':'在线结果；复用全库身份' if r['existed_in_global_library'] else '在线结果；新身份待筛选','old_scope_judgment_changed':False,'source_links':[{'path':prefix+'/'+s['raw_path'],'id':s['record_locator'],'identity_field':'record_locator'} for s in r['online_source_links']]}
  else:R[rid].update(online_source_links=r['online_source_links'],prior_topical_row_sha256=sha(B[rid]),initial_screen_status_preserved=B[rid]['status'])
 for rid,s in S.items():
  assert rid not in B,'Never repeat old initial screenings'
  R[rid].update({k:v for k,v in s.items() if k not in ['source_links','title','id']})
 for rid,f in F.items():
  r=R[rid];r['previous_reading_depth']=r['reading_depth'];r.update(reading_depth=f['depth'],status='全文相关方法/结果核实',actual_finding=f['finding'],author_interpretation_and_alternatives=f['boundary'],system=f['system'],temperature=f['temperature'],aggregate_or_biofilm_size=f['size'],methods=f['methods'],locator=f['locator'],online_fulltext_review=f,new_external_fulltext_downloaded=True)
 rows=list(R.values());dump(T/'literature_master.json',rows);csvout(T/'literature_master_utf8.csv',rows)
 for r in read(T/'seed_audit.json'):
  rr=next(x for x in online if doi(x.get('doi'))==doi(r['supplied_doi']));r['assigned_existing_id']=rr['id'];r['fulltext_state']='全文方法/结果核实' if rr['id'] in F else '元数据/完整正式摘要核实；正文未核实';r['current_reading_source']='online_fulltext_evidence.json' if rr['id'] in F else 'seed_abstract_evidence.json'
  # Preserve verified seed state for citation register below.
  r['online_record_id']=rr['id']
  if 'seeds' not in locals():seeds=[]
  seeds.append(r)
 dump(T/'seed_audit.json',seeds);csvout(T/'seed_audit_utf8.csv',seeds)
 # Current full-text task list, deduplicated by existing ID; preserve all old gaps.
 pending={}
 for n in read(T/'fulltext_pending.json'):
  rid=n.get('id')
  if not rid and n.get('seed_key'):rid=next(s['assigned_existing_id'] for s in seeds if s['seed_key']==n['seed_key'])
  if rid and rid not in F:pending[rid]={**n,'id':rid,'status':'待全文/具体疑点核实；网络已恢复，不代表原件可得'}
 for rid,s in S.items():
  if s['status']=='摘要初筛' and s['scope_decision']!='暂不优先' and rid not in F:
   pending[rid]={'id':rid,'doi':R[rid].get('doi'),'title':R[rid]['title'],'next_action':'正文方法/结果、参数条件和空间尺度核实','specific_gap':s['author_interpretation_and_alternatives'],'status':'待全文/补证'}
 for rid in F:pending.pop(rid,None)
 needs=list(pending.values());dump(T/'fulltext_pending.json',needs);csvout(T/'fulltext_pending_utf8.csv',needs)
 remaining=[{'id':r['id'],'title':r['title'],'next_action':'按题名/完整摘要判断；不可因词匹配自动排除','status':'尚未初筛'} for r in rows if r['id'] not in B and r['id'] not in S]
 dump(T/'remaining_queue.json',remaining);csvout(T/'remaining_queue_utf8.csv',remaining)
 ap=[x for x in read(T/'abstract_pending_queue.json') if x['id'] in B]
 ap+=[{'id':r['id'],'title':r['title'],'next_action':'题名已阅，读取完整正式摘要；保留未判断状态','status':'待完整摘要'} for r in S.values() if r['status'] in ['题名已阅待摘要','片段待正式摘要','来源字段异常待摘要']]
 dump(T/'abstract_pending_queue.json',ap);csvout(T/'abstract_pending_queue_utf8.csv',ap)
 logs=read(T/'reading_log.json');oldlog={r['id']:r for r in logs}
 for rid,s in S.items():oldlog[rid]={'id':rid,'date_recorded_utc':s['screened_utc'],'basis':s['reading_depth'],'old_judgment_changed':False,'source_links':s['source_links']}
 for rid,f in F.items():oldlog.setdefault(rid,{'id':rid})['online_fulltext_update']={'date_recorded_utc':f['recorded_utc'],'basis':f['depth'],'raw_path':f['source_path'],'locator':f['locator']}
 dump(T/'reading_log.json',list(oldlog.values()));csvout(T/'reading_log_utf8.csv',list(oldlog.values()))
 # Store pre-online evidence views once; future reruns replace only this increment.
 for name in ['structure_localization_evidence','gas_transfer_parameters','cell_interfloc_exchange','oxygen_control_evidence']:
  b=T/'checkpoints'/('pre_online_'+name+'.json')
  if not b.exists():dump(b,read(T/(name+'.json')))
 structure=read(T/'checkpoints/pre_online_structure_localization_evidence.json')
 for rid in list(F)+['SG-52CE8A5581BA','SG-7226DB9C69DF','SG-5BF06DB69D22','SG-16B4E8D37756']:
  r=R[rid];f=F.get(rid);structure.append({'id':rid,'title':r['title'],'doi':r.get('doi'),'reading_depth':r['reading_depth'],'finding':f['finding'] if f else r['actual_finding'],'limits':f['boundary'] if f else r['author_interpretation_and_alternatives'],'spatial_scale':f.get('spatial_scale') if f else '摘要层；具体成像范围待正文','three_dimensional_measured':f.get('true_whole_aggregate_3D_verified') if f else '3D距离/制样需正文' if rid in ['SG-52CE8A5581BA','SG-5BF06DB69D22'] else False,'evidence_update':'20261030_online'})
 params=read(T/'checkpoints/pre_online_gas_transfer_parameters.json')
 # Keep prior scalar S02 extraction as historical; newest row supplies full conditions.
 for p in params:
  if p['id']=='ODS-PMID17905889':p['superseded_by_online_fulltext']='20261030_online_Deff_H2'
 def parameter(rid,symbol,value,unit,acquisition,condition,locator,limits):
  r=R[rid];params.append({'id':rid,'title':r['title'],'doi':r.get('doi'),'reading_depth':r['reading_depth'],'symbol':symbol,'value_original':value,'unit_original':unit,'acquisition':acquisition,'condition':condition,'source_locator':locator,'limits':limits,'evidence_update':'20261030_online'})
 s2='ODS-PMID17905889'
 parameter(s2,'Deff,H2','(2.66 ± 0.13) × 10^-5','cm² s^-1','实测瞬态＋扩散模型拟合；n=3，r=0.98','35°C；约2 mm颗粒；纯氯仿10 min失活后冲洗；中心微电极；57±3% Dw35','PMC2168234 Methods/Estimation of Deff；Results/Fig6','SI=(2.66±0.13)e-9 m²/s；不是活性反应净参数，不可直接赋O2')
 parameter(s2,'Deff,CH4','0.57 × Dw,CH4 at 35°C','相对水扩散系数','作者沿用H2相对扩散比的模型假设','同文剖面反演','PMC2168234 Methods/Microbial activity calculations','非独立CH4扩散实测；不另造绝对值')
 parameter(s2,'net CH4 production maximum','11.5 ± 3.9','µmol cm^-3 h^-1','稳态浓度剖面＋Fick二阶反演','35°C活性颗粒；约300 µm深度','PMC2168234 Results/Fig5','CH4粗探头与300 µm步长可能影响剖面；不同气体用不同颗粒')
 parameter(s2,'H2 net surface release','2.3 ± 4.3','µmol cm^-2 h^-1','浓度梯度＋Fick一阶计算','35°C活性颗粒；无NaHCO3剖面情景','PMC2168234 Results/Fig4–5','不是直接气泡排放仪实测；SD大，保留不确定度')
 parameter('SG-52CE8A5581BA','interspecies distance','5.30 → 0.29','µm','摘要称FISH＋3D图像距离测量','丙酸富集；絮体可达54 µm；温度/距离定义未读正文','PMID23233207 official abstract','不是不同絮体间距离')
 parameter('SG-52CE8A5581BA','maximum possible H2 flux','1.1 → 10.3','nmol ml^-1 min^-1','作者计算上限（摘要）','与上述距离变化关联；输入/边界待正文','PMID23233207 official abstract','体积归一化计算通量；非实测，不与面积通量直接比较')
 parameter('SG-7E776DC9CBE3','central O2 response time ratio','5','fold vs original flocs','微电极瞬态比较（正式摘要）','絮凝后；绝对尺寸、温度和反应状态待正文','PMID15932786 official abstract','不是5倍氧剂量或5倍扩散系数')
 parameter('SG-7E776DC9CBE3','Deff,O2 ratio','2.3','fold vs original flocs','正式摘要报告的有效扩散比较；拟合程序待正文','絮凝后；绝对Deff/温度/尺寸待核','PMID15932786 official abstract','系数升高但整体响应更慢，尺寸增长可主导')
 parameter('SG-4CA688CD70DB','Deff,FeEDTA','2.8 × 10^-11','m² s^-1','3D MRI瞬态＋球扩散模型（摘要）','3–4 mm产甲烷颗粒；109×109×218 µm³；11 min；温度未核','official PubMed abstract','铁络合物非气体；核心模型失配，不能借作O2系数')
 parameter('SG-7CC189F27F28','relative water diffusion','0.36–0.96','relative to free water','PFG-NMR水自扩散分布（摘要）','5类生物聚集体；多组分/异质模型','official PubMed abstract','不是H2、CH4或O2有效扩散系数')
 parameter('SG-0EC2CB10A7BE','allowable H2 exchange distance','2 / 16 / 32','µm','Fick计算允许距离（摘要）','热嗜共培养；分别丙酸/乙醇/1-丙醇','PMID16332758 official abstract','不是直接测得菌间距/絮体间距离')
 parameter('SG-28F1C544B55F','theoretical H2 transfer','2.5 (alternative scenario 0.3)','nmol ml^-1 min^-1','Fick计算；假定10 µm距离，2.5e8 cells/ml','30°C H2监测；表面浓度取21/52 nM假设，另用3.2 Pa情景','PMC124120 Discussion末段','非直接通量；18.6 nmol ml^-1 min^-1为产乙酸对应H2当量，勿混算')
 parameter('SG-0269FB002A30','kLa scenarios, gas-specific scaling','60–1000','d^-1','学位论文模型情景输入/按扩散率缩放（返回摘要）','H2/CO2/CH4 ADM1扩展；不是本轮物理测量','OpenAlex returned abstract','不得与实测kLa合并平均；参数不是微曝气AD实测范围')
 parameter('SG-2544205CF8CD','areal methane production','24.3','L m^-2 day^-1','膜生物反应器产甲烷表现（返回摘要）','PDMS膜供H2/CO2、热嗜体系；微量元素补加后','OpenAlex returned abstract','不是H2传质系数/渗透率；面积基准、气体标况/温度待正文')
 exchange=read(T/'checkpoints/pre_online_cell_interfloc_exchange.json')
 for rid in ['SG-52CE8A5581BA','SG-28F1C544B55F','SG-0EC2CB10A7BE','SG-69FE36B49E26','SG-C8E80693DD79','SG-FA6F3124E19D','SG-1411D1E4576E']:
  r=R[rid];exchange.append({'id':rid,'title':r['title'],'doi':r.get('doi'),'reading_depth':r['reading_depth'],'finding':r['actual_finding'],'limits':r['author_interpretation_and_alternatives'],'compartment':'菌间/絮体内或分散细胞模型；未核实絮体间通量','direct_interfloc_flux_verified':False,'evidence_update':'20261030_online'})
 controls=read(T/'checkpoints/pre_online_oxygen_control_evidence.json')
 for rid in ['SG-A87DE61E8FAB','SG-D84CF676664B','SG-E3EDCB94A4B6']:
  r=R[rid];controls.append({'id':rid,'title':r['title'],'doi':r.get('doi'),'reading_depth':r['reading_depth'],'finding':r['actual_finding'],'limits':r['author_interpretation_and_alternatives'],'type':'ORP反馈供氧；正式摘要证据，局部DO与完整控制逻辑待全文','evidence_update':'20261030_online'})
 r=R['SG-3281766B8B59'];controls.append({'id':r['id'],'title':r['title'],'doi':r.get('doi'),'reading_depth':r['reading_depth'],'finding':r['actual_finding'],'limits':r['author_interpretation_and_alternatives'],'type':'局部ORP/DO实测，非反馈控制；非产甲烷AD','evidence_update':'20261030_online'})
 controls.append({'id':s2,'title':R[s2]['title'],'reading_depth':R[s2]['reading_depth'],'finding':'流动池约-200 mV，铂微电极对Ag/AgCl；不是ORP供氧设定点','limits':'实验N2鼓泡+还原剂维持厌氧，无O2剖面或DO控制','type':'ORP测量，非氧反馈控制','evidence_update':'20261030_online'})
 for collection in [structure,params,exchange,controls]:
  for item in collection:
   if item.get('evidence_update')=='20261030_online':
    rr=R[item['id']];item.setdefault('source_links',rr.get('online_source_links',[]));item.setdefault('source_locator',rr.get('locator') or 'See online_source_links and the complete returned abstract')
 for name,data in [('structure_localization_evidence',structure),('gas_transfer_parameters',params),('cell_interfloc_exchange',exchange),('oxygen_control_evidence',controls)]:dump(T/(name+'.json'),data);csvout(T/(name+'_utf8.csv'),data)
 # Raw Crossref references and online citation work records remain different layers.
 crossrefs=[];idx={doi(r.get('doi')):r['id'] for r in online if r.get('doi')}
 for s in read(T/'seed_crossref_results.json'):
  for i,ref in enumerate(s['work'].get('reference',[]),1):crossrefs.append({'trace_id':s['seed_key']+'-CR-'+str(i).zfill(3),'seed_key':s['seed_key'],'direction':'backward','reference_as_returned':ref,'assigned_id':idx.get(doi(ref.get('DOI'))),'status':'metadata identity resolved, reading state follows literature_master' if idx.get(doi(ref.get('DOI'))) else 'unresolved Crossref reference; retained lead, not an included article','raw_path':s['raw_path']})
 dump(T/'crossref_backward_reference_register.json',crossrefs)
 traces=read(T/'online_citation_trace_log.json');cq=read(T/'online_citation_queue.json')
 for r in cq:
  r.update(state='complete first-generation OpenAlex metadata pages; Crossref references retained; reading not complete',identity_state='Crossref+PubMed verified',backward_records=next(x['retrieved_records'] for x in traces if x['seed_key']==r['seed_key'] and x['direction']=='backward'),forward_records=next(x['retrieved_records'] for x in traces if x['seed_key']==r['seed_key'] and x['direction']=='forward'))
 dump(T/'online_citation_queue.json',cq)
 # Update the earlier 42 saved S02 reference tasks only when PMID matches.
 bypm={str(r.get('pmid')):r for r in online if r.get('pmid')}
 follow=read(T/'citation_followup_queue.json')
 for r in follow:
  x=bypm.get(str(r.get('pmid')))
  if x:r.update(existing_id=x['id'],processing_state=R[x['id']]['status'],online_metadata_available=True)
 dump(T/'citation_followup_queue.json',follow)
 ris(T/'spatial_gas_transfer_candidates.ris',rows)
 done=[r for r in rows if r['id'] in B or r['id'] in S or r['id'] in F];ris(T/'spatial_gas_transfer_reviewed.ris',done)
 for filename in ['spatial_gas_transfer_candidates.ris','spatial_gas_transfer_reviewed.ris']:
  path=T/filename;lines=[]
  for line in path.read_text().splitlines():
   if line.strip()=='TI  -':lines.append('N1  - Source metadata lacks a title; retain identifier and verify before use.')
   elif line.strip()=='AU  -':continue
   else:lines.append(line.rstrip())
  path.write_text('\n'.join(lines).rstrip()+'\n')
 for name in ['online_screening','online_search_log','online_citation_trace_log','online_fulltext_evidence','seed_abstract_evidence','online_identity_conflicts']:csvout(T/(name+'_utf8.csv'),read(T/(name+'.json')))
 wb=load_workbook(T/'spatial_gas_transfer_cumulative.xlsx')
 sheets={'文献与专题筛选':rows,'三维结构与定位':structure,'传质参数及剖面':params,'细胞间与絮体间':exchange,'供氧控制':controls,'在线待执行检索':read(T/'online_query_queue.json'),'在线查询实际状态':read(T/'online_search_log.json'),'在线引用追溯':traces,'种子核对':seeds,'待全文与疑点':needs,'未处理续跑':remaining,'在线逐篇判断':list(S.values()),'本轮全文核实':list(F.values()),'待完整摘要':ap}
 oversize=[]
 for name,data in sheets.items():
  if name in wb.sheetnames:del wb[name]
  ws=wb.create_sheet(name);keys=list(dict.fromkeys(k for r in data for k in r)) or ['status'];ws.append(keys)
  for r in data:
   vals=[]
   for k in keys:
    v=r.get(k);v=json.dumps(v,ensure_ascii=False) if isinstance(v,(list,dict)) else v
    if isinstance(v,str) and len(v)>32767:oversize.append({'sheet':name,'id':r.get('id'),'field':k,'characters':len(v)});v='Excel单元格长度限制：完整字段保存在同名JSON及online_records.json；ID='+str(r.get('id'))
    vals.append(v)
   ws.append(vals)
  ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
  for col in ws.columns:ws.column_dimensions[col[0].column_letter].width=28
  for row in ws.iter_rows():
   for c in row:c.alignment=Alignment(wrap_text=True,vertical='top')
 wb.save(T/'spatial_gas_transfer_cumulative.xlsx');dump(T/'excel_long_field_links.json',oversize)
 summary=read(T/'online_merge_summary.json');summary.update(updated_utc=now(),cumulative_spatial_records=len(rows),original233_initial_screens_preserved=233,new_initial_screenings=len(S),new_abstracts_read=sum(x['status']=='摘要初筛' for x in S.values()),new_title_only_pending_abstract=sum(x['status']=='题名已阅待摘要' for x in S.values()),new_fragment_or_source_field_gaps=sum(x['status'] in ['片段待正式摘要','来源字段异常待摘要'] for x in S.values()),new_related_abstract_screenings=sum(s['status']=='摘要初筛' and s['scope_decision']!='暂不优先' for s in S.values()),new_related_records=sum((s['status']=='摘要初筛' and s['scope_decision']!='暂不优先') or s['id'] in F for s in S.values()),new_fulltexts_methods_results_verified=len(F),new_3D_whole_aggregate_reconstruction_verified=1,citation_backward_metadata_records=sum(r['retrieved_records'] for r in traces if r['direction']=='backward'),citation_forward_metadata_records=sum(r['retrieved_records'] for r in traces if r['direction']=='forward'),crossref_reference_leads=len(crossrefs),unresolved_crossref_reference_leads=sum(not r['assigned_id'] for r in crossrefs),new_initial_screening_remaining=len(remaining),abstract_pending=len(ap),fulltext_evidence_tasks_remaining=len(needs),direct_interfloc_flux_verified=0,parameter_records=len(params),control_records=len(controls),excel_oversize_fields=len(oversize))
 dump(T/'online_increment_summary.json',summary)
 progress=read(T/'progress.json');progress.update(updated_utc=now(),candidates=len(rows),processed=len(done),remaining=len(remaining),new_external_articles_discovered=summary['new_to_global_library'],external_discovery_is_not_inclusion=True,new_external_queries_completed=12,online_status='original12 and first-generation citation pages complete; '+str(len(S))+' additional reading decisions with explicit depth; remaining not silently screened',seed_metadata_verified_online=5,seed_metadata_unverified=0,full_3D_reconstruction_papers_verified=1,new_fulltext_methods_results_verified=len(F),parameter_or_profile_records=len(params),control_records=len(controls),abstract_pending_after_title=len(ap),original233_initial_screen_ids=[r['id'] for r in base],completed_initial_screen_ids=[r['id'] for r in done],new_completed_initial_screen_ids=list(S),completed_ids=[r['id'] for r in done],next_id=remaining[0]['id'] if remaining else None,resume='Do not rerun prepare/render/update_guide. Use render_online for this increment; continue remaining_queue and abstract_pending; 233 initial decisions preserved in checkpoints/pre_online_233_master.json')
 progress.update(abstract_only_processed=sum(r['status']=='摘要初筛' for r in rows),title_only_processed=sum(r.get('reading_depth','').startswith('仅题名') for r in rows),new_fragment_or_source_field_gaps=summary['new_fragment_or_source_field_gaps'],new_external_bibliographic_identities=summary['new_to_global_library'],external_count_unit='provisional bibliographic identities including unverified types; not inclusions')
 dump(T/'progress.json',progress)
 oldsummary=read(T/'summary.json');oldsummary['online_increment']=summary;dump(T/'summary.json',oldsummary)
 checkpoint='ONLINE003.json' if (T/'followup_online_screening.json').exists() else 'ONLINE002.json'
 dump(T/'checkpoints'/checkpoint,summary)
 print(json.dumps(summary,ensure_ascii=False))
if __name__=='__main__':main()
