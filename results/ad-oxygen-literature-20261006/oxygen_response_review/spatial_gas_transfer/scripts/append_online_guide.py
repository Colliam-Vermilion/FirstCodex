"""Append new evidence to all four homepage parts and current citation maps."""
import json,re,hashlib,shutil
from pathlib import Path
from openpyxl import load_workbook
from prepare import T,M,R,read,dump,now
from render import csvout
G=R/'scientific_guide';PREFIX=T.relative_to(R).as_posix();MARK='20261030_online'
LINKS={'P1c':['ODS-PMID17905889','SG-7E776DC9CBE3','SG-2DCD3D787125'], 'P1d':['SG-A87DE61E8FAB','SG-D84CF676664B','SG-E3EDCB94A4B6'], 'P2d':['SG-7E776DC9CBE3','SG-A509A75F5507'], 'P3c':['SG-822165E8E403','ODS-PMID17905889','SG-52CE8A5581BA','SG-7226DB9C69DF','SG-5BF06DB69D22'], 'P4d':['SG-52CE8A5581BA','SG-28F1C544B55F','SG-69FE36B49E26','SG-1411D1E4576E'], 'P4e':['SG-52CE8A5581BA','SG-28F1C544B55F']}
def append_block(path,key,block,before=None):
 s=path.read_text();start='<!-- '+key+' START -->';end='<!-- '+key+' END -->';insert=start+'\n'+block.rstrip()+'\n'+end+'\n\n'
 if start in s:s=re.sub(re.escape(start)+r'.*?'+re.escape(end)+r'\n*',lambda m:insert,s,flags=re.S)
 elif before:s=s.replace(before,insert+before,1)
 else:s=s.rstrip()+'\n\n'+insert
 path.write_text(s)
def main():
 C={r['id']:r for r in read(T/'literature_master.json')};F={r['id']:r for r in read(T/'online_fulltext_evidence.json')}
 ids=set(sum(LINKS.values(),[]));originals=[]
 for rid in ids:
  r=C[rid];path=T/'evidence'/(rid+'.md')
  if not path.exists():path.write_text('# '+r['title']+'\n\n文献ID：`'+rid+'`。\n')
  f=F.get(rid);text='## 2026-10-30 在线补证\n\n阅读深度：**'+r['reading_depth']+'**。\n\n'+r['actual_finding']+'。\n\n边界：'+r['author_interpretation_and_alternatives']+'。\n'
  if f:
   text+='\n对象及状态：'+f['system']+'。\n\n尺度：'+f['spatial_scale']+'。\n\n方法：'+f['methods']+'。\n\n温度和尺寸：'+f['temperature']+'；'+f['size']+'。\n\n参数：'+f['quantitative']+'。\n\n取得方式：'+f['measurement_class']+'。\n\n应用：'+f['applicability']+'。\n\n定位：'+f['locator']+'。\n\n[官方HTML原始返回](../'+f['source_path']+') · [正文分节文本](../'+f['sections_path']+')。\n'
   pmcid={'SG-822165E8E403':'PMC11337819','ODS-PMID17905889':'PMC2168234','SG-28F1C544B55F':'PMC124120'}[rid]
   dest=M/'raw/pmc_fulltext'/(pmcid+'_spatial_online.html');dest.parent.mkdir(parents=True,exist_ok=True);shutil.copyfile(T/f['source_path'],dest)
   originals.append({'id':rid,'pmcid':pmcid,'path':dest.relative_to(R).as_posix(),'sha256':hashlib.sha256(dest.read_bytes()).hexdigest(),'raw_response':PREFIX+'/'+f['source_path'],'reading_depth':f['depth']})
   text+='\n[规范原件目录中的HTML](../../raw/pmc_fulltext/'+dest.name+')。原件为公开正文HTML，本轮没有将前置XML或验证页算作全文，也未声称下载了PDF。\n'
  else:
   src=(r.get('preferred_abstract_source') or (r.get('online_source_links') or [{}])[0]);p=src.get('raw_path')
   if p:text+='\n[本轮原始元数据/摘要](../'+p+')，定位 `'+str(src.get('record_locator'))+'`。\n'
  append_block(path,MARK,text)
 dump(T/'online_original_file_register.json',originals)
 def cite(rid,label):return '['+label+']('+PREFIX+'/evidence/'+rid+'.md)'
 s2=cite('ODS-PMID17905889','Satoh等，2007（全文方法/结果）');s1=cite('SG-822165E8E403','Urasaki等，2024（全文方法/结果）');s3=cite('SG-52CE8A5581BA','Felchner-Zwirello等，2013（正式摘要）');s4=cite('SG-7E776DC9CBE3','Chu等，2005（正式摘要）');s114=cite('SG-28F1C544B55F','de Bok等，2002（全文方法/结果）')
 blocks={
 'Part1':f'''### 在线补证 · 2026-10-30：系数、尺度与控制目标要一起读

{s2}把此前的H₂系数条件补齐：约2 mm颗粒先经纯氯仿失活10 min、冲洗，再于35°C用中心H₂微电极记录阶跃响应并拟合，得到 **(2.66±0.13)×10⁻⁵ cm²/s**（n=3；水中值的57±3%）。它是失活材料中的拟合扩散参数。活性颗粒的H₂/CH₄剖面另行测量；CH₄的相对扩散比57%是沿用H₂的假设，生成/释放速率由剖面反演。不能将这些值直接写成O₂的扩散或细胞暴露剂量。

{s4}报告絮凝后中心O₂响应时间约为原来的5倍，同时有效扩散系数却约增至2.3倍。尺寸、材料与反应项必须共同解释；目前仅读完整正式摘要，绝对尺寸、温度及拟合程序待正文。外液膜证据也单列：{cite('SG-2DCD3D787125','液膜DO边界层微电极研究（返回摘要）')}操纵流速等条件，不能把内部Deff、外部kf与气液kLa当同一参数。

ORP反馈服务于具体目标。{cite('SG-D84CF676664B','Khanal与Huang，2003（正式摘要）')}在UAF中把自然ORP约−290 mV提高25 mV：高硫进水时消硫改善甲烷，低硫时甲烷产率反而下降。另一个{cite('SG-E3EDCB94A4B6','ORP供氧实验（正式摘要）')}使用−230/−180 mV并发现异养耗COD竞争。这些是硫化物控制情景，不能推出跨反应器的通用安全ORP或局部DO。完整供氧迟滞、参比、输入量、耗氧和局部剖面仍待全文。
''',
 'Part2':f'''### 在线补证 · P2d：颗粒的氧响应能代表单细胞剂量吗？

{s4}提供聚集体中心O₂响应的证据；它没有在本轮已读材料中把每个细胞的实际O₂暴露、胞内ROS或耐氧阈值一并测出。{cite('SG-A509A75F5507','活性污泥O₂剖面研究（正式摘要）')}还显示，大絮体可出现近零O₂中心，小絮体却可能出现模型不能再现的非零平台。该体系不是产甲烷AD，只作传质方法参照。

因此，本轮给Individual-cell response增加的是**暴露解释的限制**：投氧量、主体DO和单点ORP不能直接替代细胞附近的浓度—时间历程。未用这些空间研究新增“某古菌在某DO下耐氧”的生理结论；这仍需细胞级测量或受控暴露实验。
''',
 'Part3':f'''### 在线补证 · 2026-10-30：已有真正三维重建，但没有测到孔隙通量

{s1}通过双能量同步辐射μCT重建厌氧颗粒孔洞。原始体素/像素为0.1988 µm，3D显示为0.7952 µm；体素不等同独立验证的成像分辨率。孔洞截面平均面积随深度为3.8±0.67至42±8.7 µm²，截面面积占比0.29±0.11%至12±4.2%；n=3指截取柱段，不能改写成3个独立生物重复。

这项方法的边界同样关键：**Os信号同时含细胞和EPS；gold-ISH未提供可检测的CT菌群特异信号**。样品经固定、脱水和树脂包埋，不是活体灌流。作者未观察到孔洞间连接通道，并指出许多小通道可能低于体素尺度；图像也不能区分孔中的水、气或其他轻元素物质。孔洞面积/体积不等于O₂流量、扩散系数或通透性。

{s2}则是20 µm切片CLSM/FISH与不同颗粒上的微剖面相互参照，不能称为同一颗粒的完整三维共注册。{cite('SG-7226DB9C69DF','Rocheleau等，1999（正式摘要）')}显示蛋白进水颗粒可以不分层，而甲醇富集颗粒出现分层；结构取决于底物和来源。{cite('SG-5BF06DB69D22','3D-FISH制样方法（正式摘要）')}也提醒固定/制样可能改变相对细胞位置。
''',
 'Part4':f'''### 在线补证 · 2026-10-30：距离实测、可能通量与真实交换要分开

{s3}的正式摘要报告FISH加3D图像所得菌间距离从5.30降至0.29 µm，而“最大可能H₂通量”从1.1升至10.3 nmol ml⁻¹ min⁻¹。前者是摘要所述图像距离，后者是**计算上限**，不是直接测得的菌间或絮体间通量；温度、距离定义和模型边界尚未从正文核实。该文目前不能标为全文确认。

{s114}在丙酸双菌培养中结合酶活、重组培养和气体时序支持甲酸的重要性。它采用**假定10 µm距离**与表面浓度计算H₂通量2.5 nmol ml⁻¹ min⁻¹（另一浓度情景0.3），再与乙酸生成对应的H₂当量18.6比较；这些不是三项同类实测通量。常规培养/酶测37°C、H₂监测试验30°C也不能混写。相反，{cite('SG-1411D1E4576E','H₂摄取传质限制研究（返回摘要）')}在不同动力学/传质假设下得出H₂扩散可承担所需电子通量，仍未排除甲酸。{cite('SG-69FE36B49E26','动态扩散—反应絮体模型（正式摘要）')}也应作为模型证据阅读。

这些结果支持“距离、载体、底物、温度与动力学共同约束互营”，不支持“H₂总不足”或“甲酸总主导”的通则。本轮仍未核实到不同絮体之间的距离—通量直接测量，也没有由聚集/甲烷产率变化直接证明氧保护区。
'''}
 root=R/'README.md'
 append_block(root,MARK+'_status','> **2026-10-30在线补证已恢复：**原12组专题查询已完整分页取回，5篇种子均核对。以下保留的较早专题段落是当时阅读深度的历史说明；各部分新增的“在线补证”及[现行专题进度]('+PREFIX+'/README.md)记录最新条件。取回完成不等于筛选或全文核实全部完成。',before='## Part 1.')
 for k,b in blocks.items():
  before={'Part1':'## Part 2.','Part2':'## Part 3.','Part3':'## Part 4.','Part4':'## 现有库中，哪些研究'}[k]
  append_block(root,MARK+'_'+k,b,before)
 # All prior map entries survive. Only this dated increment is replaced on rerun.
 props=read(G/'propositions.json')
 if not any(x['proposition_id']=='P2d' for x in props):props.append({'proposition_id':'P2d','part':'Part 2','question':'聚集体的氧响应能否代表单细胞实际氧剂量？','evidence_rule':'区分整体瞬态、局部浓度—时间与胞内生理；非AD方法参照单列'})
 prop={r['proposition_id']:r for r in props};relations=[]
 for pid,ids_ in LINKS.items():
  for rid in ids_:
   r=C[rid];relations.append({'id':rid,'title':r['title'],'year':r.get('year'),'journal':r.get('journal'),'doi':r.get('doi'),'part':prop[pid]['part'],'proposition_id':pid,'proposition_question':prop[pid]['question'],'finding':r['actual_finding'],'boundary':r['author_interpretation_and_alternatives'],'reading_depth':r['reading_depth'],'mapping_status':'在线补证；不改变旧全库范围判断','support_role':'范围、测量与限制；不是全部命题的独立机制验证','evidence_card':PREFIX+'/evidence/'+rid+'.md','evidence_update':MARK,'source_links':r.get('online_source_links',[])})
 sp=[r for r in read(G/'spatial_gas_proposition_links.json') if r.get('evidence_update')!=MARK]+relations
 mp=[r for r in read(G/'literature_proposition_map.json') if r.get('evidence_update')!=MARK];oldids={r['id'] for r in mp};linked={r['id'] for r in relations}
 mp+=relations
 for rid,r in C.items():
  if rid not in oldids and rid not in linked:mp.append({'id':rid,'title':r['title'],'doi':r.get('doi'),'year':r.get('year'),'proposition_id':None,'mapping_status':'在线结果尚无命题关联；'+r['status'],'reading_depth':r['reading_depth'],'evidence_update':MARK})
 for name,data in [('spatial_gas_proposition_links',sp),('literature_proposition_map',mp),('propositions',props)]:dump(G/(name+'.json'),data);csvout(G/(name+'_utf8.csv'),data)
 wb=load_workbook(G/'scientific_reading_guide.xlsx')
 for name,data in [('全部文献—命题',mp),('科学命题',props),('空间专题关联',sp),('在线补证逐篇',read(T/'online_screening.json')),('在线全文证据',list(F.values())),('在线查询状态',read(T/'online_search_log.json'))]:
  if name in wb.sheetnames:del wb[name]
  ws=wb.create_sheet(name);keys=list(dict.fromkeys(k for r in data for k in r));ws.append(keys)
  for r in data:ws.append([json.dumps(r.get(k),ensure_ascii=False) if isinstance(r.get(k),(list,dict)) else r.get(k) for k in keys])
  ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
 wb.save(G/'scientific_reading_guide.xlsx')
 summary=read(G/'summary.json');summary.update(spatial_external_discovery_completed=True,spatial_screening_completed=False,spatial_online_queries_completed=12,spatial_online_fulltext_verified=3,spatial_candidates=len(C),spatial_online_new_links=len(relations),spatial_proposition_links=len(sp),expanded_map_rows=len(mp),expanded_map_unique_ids=len({r['id'] for r in mp}));dump(G/'summary.json',summary)
 append_block(G/'README.md',MARK,'## 2026-10-30 在线专题补证\n\n原12组查询已完整取回，5篇种子双源核对，3篇读取了正文方法/结果。新增'+str(len(relations))+'条带阅读深度的命题关系；Part2新增P2d，仅界定细胞暴露解释的边界。旧文献与命题全部保留，未关联的在线命中仍标“尚无命题关联”，不计正式纳入。\n\n[最新专题进度](../'+PREFIX+'/README.md) · [当前累计Excel](../'+PREFIX+'/spatial_gas_transfer_cumulative.xlsx) · [新全文证据](../'+PREFIX+'/online_fulltext_evidence.json)。')
 dump(T/'guide_online_update.json',{'updated_utc':now(),'new_proposition_links':len(relations),'homepage_locations':['Part1末尾 在线补证','Part2末尾 P2d','Part3末尾 在线补证','Part4末尾 在线补证'],'old_homepage_content_preserved':True,'proposition_map_rows':len(mp)})
 print('new dated proposition links',len(relations))
if __name__=='__main__':main()
