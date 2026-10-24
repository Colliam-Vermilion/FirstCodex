#!/usr/bin/env python3
"""Validate deliverables against source bytes, locators and unchanged AD data."""
import csv,hashlib,json,re,xml.etree.ElementTree as ET
from pathlib import Path
from collections import Counter
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parents[1];PROJECT=ROOT.parent
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def norm(t):return ' '.join(t.split())
def run():
 baseline=json.loads((ROOT/'baseline.json').read_text())
 old_checks={p:sha(PROJECT/p)==expected for p,expected in baseline['sha256'].items()}
 assert all(old_checks.values()),old_checks
 master=list(csv.DictReader((PROJECT/'final_results/ad_oxygen_download_review_utf8.csv').open(encoding='utf-8-sig')))
 cat=json.loads((ROOT/'catalog.json').read_text());ids=[r['id'] for r in cat];assert len(ids)==len(set(ids))
 assert {r['稳定记录ID'] for r in master}<={r['id'] for r in cat}
 old_by_id={r['稳定记录ID']:r for r in master}
 for r in cat:
  if r['id'] in old_by_id:
   o=old_by_id[r['id']];assert r.get('old_ad_classification')==o['最终分类'];assert r.get('old_human_classification','')==o.get('人工最终分类','')
 dec=json.loads((ROOT/'screening.json').read_text());selection=json.loads((ROOT/'batch_B001_selection.json').read_text());assert len(dec)==20;assert {r['id'] for r in dec}=={r['id'] for r in selection}
 quote_checks=[];source_checks=[]
 for r in dec:
  src=ROOT/r['source_file'];assert src.exists() and sha(src)==r['source_sha256'],r['id']
  s=json.loads((ROOT/'evidence_sources'/f'{r["id"]}.json').read_text());loc={p['locator']:p for p in s['paragraphs']}
  if s['source_kind']=='JATS':assert ET.fromstring(src.read_bytes()).find('.//body') is not None
  if r['reading_level']=='仅题名/摘要':assert s['source_kind']=='abstract';assert r['screening_stage']!='全文核实'
  e=r['evidence']
  for q,l in [(e['quote'],e['quote_location']),(e['additional_quote'],e['additional_quote_location'])]:
   assert l in loc and norm(q) in norm(loc[l]['text']),(r['id'],l)
   quote_checks.append({'id':r['id'],'locator':l,'exact_excerpt_in_saved_source':True})
  source_checks.append({'id':r['id'],'sha256':r['source_sha256'],'kind':s['source_kind'],'path':r['source_file'],'verified':True})
 events=[json.loads(t) for t in (ROOT/'request_log.jsonl').read_text().splitlines()]
 for e in events:
  f=ROOT/e['path'];assert f.stat().st_size==e['bytes'] and sha(f)==e['sha256'],e['path']
  assert not any(x in e['url_without_contact'].lower() for x in ['api_key=','email=','mailto=','token='])
 logs=json.loads((ROOT/'search_log.json').read_text());inputs=list(csv.DictReader((ROOT/'dedup_input_mapping_utf8.csv').open(encoding='utf-8-sig')))
 returned=[r for r in inputs if r['query_id']!='OLD554']
 assert sum(r.get('retrieved_records',0) for r in logs)==len(returned)
 for l in logs:
  q=l['query_id'];assert sum(r['query_id']==q for r in returned)==l['retrieved_records'],q
  if q.startswith('P'):assert l['reported_count']==l['retrieved_records']==l['returned_ids']
 d=json.loads((ROOT/'dedup_summary.json').read_text());assert len({r['assigned_id'] for r in returned})==d['unique_returned_records']
 queue=json.loads((ROOT/'remaining_queue.json').read_text());order=json.loads((ROOT/'processing_order.json').read_text())
 done={r['id'] for r in dec};assert len(order)==len(cat);assert {r['id'] for r in order}==set(ids)
 assert {r['id'] for r in queue}==set(ids)-done;assert not done & {r['id'] for r in queue}
 wb=load_workbook(ROOT/'oxygen_response_review_cumulative.xlsx',read_only=True,data_only=False)
 expected={'文献总表':len(cat),'机制证据':len(dec),'重点阅读':len(dec),'检索日志':len(logs),'引文追溯':len(json.loads((ROOT/'citation_trace.json').read_text())),'续跑队列':len(queue)}
 xchecks={}
 for name,count in expected.items():
  ws=wb[name];assert ws.max_row-1==count,(name,ws.max_row,count);xchecks[name]=count
  for row in ws.iter_rows():
   assert not any(c.data_type=='e' for c in row),name
 ws=wb['文献总表'];headers=[c.value for c in next(ws.rows)];idcol=headers.index('id');assert {r[idcol] for r in ws.iter_rows(min_row=2,values_only=True)}==set(ids)
 wb.close()
 rischecks={}
 for fname,expectedids in [('oxygen_response_all_candidates.ris',set(ids)),('oxygen_response_B001.ris',done)]:
  t=(ROOT/fname).read_text();risids=re.findall(r'^ID  - (.+)$',t,re.M);assert len(risids)==len(set(risids));assert set(risids)==expectedids;assert len(re.findall(r'^TY  - ',t,re.M))==len(risids)==len(re.findall(r'^ER  - ',t,re.M));rischecks[fname]=len(risids)
 result={'baseline_commit':baseline['source_commit'],'old_file_sha_checks':old_checks,'original_master_records':len(master),'original_columns':len(master[0]),'all_original_ids_preserved':True,'all_original_ad_and_human_decisions_preserved':True,'screened_candidates':len(dec),'source_file_checks':source_checks,'exact_quote_checks':quote_checks,'raw_successful_response_files_checked':len(events),'normalized_source_returns':len(returned),'unique_returned_ids':d['unique_returned_records'],'old_ids_linked_to_returns':d['original_ids_linked_to_returns'],'new_to_old_library_candidates':d['unique_returned_records']-d['original_ids_linked_to_returns'],'catalog_records':len(cat),'queue_records':len(queue),'queue_has_no_completed_records':True,'workbook_sheets':xchecks,'ris_records':rischecks,'screening_stages':dict(Counter(r['screening_stage'] for r in dec)),'read_levels':dict(Counter(r['reading_level'] for r in dec)),'full_read_of_every_page_claimed':False,'sequencing_raw_data_downloaded':False,'validation':'passed'}
 (ROOT/'validation.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8');print(json.dumps({k:v for k,v in result.items() if k not in ['source_file_checks','exact_quote_checks']},ensure_ascii=False))
if __name__=='__main__':run()
