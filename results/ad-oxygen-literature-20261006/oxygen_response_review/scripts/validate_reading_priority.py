#!/usr/bin/env python3
"""Check retained results, source metrics, claim links, workbook and RIS."""
import collections,datetime as dt,hashlib,json,re,xml.etree.ElementTree as ET
from pathlib import Path
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parents[1];OUT=ROOT/'reading_priority'
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def run():
 b=json.loads((OUT/'baseline.json').read_text());protected={p:sha(ROOT/p)==expected for p,expected in b['review_file_sha256'].items()};assert all(protected.values()),protected
 original=json.loads((ROOT/'baseline.json').read_text());ad_checks={p:sha(ROOT.parent/p)==expected for p,expected in original['sha256'].items()};assert all(ad_checks.values())
 ev=json.loads((OUT/'evaluations.json').read_text());screen=json.loads((ROOT/'screening.json').read_text());sm={r['id']:r for r in screen};assert len(ev)==len(screen)==20;assert {r['文献ID'] for r in ev}==set(sm)
 selected=[r for r in ev if r['正式精读推荐']];held=[r for r in ev if not r['正式精读推荐']]
 for r in ev:
  s=sm[r['文献ID']];assert r['原筛选状态']==s['screening_stage'];assert r['原范围结论']==s['screening_decision']
  if r['正式精读推荐']:
   assert not s['screening_decision'].startswith('待确认')
   if s['reading_level']=='仅题名/摘要':assert r['推荐确认状态'].startswith('暂定')
  if r['推荐主层级']=='核心机制文献':assert r['AD比较等级']==3 and r['机制直接性等级']>=3
  if r['JIF2024'] is not None:
   assert r['JIF指标年份']==2024 and r['JCR发布年份']==2025 and r['期刊指标快照'] and r['期刊指标来源']
   assert (ROOT/r['期刊指标快照']).exists()
  else:assert '留空' in r['期刊指标状态']
 raw=json.loads((OUT/'raw/openalex_citations_20.json').read_text());oa={(r['doi'].lower().split('doi.org/')[-1]):r for r in raw['results']}
 for r in ev:
  source=oa.get(r['DOI']);assert source is not None;assert source['cited_by_count']==r['引用次数'];assert r['引用数据来源']=='OpenAlex'
  days=(dt.date(2026,10,24)-dt.date.fromisoformat(source['publication_date'])).days;assert r['年均引用']==round(source['cited_by_count']/(days/365.2425),3)
 req=[json.loads(x) for x in (OUT/'request_log.jsonl').read_text().splitlines()]
 for r in req:assert (ROOT/r['path']).stat().st_size==r['bytes'] and sha(ROOT/r['path'])==r['sha256']
 meta=ET.fromstring((OUT/'raw/pubmed_selected_and_support.xml').read_bytes());articles={a.findtext('MedlineCitation/PMID'):a for a in meta.findall('PubmedArticle')}
 links=json.loads((OUT/'mechanism_validation_links.json').read_text());assert {r['focal_id'] for r in links}==set(sm)
 for r in links:
  assert r['strict_same_object_independent_replication_confirmed'] is False
  if r.get('related_doi'):
   a=articles[r['related_pmid']];texts=[' '.join(''.join(z.itertext()).split()) for z in a.findall('MedlineCitation/Article/Abstract/AbstractText')]
   assert any(r['quote'] in t for t in texts),(r['claim_id'],r['related_pmid'])
   doi=next((''.join(z.itertext()).lower() for z in a.findall('PubmedData/ArticleIdList/ArticleId') if z.get('IdType')=='doi'),'')
   if doi:assert doi==r['related_doi']
   else:assert r['related_pmid']=='8226748' and r['related_doi']=='10.1016/s0021-9258(18)41538-4'
   if r['validation_relation_code']=='same_team_follow_up':assert r['shared_listed_authors'],r['claim_id']
   if '不同团队' in r['relation']:assert not r['shared_listed_authors'],r['claim_id']
 fullproof=json.loads((OUT/'support_fulltext_excerpts.json').read_text())
 for r in fullproof:
  f=ROOT/r['source'];assert sha(f)==r['source_sha256'];xml=ET.fromstring(f.read_bytes());m=re.search(r'@id="([^"]+)"\]/p\[(\d+)\]',r['locator']);assert m
  s=next(s for s in xml.findall('.//sec') if s.get('id')==m.group(1));p=s.findall('p')[int(m.group(2))-1];assert r['quote'] in ' '.join(''.join(p.itertext()).split())
 wb=load_workbook(ROOT/'oxygen_response_review_cumulative.xlsx',read_only=True);sheets={}
 for name,info in b['existing_workbook_sheets'].items():
  ws=wb[name];digest=hashlib.sha256(json.dumps(list(ws.values),ensure_ascii=False,default=str).encode()).hexdigest();assert digest==info['values_sha256'];sheets[name]={'retained_cell_values':True,'rows':ws.max_row-1}
 for name,n in [('优先精读',len(ev)),('机制验证关联',len(links)),('评价口径',len(json.loads((OUT/'evaluation_policy.json').read_text())))]:
  assert wb[name].max_row-1==n;sheets[name]={'rows':n}
 ws=wb['优先精读'];headers=[c.value for c in next(ws.rows)];idcol=headers.index('文献ID');assert {r[idcol] for r in ws.iter_rows(min_row=2,values_only=True)}==set(sm)
 for ws in wb:
  for row in ws:
   assert not any(c.data_type=='e' for c in row)
 wb.close()
 ris=(OUT/'priority_reading_set.ris').read_text();ids=re.findall(r'^ID  - (.+)$',ris,re.M);assert len(ids)==len(set(ids))==len(selected);assert set(ids)=={r['文献ID'] for r in selected};assert len(re.findall(r'^TY  - ',ris,re.M))==len(re.findall(r'^ER  - ',ris,re.M))==len(selected)
 for r in selected:
  if r['推荐确认状态'].startswith('暂定'):assert r['推荐确认状态'] in ris
 result={'validation':'passed','date_Asia_Shanghai':'2026-10-24','protected_review_file_checks':protected,'original_AD_file_checks':ad_checks,'all_B001_ids_assessed_without_reclassification':True,'evaluations':len(ev),'recommended':len(selected),'held':len(held),'recommendation_primary_tiers':dict(collections.Counter(r['推荐主层级'] for r in ev)),'tentative_recommendations':sum(r['推荐确认状态'].startswith('暂定') for r in selected),'citations_matched_to_DOI':len(ev),'verified_annualized_calculations':len(ev),'jif2024_available':sum(r['JIF2024'] is not None for r in ev),'raw_API_snapshots_sha_checked':len(req),'claim_relation_rows':len(links),'exact_support_abstract_excerpts_checked':sum(bool(r.get('quote')) for r in links),'fulltext_excerpts_checked':len(fullproof),'original_workbook_cells_preserved':True,'workbook_sheets':sheets,'ris_records':len(ids),'old_screening_queue_unchanged':True,'strict_independent_replication_not_inferred_from_citations':True}
 (OUT/'validation.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8');print(json.dumps(result,ensure_ascii=False))
if __name__=='__main__':run()
