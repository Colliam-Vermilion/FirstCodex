#!/usr/bin/env python3
"""Verify GitHub tree objects for all committed round files and read key bytes.

Uses the existing GH_TOKEN binding through HTTPS; never writes its value.
Does not require a second full clone of the existing large AD corpus.
"""
import argparse,concurrent.futures,datetime as dt,hashlib,json,os,subprocess
from pathlib import Path
from urllib.parse import quote,urlencode
from urllib.request import Request,urlopen
ROOT=Path(__file__).resolve().parents[1];REPO=ROOT.parents[2]
PREFIX='results/ad-oxygen-literature-20261006/oxygen_response_review/'
API='https://api.github.com/repos/Colliam-Vermilion/FirstCodex/'
def get(path,raw=False):
 headers={'Accept':'application/vnd.github.raw+json' if raw else 'application/vnd.github+json','User-Agent':'OxygenReviewDeliveryVerification/1.0'}
 if os.environ.get('GH_TOKEN'):headers['Authorization']='Bearer '+os.environ['GH_TOKEN']
 with urlopen(Request(API+path,headers=headers),timeout=60) as r:return r.read()
def run(commit):
 c=json.loads(get('git/commits/'+commit));tree=json.loads(get('git/trees/'+c['tree']['sha']+'?recursive=1'));assert not tree.get('truncated')
 remote={r['path']:r['sha'] for r in tree['tree'] if r['type']=='blob'}
 files=subprocess.check_output(['git','ls-tree','-r','--name-only',commit,'--',PREFIX],cwd=REPO,text=True).splitlines()
 objects=[]
 for rel in files:
  b=(REPO/rel).read_bytes();oid=hashlib.sha1(b'blob '+str(len(b)).encode()+b'\0'+b).hexdigest();assert remote.get(rel)==oid,rel
  objects.append({'path':rel,'bytes':len(b),'sha256':hashlib.sha256(b).hexdigest(),'git_blob_oid':oid,'remote_tree_matches':True})
 names=['README.md','PROTOCOL_zh.md','MECHANISM_SYNTHESIS_zh.md','oxygen_response_review_cumulative.xlsx','oxygen_response_B001.ris','oxygen_response_all_candidates.ris','catalog.json','literature_master_utf8.csv','mechanism_evidence.json','mechanism_evidence_utf8.csv','batch_B001_results_utf8.csv','screening.json','priority_reading_utf8.csv','remaining_queue.json','remaining_queue_utf8.csv','processing_order.json','search_log.json','search_log_utf8.csv','citation_trace.json','summary.json','validation.json','baseline.json','source_document_register.json']
 dest=Path('/tmp/oxygen-response-remote-readback')/commit;dest.mkdir(parents=True,exist_ok=True)
 def read_one(name):
  path=PREFIX+name;b=get('contents/'+quote(path,safe='/')+'?'+urlencode({'ref':commit}),True)
  # A proxy/server ignoring raw Accept may return the content envelope.
  if b.startswith(b'{'):
   try:j=json.loads(b)
   except Exception:j={}
   if j.get('encoding')=='base64' and 'content' in j:
    import base64;b=base64.b64decode(j['content'])
  local=(ROOT/name).read_bytes();assert len(b)==len(local) and hashlib.sha256(b).digest()==hashlib.sha256(local).digest(),name
  f=dest/name;f.parent.mkdir(parents=True,exist_ok=True);f.write_bytes(b)
  return {'path':path,'bytes':len(b),'sha256':hashlib.sha256(b).hexdigest(),'actual_remote_bytes_verified':True}
 with concurrent.futures.ThreadPoolExecutor(max_workers=2) as pool:readback=list(pool.map(read_one,names))
 # Re-open the independently obtained workbook, JSON and RIS, not local outputs.
 from openpyxl import load_workbook
 wb=load_workbook(dest/'oxygen_response_review_cumulative.xlsx',read_only=True)
 sheets={ws.title:ws.max_row-1 for ws in wb};wb.close();assert sheets['机制证据']==20
 assert len(json.loads((dest/'screening.json').read_text()))==20
 assert (dest/'oxygen_response_B001.ris').read_text().count('TY  - ')==20
 result={'payload_commit':commit,'verified_utc':dt.datetime.now(dt.timezone.utc).isoformat(),'repository':'Colliam-Vermilion/FirstCodex','branch':'review/oxygen-response-mechanisms','remote_tree_not_truncated':True,'all_committed_round_files_git_object_matches':len(objects),'remote_git_objects':objects,'independent_HTTP_byte_readbacks':len(readback),'readbacks':readback,'workbook_reopened_from_remote':sheets,'remote_screening_json_records':20,'remote_B001_RIS_records':20,'verification':'passed','scope_note':'All round files matched to immutable remote Git tree object IDs. The explicitly listed key deliverables were additionally downloaded over HTTP, SHA256 checked and reopened. Local /tmp path is validation workspace, not delivery link; use GitHub branch links.'}
 (ROOT/'remote_readback.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8');print(json.dumps({k:v for k,v in result.items() if k not in ['remote_git_objects','readbacks']},ensure_ascii=False))
if __name__=='__main__':
 ap=argparse.ArgumentParser();ap.add_argument('--commit');a=ap.parse_args();run(a.commit or subprocess.check_output(['git','rev-parse','HEAD'],cwd=REPO,text=True).strip())
