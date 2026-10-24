#!/usr/bin/env python3
"""Lightweight checkpoint validation; never reads unchanged paper originals."""
import collections,hashlib,json,subprocess,tempfile
from pathlib import Path
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parents[1];OUT=ROOT/'all_pool'
def read(p):return json.loads(p.read_text())
def run():
 base=read(OUT/'baseline_catalog.json');cat=read(ROOT/'catalog.json');old=read(OUT/'baseline_screening.json');screen=read(ROOT/'screening.json');dec=[read(p) for p in (OUT/'decisions').glob('*.json')];ids={r['id'] for r in base};oid={r['id'] for r in old};did={r['id'] for r in dec};assert len(cat)==len(ids)==1536;assert len(did)==len(dec)==1516;assert not (oid&did);assert oid|did==ids
 assert screen[:len(old)]==old
 assert hashlib.sha256((ROOT/'reading_priority/evaluations.json').read_bytes()).hexdigest()==read(OUT/'baseline.json')['original20_evaluations_sha256']
 baseline={r['id']:r for r in base}
 for r in cat:
  for k in ['old_library_ids','old_ad_classification','old_human_classification','original_fulltext_paths']:
   assert r.get(k)==baseline[r['id']].get(k),(r['id'],k)
 judgments=[x for p in sorted(list((OUT/'judgments').glob('*.json'))+list(OUT.glob('judgments_P*.json'))) for x in read(p)];assert len(judgments)==1516;db={x['id']:x for x in dec}
 for x in judgments:
  d=db[x['id']]
  for k in ['batch','reason','decision']:
   field={'reason':'screening_reason','decision':'screening_decision'}.get(k,k);assert x[k]==d[field],(x['id'],k)
 checkpoints=[read(p) for p in (OUT/'checkpoints').glob('*.json')];assert len(checkpoints)==31;assert sum(x['new_records'] for x in checkpoints)==1516
 meta=read(OUT/'metadata_enrichment.json');assert len(meta)==1536
 origdois={x['doi'] for x in old if x.get('doi')}
 # Old screening objects do not all carry DOI; resolve from frozen catalog.
 origdois|={baseline[x['id']].get('doi') for x in old};origdois.discard(None);origdois.discard('')
 requested={d for b in read(OUT/'metadata_requests.json') for d in b['dois']};assert not origdois&requested
 for x in meta:
  if x['citation_count'] is None:assert x.get('citation_source') is None
 p=read(OUT/'progress.json');assert p['remaining_semantic_unscreened']==0;assert read(ROOT/'remaining_queue.json')==[]
 expected={'文献总表':1536,'全池筛选':1536,'全池指标':1536,'全池暂定优先':364,'背景追溯队列':336,'待确认队列':479,'重点全文候选':52,'下一阶段新候选':35,'续跑队列':0}
 wb=load_workbook(ROOT/'oxygen_response_review_cumulative.xlsx',read_only=True,data_only=False)
 for sheet,n in expected.items():assert wb[sheet].max_row-1==n,(sheet,wb[sheet].max_row-1,n)
 prefix='results/ad-oxygen-literature-20261006/oxygen_response_review/'
 prior_bytes=subprocess.check_output(['git','show','625dac5eb84d06fb15e5ff7d884f475730f9af06:'+prefix+'oxygen_response_review_cumulative.xlsx'],cwd=ROOT)
 with tempfile.NamedTemporaryFile(suffix='.xlsx') as f:
  f.write(prior_bytes);f.flush();previous=load_workbook(f.name,read_only=True,data_only=False)
  for name in ['机制证据','重点阅读','检索日志','引文追溯','优先精读','机制验证关联','评价口径']:assert list(previous[name].values)==list(wb[name].values),name
  previous.close()
 wb.close()
 for name,n in [('provisional_priority_set.ris',364),('new_focus_set.ris',35),('focus_fulltext_set.ris',52)]:
  s=(OUT/name).read_text();assert s.count('TY  - ')==s.count('ER  -')==s.count('ID  - ')==n
 changed=subprocess.check_output(['git','diff','--name-only','625dac5eb84d06fb15e5ff7d884f475730f9af06'],cwd=ROOT).decode().splitlines();assert all(x.startswith(prefix) or x=='results/ad-oxygen-literature-20261006/README.md' for x in changed),changed
 report={'result':'pass','pool':1536,'new_decisions':1516,'checkpoints':31,'new_status_counts':dict(collections.Counter(r['screening_decision'] for r in dec)),'old20_screen_objects_equal':True,'old20_evaluation_bytes_equal':True,'seven_original_evidence_evaluation_sheets_cells_equal':True,'old_AD_manual_fields_and_original_paths_preserved':True,'old20_OpenAlex_DOIs_requeried':0,'workbook_rows_verified':expected,'changes_confined_to_review_directory_and_project_README':True,'unchanged_paper_originals_read_or_rehashed':False,'RIS_structure_pass':True,'limitations':'轻量结构/保留性核验，不把规则路由当语义审阅，也不验证候选机制因果'}
 (OUT/'validation.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8');print(json.dumps(report,ensure_ascii=False))
if __name__=='__main__':run()
