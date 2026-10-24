#!/usr/bin/env python3
"""Add reading-priority evaluations; preserve discovery/screening/queue snapshots.

No literature search, citation threshold or impact-factor inclusion rule here.
Focal assessments are explicit, reviewed judgments, not a term classifier.
"""
import collections,datetime as dt,hashlib,json,re,sys,unicodedata,xml.etree.ElementTree as ET
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.styles import Alignment,Font,PatternFill
ROOT=Path(__file__).resolve().parents[1];PROJECT=ROOT.parent;OUT=ROOT/'reading_priority';RAW=OUT/'raw'
DATE='2026-10-24';ASOF=dt.date(2026,10,24)
sys.path.insert(0,str(ROOT/'scripts'))
from build_catalog import csvout,dump,normtitle

def txt(e):return ''.join(e.itertext()) if e is not None else ''
def parse_metadata():
 path=RAW/'pubmed_selected_and_support.xml';records={}
 for a in ET.fromstring(path.read_bytes()).findall('PubmedArticle'):
  art=a.find('MedlineCitation/Article');pm=a.findtext('MedlineCitation/PMID');ids={x.get('IdType'):txt(x) for x in a.findall('PubmedData/ArticleIdList/ArticleId')}
  authors=[{'family':x.findtext('LastName',''),'given':x.findtext('ForeName',''),'initials':x.findtext('Initials',''),'affiliations':sorted(set(txt(z) for z in x.findall('.//Affiliation')))} for x in art.findall('AuthorList/Author')]
  sections=[{'locator':f'PubmedArticle[PMID={pm}]/MedlineCitation/Article/Abstract/AbstractText[{i+1}]','label':x.get('Label',''),'text':txt(x)} for i,x in enumerate(art.findall('Abstract/AbstractText'))]
  records[pm]={'pmid':pm,'doi':ids.get('doi','').lower(),'pmcid':ids.get('pmc',''),'title':txt(art.find('ArticleTitle')),'journal':art.findtext('Journal/Title',''),'issn':art.findtext('Journal/ISSN',''),'issn_type':art.find('Journal/ISSN').get('IssnType','') if art.find('Journal/ISSN') is not None else '', 'year':art.findtext('Journal/JournalIssue/PubDate/Year',''),'authors':authors,'abstract_sections':sections,'source_path':str(path.relative_to(ROOT)),'url':f'https://pubmed.ncbi.nlm.nih.gov/{pm}/'}
 # This historical PubMed record lacks a DOI. Crossref is accepted only after
 # exact title, year, pages and first author verification, not query rank.
 j=json.loads((RAW/'crossref_Flint1993.json').read_text())
 matches=[r for r in j['message']['items'] if normtitle(' '.join(r.get('title',[])))==normtitle(records['8226748']['title']) and r.get('page')=='22369-22376' and r.get('author',[{}])[0].get('family')=='Flint' and r.get('published',{}).get('date-parts',[[]])[0][0]==1993]
 assert len(matches)==1
 records['8226748'].update(doi=matches[0]['DOI'].lower(),doi_verification='Crossref exact title + Flint + 1993 + pages22369–22376',doi_source_path=str((RAW/'crossref_Flint1993.json').relative_to(ROOT)))
 return records

def journal_metrics():
 p=PROJECT/'inventory/merge_20261023/jcr_sources';records={}
 for f in sorted(p.glob('jcr_search_*.json')):
  j=json.loads(f.read_text());text='\n'.join(x.get('text','') for x in j.get('content',[]) if x.get('type')=='text')
  for block in text.split('--------------------------------------------------------------------------------'):
   m=re.search('刊名全称[:：]([^\n]+?)\\s+ISSN[:：]([^\n]*)',block);v=re.search(r'影响因子\(2024\)[:：]\s*([0-9.]+)',block)
   if not m or not v:continue
   title=m.group(1).strip();norm=normtitle(title);url=re.search(r'https://topj\.lib\.whu\.edu\.cn/[^)\s]+',block)
   cats=[]
   if 'JCR学科及分区' in block:
    section=block.split('JCR学科及分区',1)[1].split('ESI期刊信息',1)[0]
    cats=[{'category':c.strip(),'quartile':q} for c,q in re.findall(r'→(?:SCIE|SSCI)[:：]\s*([^\n(]+)\((Q[1-4])\)',section)]
   rec={'journal':title,'issn':re.sub('[^0-9Xx]','',m.group(2).split('影响因子')[0].split('投稿指南')[0]),'jif':float(v.group(1)),'jif_metric_year':2024,'jcr_release_year':2025,'categories':cats,'source_url':url.group(0) if url else '', 'source_snapshot':'../'+str(f.relative_to(PROJECT)),'source_snapshot_sha256':hashlib.sha256(f.read_bytes()).hexdigest(),'source_date':'2026-10-23','verification':'复用已保存机构JCR栏目搜索索引快照；非Clarivate账户导出，未声称本轮直连原页成功','evidence_quote':block[max(0,m.start()-10):min(len(block),v.end()+650)]}
   if norm in records:assert records[norm]['jif']==rec['jif'],title
   records[norm]=rec
 aliases={'microbiologyreadingengland':'microbiologysgm','thefebsjournal':'febsjournal','theismejournal':'ismejournal'}
 dump(OUT/'journal_metric_source_register.json',records)
 return records,aliases

# Explicit assessments: PMID, primary theme, AD comparison grade (3=direct
# process/same genus+function; 2=specified shared pathway; 1=limited analogy),
# evidence grade (4=causal genetics/enzyme;3=controlled physiology;2=supportive;
# 1=potential/correlation), primary tier, taxon match, AD basis, mechanistic
# portability, exact claim, primary evidence, boundary and recommendation.
ASSESSMENTS=[
('32301184','M1 损伤',2,4,'基础理论文献','未核实B. thetaiotaomicron为目标AD常见物种；不以同属Bacteroidota作为依据','AD发酵/产乙酸过程同样依赖低电位丙酮酸代谢；可比靶点为PFOR/PFL及ROS敏感FeS酶','具体酶/反应机制可迁移，菌株耐氧阈值不可迁移','分子O2直接使B. thetaiotaomicron PFOR失活，ROS与直接O2损伤可并存','ROS清除缺失/过表达、酶活和代谢通路干预；直接验证特定靶点','肠道培养基/还原物与AD不同；PFOR能否修复存在物种及试验条件限制','优先建立直接O2与ROS的损伤框架；先读后续类群比较'),
('24222618','M2 解毒',3,4,'核心机制文献','Methanosarcina同属与产甲烷功能对应AD；M. acetivorans具体种在目标AD的常见性未核实','原AD研究R07202包含Methanosarcina；比较的是产甲烷生理与抗氧化能力，不是只因同属古菌','甲烷代谢和H2O2/O2区分可比；催化酶系统存在种/菌株差异','EcKatG诱导提高H2O2耐受，但在该条件下不足以提高O2耐受','遗传引入/诱导、酶活、生长/存活；摘要报告的直接干预','海洋来源模式菌/诱导系统不等于污泥原位；只否定单独提高催化酶的充分性，不否定全部作用','优先读阴性O2结果，防止把catalase列为普遍耐氧解释'),
('38129989','M2 解毒',3,4,'核心机制文献','Syntrophomonas与AD互营脂肪酸氧化菌同属/同功能；并非把任意厌氧菌当近缘对象','与R13862的S. zehnderi同属且共同承担脂肪酸互营氧化；不假定两种的防御表达相同','FDP家族氧还原机制可比；具体五结构域FDP_H不可跨物种直接泛化','S. wolfei FDP_H可还原O2/H2O2；rubredoxin铁结合位点突变降低反应活性','重组蛋白、光谱、定点突变、动力学；直接离体酶学','尚未证明S. wolfei完整细胞表达量/耐氧或AD原位通量；离体酶不等于过程促进','优先读AD相关互营菌的明确生化机制，同时补完整细胞证据'),
('25706146','M6 恢复与适应',3,3,'核心机制文献','Methanosarcina同属/产甲烷功能与AD相同；具体海洋种的AD代表性未核实','产甲烷活性与反复微氧暴露直接相关；比较的是同菌首次暴露与长期培养','可迁移的是多防御协同及暴露史，不能照搬百分比O2阈值','反复培养至少6个月形成保持生长/产甲烷的耐氧表型','重复氧处理与对照、产甲烷/蛋白、抗氧化酶与还原物；生理直接、分子解释支持性','培养基硫化物/半胱氨酸和高盐改变暴露；无遗传适应证据；不证明某一酶单独致效','用于区分首次冲击、长期表型驯化与遗传适应'),
('36083837','M6 恢复与适应',2,3,'重要补充与边界文献','产甲烷及互营功能可比；ORM2苯发酵群不是已核实AD常见菌','AD依赖多功能团连续供底物；本体系直接显示伴生需氧菌扩增与慢生长厌氧功能恢复','适用的是氧耗尽与功能恢复分离、竞争/替代，非恢复时间的直接外推','较大空气侵入使产甲烷苯降解长期停滞，去氧后恢复近一年','受控氧扰动、苯/甲烷及绝对丰度时间序列；直接功能响应','苯为唯一底物，特殊慢生长富集；不能将近一年套用于AD或推断遗传适应','优先读负面与长恢复实例，避免只写促进效应'),
('33326896','M4 互作',3,1,'重要补充与边界文献','直接AD反应器；Proteiniphilum、Methanosarcina等对应目标过程','Napier grass微曝气AD，原ID R07202；水解发酵与氢营养产甲烷的具体功能组合','cytochrome bd/SOR/rbr为明确可比候选系统，但原位活性未验证','Proteiniphilum的bd和产甲烷菌ROS防御基因可能协同支撑微曝气AD','宏基因组编码/通路重建，结合既往性能；相关性/功能潜力，不是酶活或因果互作','ORP不能唯一换算picomolar O2；高丰度/编码不证明表达活性或DIET','AD关联强，适合提出待验证模型；不能升级为核心已验证机制'),
('32515155','M4 互作',3,3,'核心机制文献','Syntrophomonas zehnderi、Methanobacterium formicicum与AD脂肪酸互营功能直接相符','原AD库R13862；油酸氧化—产甲烷与兼性Pseudomonas组分干预','耗氧与H2/甲酸/乙酸竞争共同可比；不能把保护机制唯一归因一项','加入Pseudomonas的合成共培养可在微氧下加速油酸转化；保护与代谢清除可能并存','组合培养/对照和底物/产物测量直接；耗氧保护的中介归因仍支持性','连续阶段多项条件改变，存在还原剂缓冲；不证明DIET；2024跟进存在团队重叠','优先读群落干预及代谢物，分开保护、交叉营养与竞争'),
('29725720','M5 空间保护',3,3,'核心机制文献','Methanosarcina同属/乙酸产甲烷功能与AD对应；未以甲烷氧化菌为AD普遍成员','乙酸产甲烷菌与好氧甲烷氧化菌长期共培养，DO/甲烷测量及FISH','氧消耗—聚集—厌氧功能共存可比，同时有甲烷被氧化的代价','限氧反应器中产甲烷菌/甲烷氧化菌稳定共存并紧密聚集，作者提出耗氧保护','共培养生理/化学和FISH直接；保护中介/空间扩散机制未单独因果验证','液相低于检出限非零氧；无充分伙伴去除或微剖面验证；净甲烷可能受消耗','优先读空间互作与产物竞争权衡，不把FISH直接称扩散保护证明'),
('37918493','M5 空间保护',1,2,'暂缓推荐','DAMO/Anammox与厌氧环境有关，但不是AD产甲烷过程；不以同为厌氧菌判高关联','可比问题是生物膜氧扩散；主要代谢为甲烷氧化/氮转化，方向与AD产甲烷不同','保护模型可作线索，但当前摘要不足以核实分层/氧微剖面','生物膜分层保护DAMO类群免受氧冲击','目前摘要有性能/丰度/基因变化；核心空间保护主要为作者解释','尚未确认直接分层成像、氧剖面及干预；仅基因丰度不能证明保护','保留原核心范围筛选结论；精读优先级暂缓，先补空间证据再评'),
('25398869','M5 空间保护',3,3,'核心机制文献','M. harundinacea与UASB颗粒相关，为目标产甲烷丝状类群；关联依据为原研究对象而非上位分类','乙酸产甲烷、长短丝形态与AD颗粒结构形成对应','碳分配/抗氧化活性可比；厚包被阻氧仍需直接验证','长丝相较短丝表现更高氧耐受/催化酶活，并有乙酸碳分配及包被差异','氧耐受、酶活、13C乙酸、TEM/转录；实测支持但包被中介为相关性','纯培养形态与整体颗粒不同；氧挑战剂量/时长未从摘要核实','优先补全文，连接AD优势形态、碳流和耐氧，避免将包被观察直接当屏障机制'),
('37382539','M2 解毒',1,3,'重要补充与边界文献','肠道Faecalibacterium种的AD常见性未核实；不以同门或同科判断高相关','共享FDP/rbr/SOR等候选防御与还原性培养物问题，但不是AD同一菌群','可迁移的是菌株差异及外源半胱氨酸改变氧/ROS暴露，非绝对阈值','不同株O2敏感性及防御谱差异显著，半胱氨酸可降低胞外超氧并改善存活','存活/化学干预直接；调控网络和具体酶贡献仍主要转录/基因支持','基因有无不等于活性；半胱氨酸化学反应取决于体系；未核实ROS探针特异性','用作反对统一耐氧阈值/通用酶叙事的条件性研究'),
('38953644','M2 解毒',2,4,'基础理论文献','C. difficile不是AD目标常见菌；相关性来自具体SOR防御系统，不是同属Clostridia','R07202中AD产甲烷菌编码SOR，但不能据编码即声称活性','SOR的生化反应/抗氧化因果原理可迁移；病原菌剂量与生态不可直接迁移','SOR遗传失活及异源表达/清除活性支持其对超氧防御的作用','遗传失活、异源表达、裂解液活性/蛋白组；直接功能干预，当前摘要','O2与超氧挑战需分开核实；SOR在所有AD菌中的必要性未知','优先补全文作为SOR因果证据，配合2012古菌研究界定跨类群支持'),
('26431975','M3 代谢',2,3,'基础理论文献','C. ljungdahlii为产乙酸/合成气发酵对象；不假定此种在目标AD常见','Wood–Ljungdahl/低电位碳代谢、电子供体与还原力分配与AD产乙酸/发酵过程可比','AOR/rbr及还原力成本可启发AD；气体底物与富集培养条件不同','氧暴露伴乙醇增、乙酸降，AOR/rbr相关响应联系代谢与防御','发酵产物/酶活直接加转录支持；未核实相应基因干预','浓度改变非完整通量；8% O2定义与时长需正文，不能当AD最优剂量','优先补全文，连接解毒电子需求与碳代谢重排'),
('27663027','M3 代谢',2,3,'基础理论文献','A. muciniphila不是AD常见物种；只比较具体cytochrome bd氧利用系统','AD R07202中有bd候选编码；本研究生理/异源互补提供功能层级的参照','可迁移低氧终端氧化酶功能，不可由此断言AD所有严格厌氧菌均以氧获能','低氧下生长/耗氧能力及bd异源互补支持条件性氧利用','生理测量＋异源遗传互补；当前摘要','异源互补不替代本菌敲除；黏液界面/底物与AD不同；纳摩尔阈值需核实','优先与B. fragilis 2004合读低氧利用的跨类群原始证据'),
('33147001','M6 恢复与适应',1,3,'重要补充与边界文献','Anammox代谢不是AD产甲烷/发酵过程；类群相近性不作为主要理由','可比较已定义DO暴露、去氧与还原剂恢复试验设计','活性—转录联合和恢复处理可借鉴，kat/sod功能及铁化学效应不可直接外推','DO冲击后中心代谢转录变化，适量nZVI的恢复表现接近对照','DO/恢复干预及活性直接；stringent response与kat/sod作用仍为推测','铁耗氧/生物效应未分离；转录不证明能量通量；恢复时长待全文','用于恢复试验和证据层级的边界案例，暂定并优先补方法'),
('29312168','M5 空间保护',2,2,'重要补充与边界文献','硫/硫化物转化功能与AD硫循环可比；Arcobacter具体AD出现未核实','厌氧—再氧合界面有氧微剖面及硫化学，与AD氧入口分区问题相关','可迁移梯度/硫循环重组，不可直接搬海水硫酸盐水平和固氮贡献','再氧合改变微生物结构/硫氧化转录，直接测到沉积物氧梯度','氧微剖面与化学直接；类群代谢及固氮功能主要支持性转录','自然沉积物扩散、Fe/S/N反应混杂；网络不证明互作，RNA不证明固氮通量','优先读作为实际微剖面与空间模型证据差异的对照'),
('34409684','M7 剂量与情境',2,3,'重要补充与边界文献','Methanobacterium同属/氢营养产甲烷功能与AD对应，前段泥炭降解菌不等同AD菌','氧化处理改变供给甲烷菌的底物，涉及与AD可比较的H2/CO2产甲烷路径','时间顺序和底物活化可迁移，低pH泥炭及98天氧化阶段与微曝气AD差异显著','先氧化后厌氧的泥炭培养提高后续甲烷产量，伴群落/碳化学转变','阶段干预、甲烷/气体/化学直接；具体跨营养/类Fenton中介主要推测','最高倍数属特定对照；不是持续微氧促甲烷；低pH/Fe及取样层混杂','优先读剂量、阶段与底物决定促进/抑制的条件性证据'),
('38580315','M2 解毒',3,2,'暂缓推荐','AD来源氢营养产甲烷富集，高过程相关性','同为H2/CO2产甲烷及抗氧化蛋白，但氧暴露界定不充分','硒蛋白组为候选机制；不能将ORP视为O2剂量','高ORP/无AQDS组的硒相关蛋白变化与氧化应激防御有关','蛋白组/甲烷支持性；O2侵入间接推断且AQDS改变电子流','DO未直接测量；硒路径无因果干预，不能按高AD关联升级核心','保留待确认，先解决氧处理与硒路径归因再进入正式精读集'),
('40057977','M3 代谢',1,2,'重要补充与边界文献','海岸Chloroflexota/Patescibacteria具体种与AD对应未核实；不以同门判断高度相关','硫酸盐还原/发酵及活性代理—转录解耦可作为AD证据解释的方法学参照','可迁移多层证据不互替的原则，不把海水稀释培养的速率直接外推AD','氧条件下呼吸相关信号与转录方向不一致','RSG/单细胞、基因组、转录及化学；活性是代理与估算，非全量直接通量','1:100稀释；RNA样品受laminarin及数量限制；氧单位表述有歧义','保留新近低引用研究，用于检索结果到功能结论的证据边界'),
('39608160','M4 互作',3,1,'暂缓推荐','直接AD底物/工艺对象，但不等于直接氧响应机制','秸秆/真菌/生物炭与AD过程相关；原AD排除及本轮待确认均保留','尚不能把真菌/材料效应归因O2，亦不证明DIET','初始残留氧下真菌/生物炭改善转化；O2中介/DIET主张未建立','干预真菌/材料；缺已确认独立氧对照，机制归因不充分','电容/电活性菌富集不是DIET因果证据；氧消耗/无氧对照需核实','保持待确认，不因Water Research指标或引用量纳入正式精读集'),
]

def author_key(a):return normtitle(a['family'])+':'+a['initials'].upper()
def quote_for(meta,needle):
 for p in meta['abstract_sections']:
  t=' '.join(p['text'].split());m=re.search(re.escape(needle),t,re.I)
  if m:return {'quote':t[max(0,m.start()-25):min(len(t),m.end()+160)],'locator':p['locator'],'source_path':meta['source_path'],'source_url':meta['url'],'reading_level':'摘要核验；非全文'}
 raise ValueError((meta['pmid'],needle))

def claim_links(meta,cat):
 # Each relation describes a SPECIFIC claim; a similar phenomenon is not
 # automatically an independently replicated mediator or universal rule.
 specs=[
 ('CL01','32301184','B. thetaiotaomicron PFOR可被分子O2直接损伤','11260473','同团队早期支持','分子O2与超氧靶点区分','2001与2020共同Imlay/Illinois，不能计独立实验室','2001摘要与2020对PFOR恢复性的表述差异需另查','damaged in vitro by molecular oxygen'),
 ('CL01-repair','32301184','PFOR受氧损伤后是否可恢复需按酶/条件核实','11260473','同团队早晚恢复性表述差异（待核实）','2001摘要称相关酶在去氧后迅速修复，2020全文特定试验PFOR不能恢复','共同Imlay/Illinois，不能把时间间隔当团队独立','2001全文方法未核实，不能凭摘要判哪篇错误；须区分酶种类、剂量及修复试验','they were then rapidly repaired'),
 ('CL01','32301184','不能把所有PFOR概括为同样易损/不可恢复','18161989','不同团队反例/适用条件','Desulfovibrio PFOR二硫键保护与再活化','CNRS Marseille与Illinois不同已列团队；对象/同工酶不同','限制通用化，不直接反驳Bacteroides特定条件结果','disulfide bond-dependent protective mechanism'),
 ('CL02','32301184','O2暴露后内源超氧可损伤FeS代谢酶','29559534','同团队后续/互补机制','2018研究ROS支路；2020加直接O2支路','Imlay/Illinois重叠；不是不同实验室验证','不同酶靶点，不能拿ROS结果替代PFOR的直接O2结果','overproduction of superoxide dismutase substantially protected'),
 ('CL02','32301184','某些[4Fe-4S]水合裂解酶易受超氧破坏','8226748','不同团队跨对象生化支持（暂定）','纯化酶/簇铁释放，非直接厌氧背景完整细胞复现','Flint/du Pont与Imlay/Illinois不同已列团队','只有摘要；含非厌氧/哺乳动物酶，不验证全部PFOR或AD','release of iron from their clusters'),
 ('CL03','24222618','提高catalase活性不必然提高O2耐受','25706146','不同团队条件性比较，非矛盾/复制','长期驯化有多酶/还原物共同变化','Lessner团队与Jasso-Chávez/Moreno-Sánchez团队不同；培养方式不同','相关升高不证明catalase单独充分；不是互相否定','transcripts and activities of superoxide dismutase'),
 ('CL04','38129989','五结构域S. wolfei FDP_H还原O2及结构域作用','32900801','共享作者团队的跨类群支持','C. difficile FDP/逆rbr的遗传＋酶学','共享Martins MC/Folgosa F及ITQB，不计完全独立团队','酶家族/结构域不同，不能复制FDP_H特定结构域结论','growth of a fdpA mutant is affected'),
 ('CL04','38129989','FDP家族可参与厌氧微生物耐氧','23093671','不同团队跨类群原始支持（暂定）','P. furiosus FDP_A缺失增加O2敏感性','Adams/Georgia与ITQB不同已列团队','摘要核验；putative FDP_A，不同类别；不能替代S. wolfei细胞内验证','putative flavodiiron protein A showed increased sensitivity'),
 ('CL05','38953644','SOR可参与O2/超氧防御','23093671','不同团队跨类群原始支持（暂定）','P. furiosus SOR缺失提高O2敏感性','Vedantam等与Adams等作者/机构不同','当前两方相关主张以摘要核实，O2与超氧挑战不同；不验证AD古菌全部必要性','Deletion strains lacking superoxide reductase'),
 ('CL06','32515155','兼性菌帮助互营脂肪酸群落在微氧下转化','39547298','共享团队后续支持','0–10%O2、两组合成互营培养，±Pseudomonas','共享Duarte MS、Pereira、Stams及Minho团队；非独立实验室','2024后续仅摘要；不能当完全独立再现，不证明DIET','Without Pseudomonas, the syntrophic activity was inhibited'),
 ('CL06','32515155','耗氧伙伴与厌氧功能共存的原则','29725720','不同团队/体系支持性观察，非相同机制复制','产甲烷菌—甲烷氧化菌限氧共培养','Minho/Wageningen与Welte等团队不同已列人员；未系统核共享资源','底物/伙伴不同；FISH共存不单独证明耗氧保护中介','methanogens are protected by the concomitant oxygen consumption'),
 ('CL07','29725720','聚集/耗氧可能保护产甲烷菌','32515155','不同团队/体系支持性观察，非相同机制复制','兼性Pseudomonas与互营甲烷培养','作者/研究组合不同；保护中介未完全分离','没有严格相同对象的独立复现；不能从聚集推出扩散机制','protecting strict anaerobes from oxygen toxicity'),
 ('CL08','27663027','某些厌氧菌可通过bd氧化酶利用低水平O2','14749831','不同团队跨类群原始支持（暂定）','B. fragilis bd对耗氧/特定条件促生长必要','Baughn/Malamy Tufts与Ouwerkerk/Belzer/Wageningen不同已列团队','2004基础研究仅摘要核实；非所有厌氧菌均利用O2获能；非AD原位证明','cytochrome bd oxidase that is essential for O2 consumption'),
 ('CL09','25706146','长期耐氧培养不等于单一catalase因果','24222618','不同团队阴性结果限定机制归因','提高催化酶不增加O2耐受','不同已列团队，且一个短期诱导一个多月培养','不是对长期表型耐受的否定，更未验证遗传适应','did not exhibit increased tolerance to O2'),
 ('CL10','33326896','AD中bd/SOR/rbr编码不能直接等同活性','38129989','跨对象直接酶学参照，非AD独立验证','互营菌FDP_H可实际还原氧','本研究对象/酶系统不同，不能为本条bd/SOR基因赋予活性','研究只提供证据层级参照，不算原AD机制独立验证','O2 reduction activity'),
 ('CL11','34409684','氧对甲烷的净效应受阶段/底物/剂量限制','36083837','不同体系的负向边界，非矛盾/复制','苯富集较大冲击后长期抑制','已列作者团队不同；底物和处理历程不同','两体系并不构成互相矛盾，不能比较原始效应倍数','subsequent 100 times larger injection of air'),
 ('CL12','26431975','氧解毒可与还原力/电子流重新分配相联系','23093671','不同团队跨体系支持（暂定），非AOR/rbr特定机制验证','P. furiosus O2消耗伴H2产量下降，基因删除支持解毒耗电子','Grunden等与Adams等不同已列团队；非同代谢通路复制','仅摘要；发酵底物/温度不同，不能验证C. ljungdahlii的AOR或乙醇路径','for each mole of O(2) consumed'),
 ]
 # Add full-text confirmation for the shared-team C. difficile support; the
 # abstract quote remains traceable and is not relabeled as a body passage.
 rows=[]
 pmap={r.get('pmid'):r for r in cat.values() if r.get('pmid')}
 for claim,focal,claimtext,related,relation,scope,teamnote,limit,needle in specs:
  a=meta[focal];b=meta[related];overlap=sorted({author_key(z) for z in a['authors']} & {author_key(z) for z in b['authors']})
  q=quote_for(b,needle)
  rows.append({'claim_id':claim,'focal_id':pmap[focal]['id'],'focal_doi':a['doi'],'specific_claim':claimtext,'related_id':pmap.get(related,{}).get('id','SUP-PMID'+related),'related_title':b['title'],'related_year':b['year'],'related_doi':b['doi'],'related_pmid':related,'relation':relation,'what_is_supported':scope,'shared_listed_authors':overlap,'team_assessment':teamnote,'independence_caveat':'作者/机构差异不证明无共享材料/人员历史；未系统排查实验资源共享','exact_same_object_independent_replication':'未核实','limitations':limit,**q})
 linked={r['focal_id'] for r in rows}
 for r in json.loads((ROOT/'screening.json').read_text()):
  if r['id'] not in linked:
   rows.append({'claim_id':'UNVERIFIED-'+r['id'],'focal_id':r['id'],'focal_doi':cat[r['id']]['doi'],'specific_claim':r['evidence']['major_response'],'related_id':'','related_title':'','related_year':'','related_doi':'','related_pmid':'','relation':'未核实独立验证','what_is_supported':'本轮未找到可确认的同机制独立验证链；不表示不存在','shared_listed_authors':[],'team_assessment':'未核实','independence_caveat':'未用引用次数或综述重复引用代替独立原始证据','exact_same_object_independent_replication':'未核实','limitations':r['uncertainties'],'quote':'','locator':'','source_path':'','source_url':'','reading_level':'未核实'})
 for r in rows:
  rel=r['relation']
  r['validation_relation_code']=('same_team_follow_up' if (rel.startswith('同团队') or '共享' in rel) else 'independent_cross_taxon_provisional' if '不同团队跨类群' in rel else 'counterexample_or_condition' if ('反例' in rel or '限定' in rel or '负向边界' in rel) else 'cross_system_support_not_replication' if ('不同团队' in rel or '跨体系' in rel) else 'unverified_or_context_only')
  r['strict_same_object_independent_replication_confirmed']=False
  r['additional_fulltext_reading']='相关Results/Methods已读，仍不表示S. wolfei FDP_H的严格独立复制；本行短引保留其实际摘要位置' if r.get('related_pmid')=='32900801' else ''
  r['additional_fulltext_excerpt_file']='reading_priority/support_fulltext_excerpts.json' if r.get('related_pmid')=='32900801' else ''
 return rows

def evaluation_policy():
 return [
 {'dimension':'范围与优先级分离','rule':'保留B001全部20条既有筛选结论；精读暂缓不等于范围排除。1516未筛队列不改变；未评估文献不记低优先级。'},
 {'dimension':'AD相关性','rule':'分别记录同种/同属/功能与特定通路/靶点/防御系统；不凭同门同科或都厌氧认定高度相关。A3：直接AD或有明确同属＋同功能；A2：具体共同过程/系统；A1：有限方法/生态类比。'},
 {'dimension':'机制直接性','rule':'D4：遗传/直接酶学；D3：受控生理、共培养、同位素等；D2：支持性测量/组学与解释混合；D1：功能潜力/相关性或缺氧处理对照。一文内各主张仍单独注明，整体等级不把所有推测升级。'},
 {'dimension':'独立验证','rule':'以具体主张及原始DOI为单位。共同作者/机构须提示；不同团队跨类群支持不等于同对象同条件独立复制。仅摘要的支持标暂定；没有证据为未核实，不作无验证断言。'},
 {'dimension':'引用数据','rule':'统一OpenAlex、2026-10-24上海日期、按DOI匹配。年均=引用数/从OpenAlex publication_date至查询日的天数×365.2425；数据库日期可能为卷期/默认日期，非保证首个在线日。缺失留空，不记0。'},
 {'dimension':'新文献与指标作用','rule':'年均引用只辅助比较，年轻文章年均波动大；引用/JIF均不作为准入、降级或排除阈值，不计算可抵消科学缺陷的加权总分。'},
 {'dimension':'期刊指标','rule':'统一JIF2024指标年/JCR2025发布版；只复用有明确年份和期刊身份的既有机构JCR栏索引快照。逐学科保留JCR分区，不混CAS/SJR；缺失留空。二级快照非Clarivate账户导出。'},
 {'dimension':'核心机制文献','rule':'AD相关性高且有直接因果/生理干预，或确有较强独立证据；本轮以A3且D≥3为基本必要条件，另核具体主张。仅摘要标暂定。'},
 {'dimension':'基础理论文献','rule':'对象与AD不同，但特定靶点/防御/代谢机制有直接或较强支持并可明确迁移；不因非AD或较低引用丢弃。'},
 {'dimension':'重要补充与边界文献','rule':'纳入有意义的阴性、条件性、恢复、新近/证据解释研究；保留限制，不将支持性证据改写成因果。多标签允许。'},
 {'dimension':'暂缓与全文队列','rule':'氧处理/机制归因待确认者不进入正式精读RIS；空间主张仅摘要且缺直接证据的DAMO研究暂缓。所有摘要推荐为暂定并列优先补全文；网页挑战HTTP200不计成功正文。'},
 {'dimension':'主题与顺序','rule':'按M1–M7组织；主题内先核心/基础/边界，再AD比较、证据直接性、验证关联级别。V2为不同团队跨类群的暂定支持，V1为同团队/支持性关联，V0未核实；V2不是严格复制。引用/JIF供同层级辅助比较，不对不同科学主张机械按指标强排，不决定层级。'},
 {'dimension':'历史与续跑','rule':'B001旧重点阅读粗排不覆盖本次五维评价。原六工作表逐格保留，新增优先精读/机制验证关联/评价口径。后续从已筛选但尚未评价的ID增量更新，不重做已评估或重检索。'},
 ]

def attach_sheets(root=ROOT):
 out=root/'reading_priority';workbook=root/'oxygen_response_review_cumulative.xlsx'
 if not (out/'evaluations.json').exists():return
 wb=load_workbook(workbook)
 for name,filename in [('优先精读','evaluations.json'),('机制验证关联','mechanism_validation_links.json'),('评价口径','evaluation_policy.json')]:
  if name in wb:del wb[name]
  ws=wb.create_sheet(name);rows=json.loads((out/filename).read_text());fields=list(dict.fromkeys(k for r in rows for k in r));ws.append(fields)
  for r in rows:
   vs=[]
   for k in fields:
    v=r.get(k);v=json.dumps(v,ensure_ascii=False) if isinstance(v,(dict,list)) else v
    if isinstance(v,str) and len(v)>32767:v=v[:32600]+' [full value retained in JSON]'
    if isinstance(v,str) and v.startswith(('=','+','@')):v="'"+v
    vs.append(v)
   ws.append(vs)
  ws.freeze_panes='C2';ws.auto_filter.ref=ws.dimensions
  for c in ws[1]:c.fill=PatternFill('solid',fgColor='244766');c.font=Font(bold=True,color='FFFFFF')
  for col in ws.columns:ws.column_dimensions[col[0].column_letter].width=50 if any(x in str(col[0].value) for x in ['依据','主张','理由','边界','验证','证据','rule','claim','limitations','assessment']) else 25
  for row in ws.iter_rows(min_row=2):
   for c in row:c.alignment=Alignment(wrap_text=True,vertical='top')
 wb.save(workbook)

def run():
 meta=parse_metadata();cat={r['id']:r for r in json.loads((ROOT/'catalog.json').read_text())};by_pm={r.get('pmid'):r for r in cat.values() if r.get('pmid')};screen={r['pmid']:r for r in json.loads((ROOT/'screening.json').read_text())}
 journal,aliases=journal_metrics();links=claim_links(meta,cat);evals=[]
 oa=json.loads((RAW/'openalex_citations_20.json').read_text()) if (RAW/'openalex_citations_20.json').exists() else {'results':[]}
 citations={(r.get('doi') or '').lower().split('doi.org/')[-1]:r for r in oa['results']}
 event=next((json.loads(x) for x in (OUT/'request_log.jsonl').read_text().splitlines() if json.loads(x)['name']=='openalex_citations_20'),{})
 for row in ASSESSMENTS:
  pm,theme,adgrade,evgrade,tier,taxon,ad,portable,claim,ev,boundary,why=row;c=by_pm[pm];s=screen[pm];m=meta[pm]
  recommendation=tier!='暂缓推荐';tentative=recommendation and s['reading_level']=='仅题名/摘要';tags=[tier] if recommendation else []
  if pm in ['24222618','25706146','37382539'] and '重要补充与边界文献' not in tags:tags.append('重要补充与边界文献')
  metric=journal.get(aliases.get(normtitle(c['journal']),normtitle(c['journal'])))
  # When a snapshot has an ISSN, a title alias must not silently cross journals.
  if metric and metric['issn'] and m['issn'] and re.sub('[^0-9Xx]','',m['issn'])!=metric['issn']:
   metric={**metric,'identity_note':'精确规范刊名匹配；PubMed ISSN类型为'+m['issn_type']+'，与快照所列ISSN不同；未由ISSN差异改换期刊，未作模糊题名匹配'}
  o=citations.get(c['doi']);count=o.get('cited_by_count') if o else None;pubdate=o.get('publication_date') if o else None;age=None;annual=None
  if count is not None and pubdate:
   days=(ASOF-dt.date.fromisoformat(pubdate)).days
   if days>0:age=round(days/365.2425,4);annual=round(count/(days/365.2425),3)
  ls=[x for x in links if x['focal_id']==c['id']]
  src=json.loads((ROOT/'evidence_sources'/f'{c["id"]}.json').read_text())
  d={'文献ID':c['id'],'题名':c['title'],'年份':c['year'],'DOI':c['doi'],'期刊':c['journal'],'机制主题':theme,'正式精读推荐':recommendation,'推荐主层级':tier,'推荐标签':tags,'推荐确认状态':'暂定：仅摘要，优先补全文' if tentative else ('已核实相关全文段落；非所有机制已因果验证' if recommendation else '补证后再评，不改变原范围结论'),'研究微生物':s['research_system'],'AD对象对应层级及依据':taxon,'AD过程关联依据':ad,'可比较的通路靶点防御':portable,'AD比较等级':adgrade,'机制直接性等级':evgrade,'具体机制主张':claim,'证据类型及直接性':ev,'独立验证状态':'；'.join(dict.fromkeys(x['relation'] for x in ls)),'独立支持或反例文献DOI':[x['related_doi'] for x in ls if x.get('related_doi')],'同对象严格独立重复':'未核实；不由引用或综述推断','验证关联ID':[x['claim_id'] for x in ls],'原筛选状态':s['screening_stage'],'原范围结论':s['screening_decision'],'全文核实状态':s['reading_level'],'推荐理由':why,'AD外推边界':boundary,'原始短引':s['evidence']['quote'],'原文定位':s['evidence']['quote_location'],'原件或摘要来源':s['source_file'],'原件来源SHA256':s['source_sha256'],'来源链接':s['official_url'],'JIF2024':metric['jif'] if metric else None,'JIF指标年份':2024,'JCR发布年份':2025,'JCR逐学科分区':metric['categories'] if metric else [],'期刊指标状态':metric['verification'] if metric else '未取得可核实2024值，留空；不以印象或其他年份/体系代填','期刊指标来源':metric['source_url'] if metric else '', '期刊指标快照':metric['source_snapshot'] if metric else '', '期刊指标快照日期':metric['source_date'] if metric else '', '引用次数':count,'引用数据来源':'OpenAlex' if o else 'OpenAlex未核实，留空非0','引用口径':'cited_by_count数据库总计；未剔除自引，未区分支持或批评的引用','引用查询日期_上海':DATE if o else '', '引用查询时间_上海':event.get('queried_Asia_Shanghai') if o else '', 'OpenAlexID':o.get('id') if o else '', '引用年限起点_数据库日期':pubdate,'数据库日期计算年限':age,'年均引用':annual,'年限校正限制':'OpenAlex日期不保证首个在线日，尤其1月1日可能为年份默认；新研究年均不稳定；仅辅助，不用于降级/准入','评价日期_上海':DATE,'后续补证任务':s['uncertainties'],'下一步动作':'优先获取/核实全文方法与结果' if tentative or not recommendation else '精读并按具体主张追查验证/反例；不重复当前范围筛选'}
  d['期刊身份匹配备注']=metric.get('identity_note','精确规范刊名匹配') if metric else '目标2024指标未核实'
  d['同团队支持原始研究DOI']=[x['related_doi'] for x in ls if x.get('related_doi') and x['validation_relation_code']=='same_team_follow_up']
  d['独立支持或反例文献DOI']=[x['related_doi'] for x in ls if x.get('related_doi') and x['validation_relation_code'] not in ['same_team_follow_up','unverified_or_context_only']]
  d['其他证据参照DOI']=[x['related_doi'] for x in ls if x.get('related_doi') and x['validation_relation_code']=='unverified_or_context_only']
  d['验证关联排序级别']=2 if any(x['validation_relation_code']=='independent_cross_taxon_provisional' for x in ls) else 1 if any(x.get('related_doi') for x in ls) else 0
  d['验证级别限制']='本轮无已核实的同对象同条件严格独立复制；V2也只是不同团队跨类群暂定支持，不等于广泛验证'
  evals.append(d)
 tier_order={'核心机制文献':0,'基础理论文献':1,'重要补充与边界文献':2,'暂缓推荐':3}
 evals.sort(key=lambda r:(r['机制主题'],tier_order[r['推荐主层级']],-r['AD比较等级'],-r['机制直接性等级'],-r['验证关联排序级别'],r['文献ID']))
 for i,r in enumerate(evals,1):r['主题组织序号']=i
 dump(OUT/'evaluations.json',evals);csvout(OUT/'priority_reading_utf8.csv',evals)
 dump(OUT/'mechanism_validation_links.json',links);csvout(OUT/'mechanism_validation_links_utf8.csv',links)
 policy=evaluation_policy();dump(OUT/'evaluation_policy.json',policy);csvout(OUT/'evaluation_policy_utf8.csv',policy)
 selected=[r for r in evals if r['正式精读推荐']];ris=[]
 for r in selected:
  c=cat[r['文献ID']];ris+=['TY  - JOUR','ID  - '+c['id'],'TI  - '+c['title']]
  for a in c.get('authors',[]):ris.append('AU  - '+a)
  for tag,val in [('PY',c['year']),('JO',c['journal']),('DO',c['doi']),('AN',c.get('pmid')),('UR',r['来源链接']),('AB',c.get('abstract'))]:
   if val:ris.append(tag+'  - '+str(val).replace('\n',' '))
  for label in r['推荐标签']:ris.append('KW  - '+label)
  ris+=['KW  - '+r['机制主题'],'N1  - '+r['推荐确认状态']+'; '+r['具体机制主张']+'; '+r['AD外推边界'],'ER  - ','']
 (OUT/'priority_reading_set.ris').write_text('\n'.join(ris),encoding='utf-8')
 fullqueue=[{'id':r['文献ID'],'title':r['题名'],'doi':r['DOI'],'recommended':r['正式精读推荐'],'primary_tier':r['推荐主层级'],'status':r['推荐确认状态'],'needs_fulltext':r['全文核实状态']=='仅题名/摘要','reason':r['后续补证任务'],'action':'优先补全文方法/结果' if r['全文核实状态']=='仅题名/摘要' else '全文已读相关段落，补氧处理/因果证据，不盲目重复下载','source_candidate':r['来源链接']} for r in evals if r['全文核实状态']=='仅题名/摘要' or not r['正式精读推荐']]
 dump(OUT/'fulltext_confirmation_queue.json',fullqueue);csvout(OUT/'fulltext_confirmation_queue_utf8.csv',fullqueue)
 support_pm=['23093671','32900801','39547298','14749831','11260473','18161989','29559534','8226748']
 support=[]
 for pm in support_pm:
  m=meta[pm];support.append({**m,'existing_catalog_id':by_pm.get(pm,{}).get('id'),'role':'仅具体机制关联核验；不计新增正式候选筛选或精读集','early_foundational':bool(m['year'] and int(m['year'])<2012),'fulltext_verified':pm=='32900801','priority_primary_set':False})
 dump(OUT/'support_studies.json',support)
 attach_sheets()
 summary={'evaluation_date_Asia_Shanghai':DATE,'baseline_commit':'8a64a11ed0b79175eb05254ed51d9be898801f1a','evaluated_screened_papers':len(evals),'formally_recommended':len(selected),'held_for_confirmation':len(evals)-len(selected),'primary_tier_counts':dict(collections.Counter(r['推荐主层级'] for r in evals)),'tentative_abstract_only_recommendations':sum(r['推荐确认状态'].startswith('暂定') for r in selected),'fulltext_based_recommendations':sum(r['全文核实状态']!='仅题名/摘要' for r in selected),'citation_counts_verified':sum(r['引用次数'] is not None for r in evals),'jif2024_available':sum(r['JIF2024'] is not None for r in evals),'jcr_quartiles_available':sum(bool(r['JCR逐学科分区']) for r in evals),'claim_relation_rows':len(links),'support_studies_metadata_checked':len(support),'strict_same_object_independent_replications_confirmed':0,'strict_replication_count_note':'0 means none confirmed in this limited audit, not absence from the literature','new_primary_scope_screenings':0,'new_broad_search_queries':0,'original_remaining_screening_queue':1516,'newly_screened_but_not_priority_evaluated':[],'next_update':'新增已筛文献后先按ID检查本表；已评价不重做。按主张补全文/验证，原候选筛选顺序不改变。'}
 summary['newly_screened_but_not_priority_evaluated']=sorted(set(s['id'] for s in screen.values())-{r['文献ID'] for r in evals})
 dump(OUT/'progress.json',summary);print(json.dumps(summary,ensure_ascii=False))
if __name__=='__main__':run()
