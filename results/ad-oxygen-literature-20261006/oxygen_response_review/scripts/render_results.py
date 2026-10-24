#!/usr/bin/env python3
"""Regenerate the SAME cumulative workbook and exports from saved decisions."""
import collections,hashlib,json,re,sys
from pathlib import Path
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Font,PatternFill,Alignment
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT/'scripts'))
from build_catalog import csvout,dump

def hint(r):
 t=(r['title']+' '+r.get('abstract','')).lower()
 groups=[('M1_damage',r'iron.sulfur|inactivat|damage|radical'),('M2_detox',r'rubrerythrin|superoxide|flavodiiron|detox|catalase'),('M3_metabolism',r'metabolic|electron flow|ferment'),('M4_interaction',r'syntroph|cross.feeding|co.cultur|scaveng'),('M5_spatial',r'gradient|stratif|biofilm|granul|filament'),('M6_recovery',r'recover|adapt|shock|transient'),('M7_conditions',r'redox|exposure|oxygen')]
 return next((m for m,pat in groups if re.search(pat,t)),'unassigned')

def run():
 decisions=[json.loads(p.read_text()) for p in sorted((ROOT/'decisions').glob('*.json'))] if (ROOT/'decisions').exists() else []
 decisions.sort(key=lambda r:r['batch_order']);dump(ROOT/'screening.json',decisions)
 cat=json.loads((ROOT/'catalog.json').read_text()); index={r['id']:r for r in cat}; evidence=[];priority=[]
 for c in cat:
  if not c.get('document_type'):c['document_type']=c.get('original_document_type') or '; '.join(c.get('publication_types',[])) or '来源未明确，待核实'
 for r in decisions:
  c=index[r['id']]
  for k in ['screening_stage','screening_decision','screening_reason','research_system','exposure_scenarios','reading_level','document_type','uncertainties']:
   if k in r:c[k]=r[k]
  c['processed_batch']='B001';c['processed_utc']=r['processed_utc'];c['evidence_ids']=[r['id']+'.E01']
  e=dict(r['evidence']);e.update(evidence_id=r['id']+'.E01',id=r['id'],title=c['title'],doi=c.get('doi',''),reading_level=r['reading_level'],screening_stage=r['screening_stage'],decision=r['screening_decision'],source_file=r['source_file'],source_sha256=r['source_sha256'],official_url=r['official_url'])
  evidence.append(e)
  priority.append({'priority_score':r['priority_score'],'id':r['id'],'title':c['title'],'mechanisms':e['mechanism_categories'],'reason':r['priority_reason'],'decision':r['screening_decision'],'reading_level':r['reading_level'],'next_reading_task':r['uncertainties'],'doi':c.get('doi',''),'url':r['official_url']})
 # Freeze a complete order ONCE. Subsequent batches filter completed IDs only.
 order_file=ROOT/'processing_order.json'
 if not order_file.exists():
  chosen=[r['id'] for r in json.loads((ROOT/'batch_B001_selection.json').read_text())]
  buckets=collections.defaultdict(list)
  for r in cat:
   if r['id'] not in chosen:buckets[hint(r)].append(r)
  def score(r):
   t=r['title'].lower();return sum(5 for w in ['oxygen exposure','oxygen shock','aerotoler','oxygen stress','oxygen tolerance','transient oxygen','oxygen reduction'] if w in t)+int(bool(r.get('abstract')))+2*int(any(q!='OLD554' for q in r['queries']))
  for b in buckets.values():b.sort(key=lambda r:(-score(r),r.get('year',''),r['id']))
  ids=list(chosen)
  while any(buckets.values()):
   for k in sorted(buckets):
    if buckets[k]:ids.append(buckets[k].pop(0)['id'])
  dump(order_file,[{'order':i+1,'id':rid,'mechanism_hint_not_a_decision':hint(index[rid])} for i,rid in enumerate(ids)])
 # Hit-driven expansions can add records after a checkpoint. Append; never
 # reorder the already frozen queue or assign a new ID to processed articles.
 order=json.loads(order_file.read_text());known={r['id'] for r in order}
 for r in sorted(cat,key=lambda r:r['id']):
  if r['id'] not in known:order.append({'order':len(order)+1,'id':r['id'],'mechanism_hint_not_a_decision':hint(r),'appended_after_hit_expansion':True})
 dump(order_file,order)
 done={r['id'] for r in decisions};queue=[]
 for o in json.loads(order_file.read_text()):
  r=index.get(o['id'])
  if r and r['id'] not in done:queue.append({**o,'title':r['title'],'year':r.get('year',''),'doi':r.get('doi',''),'pmid':r.get('pmid',''),'old_library_ids':r['old_library_ids'],'old_ad_classification':r.get('old_ad_classification',''),'new_scope_stage':'未开始','abstract_available':bool(r.get('abstract')),'source_queries':r['queries']})
 dump(ROOT/'catalog.json',cat);dump(ROOT/'mechanism_evidence.json',evidence);dump(ROOT/'priority_reading.json',sorted(priority,key=lambda r:(-r['priority_score'],r['id'])));dump(ROOT/'remaining_queue.json',queue)
 csvout(ROOT/'batch_B001_results_utf8.csv',[{**{k:v for k,v in r.items() if k!='evidence'},'title':index[r['id']]['title'],'doi':index[r['id']].get('doi','')} for r in decisions])
 columns=['id','title','year','journal','doi','pmid','authors','sources','queries','research_system','document_type','old_library_ids','old_ad_classification','old_human_classification','screening_stage','screening_decision','screening_reason','reading_level','exposure_scenarios','early_key_study','uncertainties','abstract','url','original_fulltext_paths','provenance']
 csvout(ROOT/'literature_master_utf8.csv',cat,columns);csvout(ROOT/'mechanism_evidence_utf8.csv',evidence);csvout(ROOT/'priority_reading_utf8.csv',sorted(priority,key=lambda r:(-r['priority_score'],r['id'])));csvout(ROOT/'remaining_queue_utf8.csv',queue)
 logs=json.loads((ROOT/'search_log.json').read_text());csvout(ROOT/'search_log_utf8.csv',logs)
 traces=json.loads((ROOT/'citation_trace.json').read_text());csvout(ROOT/'citation_trace_utf8.csv',traces)
 # Full catalog is an unscreened pool. Use B001 RIS for this batch; do not import both.
 for filename,records in [('oxygen_response_all_candidates.ris',cat),('oxygen_response_B001.ris',[index[r['id']] for r in decisions])]:
  ris=[]
  for r in records:
   types=' '.join(r.get('publication_types',[])).lower()
   oldtype=r.get('original_document_type','')
   ty='THES' if ('dissertation' in types or '学位' in oldtype) else 'RPRT' if ('报告' in oldtype) else 'CHAP' if 'book-chapter' in types else 'GEN' if any(t in types for t in ['peer-review','posted-content','proceedings','dataset']) else 'JOUR'
   ris+=['TY  - '+ty,'ID  - '+r['id'],'TI  - '+r['title']]
   for a in r.get('authors',[]):
    if a:ris.append('AU  - '+a)
   for k,v in [('PY',r.get('year')),('JO',r.get('journal')),('DO',r.get('doi')),('AN',r.get('pmid')),('UR',r.get('url')),('AB',r.get('abstract'))]:
    if v:ris.append(k+'  - '+str(v).replace('\n',' '))
   ris+=['N1  - '+r.get('screening_stage','未开始')+'; '+r.get('screening_decision','未判断')+'; old IDs='+','.join(r['old_library_ids'])+'; independent oxygen-response scope','ER  - ','']
  (ROOT/filename).write_text('\n'.join(ris),encoding='utf-8')
 wb=Workbook();wb.remove(wb.active)
 for name,data,fields in [('文献总表',cat,columns),('机制证据',evidence,None),('重点阅读',sorted(priority,key=lambda r:(-r['priority_score'],r['id'])),None),('检索日志',logs,None),('引文追溯',traces,None),('续跑队列',queue,None)]:
  ws=wb.create_sheet(name);fields=fields or list(dict.fromkeys(k for r in data for k in r))
  if not fields:fields=['status']
  ws.append(fields)
  for r in data:
   vals=[]
   for k in fields:
    v=r.get(k,'');v=json.dumps(v,ensure_ascii=False) if isinstance(v,(list,dict)) else v
    if isinstance(v,str) and len(v)>32767:v=v[:32600]+' [Excel limit; complete value retained in JSON/CSV]'
    if isinstance(v,str) and v.startswith(('=','+','@')):v="'"+v
    vals.append(v)
   ws.append(vals)
  ws.freeze_panes='C2';ws.auto_filter.ref=ws.dimensions
  for cell in ws[1]:cell.font=Font(color='FFFFFF',bold=True);cell.fill=PatternFill('solid',fgColor='244766')
  for col in ws.columns:
   ws.column_dimensions[col[0].column_letter].width=48 if any(t in str(col[0].value).lower() for t in ['title','reason','response','limitation','quote','abstract','uncertainties']) else 24
  for row in ws.iter_rows(min_row=2):
   for cell in row:cell.alignment=Alignment(vertical='top',wrap_text=True)
 wb.save(ROOT/'oxygen_response_review_cumulative.xlsx')
 summary={'date_utc':'2026-10-24','catalog_records':len(cat),'completed_first_batch_candidates':len(decisions),'evidence_rows':len(evidence),'screening_stages':dict(collections.Counter(r['screening_stage'] for r in decisions)),'decisions':dict(collections.Counter(r['screening_decision'] for r in decisions)),'reading_levels':dict(collections.Counter(r['reading_level'] for r in decisions)),'remaining_unscreened_records':len(queue),'next_id':queue[0]['id'] if queue else None,'mechanism_counts_nonexclusive':dict(collections.Counter(m for e in evidence for m in e['mechanism_categories'])),'source_scope':'Unscreened records remain candidates, not included studies; original AD classifications untouched'}
 dump(ROOT/'summary.json',summary);print(json.dumps(summary,ensure_ascii=False))

if __name__=='__main__':run()
