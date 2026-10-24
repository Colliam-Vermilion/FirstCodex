#!/usr/bin/env python3
"""Explain the frozen 35-record reading queue; never change scope decisions.
Uses cached metrics and explicit per-record assessments, with no API requests.
"""
import collections,datetime as dt,hashlib,json
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.styles import Alignment,Font,PatternFill
from build_catalog import dump,csvout
ROOT=Path(__file__).resolve().parents[1];OUT=ROOT/'reading_queue35';POOL=ROOT/'all_pool'
SOURCE_COMMIT='22374f2e4526366b785a5e891f100dbe09296a8f'
RULES=[
('队列是什么','35条是从新增347条范围/机制潜力候选中由助手按具体靶点、潜在干预证据及主题覆盖人工指定的阅读队列，不是按引用/JIF选出的前35，也不是四项全部达标的终选集；固定FOCUS列表见scripts/pool_finalize.py。'),
('来源','基于22374f2e4提交冻结35个原ID及顺序；30条带PubMed来源、6条带OpenAlex来源、6条关联原AD库，来源可重叠。原AD库条目无新检索来源时如实记为原库，不冒充本轮新检出。'),
('初筛类别','原35条均为“核心候选”；这是原标签，规范解释为“范围相关/机制潜力候选，待核”，不是筛选确认的核心文献。保留旧判断供审计，本轮不静默改成终选纳入。'),
('阅读深度','入队时是题名/摘要关键句，且可复用原20评价中的已存关联；本轮逐篇补读35条完整缓存摘要。新增全文方法/结果核实0；完整缓存摘要不是全文。'),
('高被引条件','34条有OpenAlex缓存计数，2条为0，1条缺失。没有约定高被引阈值或同年/领域基准，除0不支持高被引外，不据数字擅自宣布达标；缺值不写0。引用不是范围门槛。'),
('期刊影响力条件','9条有2024 JIF/JCR机构索引缓存，25条未核实，1条为预印本不适用。沿用同年缓存及原来源，不重查JIF；有数值不等于已满足未定义的“高影响”阈值，也不等于Clarivate账户导出核实。'),
('AD相关条件','逐篇区分AD直接体系、具体产甲烷/互营/硫循环属及功能类群、仅共享代谢靶点/结构。未核实同株常见性留未核实，不凭同门/同科或泛厌氧判高度相关；纯酶/肠道/海泥/超嗜热体系列外推边界。'),
('独立验证条件','全部35的同对象/同特定主张严格独立重复未核实。已有同团队支持和跨类群暂定支持按具体DOI复用；跨团队跨类群支持不等于本篇特定链条已独立复制，更不等于广泛验证。'),
('四项是否同时满足','当前0条能确认同时满足四项；这是核实状态，不是判定35条均不符合范围或均没有价值。非AD理论、阴性结果和新文仍可优先阅读。'),
('应当收窄的主张','生物膜分层不自动证明氧屏蔽；酸抑制后恢复不等于去氧后恢复；离体FDP金属中心不等于胞内清氧；失活突变体修复不等于活细胞氧后恢复；相关/转录不等于因果/通量；ALE表述仍需遗传稳定性证据。'),
('其余文献继续','35不是最终数量上限。其余312条新增范围相关候选另列继续核实，累计479待确认、336背景追溯及原20评价全部保留。未入35不代表排除，其他候选可随证据升级/降级并进入阅读队列。'),
('没有完成的工作','本轮没有重新检索、追查缺JIF、下载/复核论文原件或开展新的无限引用追溯；仅补读缓存摘要并解释来源和条件。')]
# These are existing links or bounded reading pointers, NOT newly verified replication.
PAIRS={
'OR-3D3C740308F5':[('10.1111/febs.17040','S.wolfei另一类别FDP；仅比较线索，不能复制Class E新FeS中心结论')],
'OR-F4B8A98806CB':[('10.1007/s00775-019-01681-2','结构/模拟氧损伤模型；与本篇体内失活突变体不同，团队独立性未核'),('10.1016/j.jbc.2021.101423','离体氧损伤修复；独立性及同主张再现未核')],
'OR-E0944369CA6F':[('10.1016/j.jbc.2021.101423','后续酶学候选；不据年份或不同DOI认定独立团队'),('10.1128/jb.00070-22','体内遗传模型候选；对象/处理不同，非直接独立复制已核')],
'OR-3A00C82F8A2B':[('10.1007/s00775-019-01681-2','摘要自述此前结构工作；同团队/独立性仍须作者与方法核查'),('10.1128/jb.00070-22','体内模型候选；不能把失活突变体与氧损伤等同')],
'OR-5D303857796F':[('10.1099/mic.0.054403-0','本池B.fragilis Frd产ROS研究标题提供反例比较指针；未核其全文及团队，不能当本篇同对象复制')],
'OR-E05033DBC651':[('10.1002/anie.1942100','本池对应通信条目；尚未读取内容，不能作为反证已成立')],
'OR-52D6CC670C6A':[('10.64898/2026.06.05.730451','PubMed主记录预印本DOI'),('10.2139/ssrn.6997560','OpenAlex前向来源DOI不同；版本身份待核，不当第二独立研究')]
}
def refresh():
 if not (OUT/'assessments.json').exists():return
 baseline=json.loads((OUT/'baseline_queue35.json').read_text());notes=json.loads((OUT/'assessments.json').read_text());cat={r['id']:r for r in json.loads((ROOT/'catalog.json').read_text())};meta={r['id']:r for r in json.loads((POOL/'metadata_enrichment.json').read_text())};jc=json.loads((POOL/'journal_metric_cache.json').read_text());links=json.loads((ROOT/'reading_priority/mechanism_validation_links.json').read_text());rows=[]
 assert len(baseline)==len(notes)==35
 for q,n in zip(baseline,notes):
  assert q['id']==n['id'];rid=q['id'];c=cat[rid];m=meta[rid];j=jc[m['journal_metric_key']];v=m['citation_count'];pre='Preprint' in str(c.get('publication_types','')) or 'preprint' in c.get('journal','').lower();related=[x for x in links if x['related_id']==rid or x['focal_id']==rid]
  if rid=='OR-C63E5057C2A6':ind='已有不同团队跨类群/体系原始支持（沿用旧表、暂定）；未确认本篇具体电子路线或同对象独立重复'
  elif related:ind='已有同团队前后续/比较关联；不构成已核实的独立验证'
  else:ind='未核实独立支持/反例链；不声称不存在，也不据引用或综述数认定广泛验证'
  high='引用缺失，无法判断' if v is None else ('当前0次，不支持高被引；新文仍可因机制相关保留' if v==0 else '引用数可核，但无统一阈值/同年领域基准，未判高被引达标')
  jstatus='预印本不适用期刊JIF，不能按拟投期刊评价' if pre else ('有2024机构缓存值；未约定高影响阈值，不判已达标' if m['jif2024'] is not None else '2024指标未核实，暂无法判断；不阻塞范围筛选')
  role=n['current_reading_role'];narrow='氧进入/直接机制范围待回查；暂作背景或理论阅读' if any(w in role for w in ['待范围核实','背景阅读；','背景优先阅读；']) else '范围相关/机制潜力初筛候选，待全文及具体主张核实'
  r={'order':q['next_stage_order'],'id':rid,'title':c['title'],'short_title_zh':n['short_title_zh'],'year':c.get('year'),'journal':c.get('journal'),'doi':c.get('doi'),'pmid':c.get('pmid'),'source_databases':c.get('sources'),'source_query_ids':c.get('queries'),'source_record_provenance':c.get('provenance'),'official_record_links_not_revisited_this_turn':([f"https://pubmed.ncbi.nlm.nih.gov/{c['pmid']}/"] if c.get('pmid') else [])+([f"https://doi.org/{c['doi']}"] if c.get('doi') else []),'old_library_ids':c.get('old_library_ids'),'old_AD_classification_preserved':c.get('old_ad_classification'),'old_manual_classification_preserved':c.get('old_human_classification'),'initial_screening_batch':c.get('batch'),'initial_screening_stage':c.get('screening_stage'),'initial_screening_category_preserved':c.get('screening_decision'),'initial_category_meaning':'范围相关/机制潜力初筛候选；非已确认核心文献','current_scope_caveat':narrow,'initial_selection_reason':q['screening_reason'],'selection_reason_refined':n['selection_reason_refined'],'current_reading_role':role,'mechanism_theme':q['fulltext_focus_theme'],'reading_depth_at_selection':q['reading_basis'],'current_reading_depth':'本轮完整缓存摘要已读；全文方法/结果未核实','cached_abstract_path':'all_pool/metadata_enrichment.json','cached_abstract_locator':'[id='+rid+']/abstract_enriched','cached_abstract_sha256':hashlib.sha256(m['abstract_enriched'].encode()).hexdigest(),'cached_abstract_characters':len(m['abstract_enriched']),'abstract_source':m['abstract_from'],'microbial_objects':n['microbial_objects'],'AD_relationship_assessment':n['AD_relationship_assessment'],'AD_relationship_basis_and_boundary':n['AD_relationship_basis_and_boundary'],'abstract_reported_potential_evidence':n['abstract_reported_potential_evidence'],'unresolved_limits':n['unresolved_limits'],'citation_count':v,'citation_source':m['citation_source'],'citation_query_date':m['citation_query_date'],'citation_raw_path':m['citation_raw_path'],'citations_per_year':q['citations_per_year'],'citation_age_method':'沿用OpenAlex出版日计算，最小1年；不是领域/同年归一化','high_citation_condition':high,'JIF2024':m['jif2024'],'JIF_metric_year':2024 if m['jif2024'] is not None else None,'JCR2024_categories':m['jcr2024_categories'],'JIF_source':j['source'],'JIF_source_snapshot':j['source_snapshot'],'journal_impact_condition':jstatus,'independent_validation_condition':ind,'existing_validation_relations':[{'claim_id':x['claim_id'],'focal_id':x['focal_id'],'focal_doi':x['focal_doi'],'related_doi':x['related_doi'],'relation':x['relation'],'specific_claim':x['specific_claim'],'limitations':x['limitations'],'source_path':x['source_path'],'reading_level':x['reading_level']} for x in related],'bounded_comparison_pointers_not_verified':[{'doi':d,'status':s} for d,s in PAIRS.get(rid,[])],'strict_same_claim_independent_replication_verified':False,'all_four_conditions_confirmed':False,'all_four_status':'当前不能确认四项同时满足；不是四项均不符合的结论','queue_status':'优先阅读队列，非已确认核心文献','fulltext_verified_this_turn':False,'new_independent_validation_claims_this_turn':False,'source_queue_commit':SOURCE_COMMIT,'source_selection_implementation':'scripts/pool_finalize.py::FOCUS；人工选定后跨主题轮转，不是指标前35','followup':'优先核全文方法/结果及本行具体限制；同时继续队列外范围候选，允许升降级/补入'}
  assert m['abstract_enriched'];rows.append(r)
 dump(OUT/'source_conditions_audit.json',rows);csvout(OUT/'source_conditions_audit_utf8.csv',rows)
 # Preserve IDs and order, augment current queue only; no scope decision rewritten.
 audit={r['id']:r for r in rows}
 for file in ['next_stage_queue.json','focus_fulltext_queue.json','provisional_priorities.json']:
  q=json.loads((POOL/file).read_text())
  for item in q:
   if item['id'] in audit and item.get('screening_stage')!='全文核实' and not item.get('fulltext_verified_this_round'):
    a=audit[item['id']];item.update(queue_status='优先阅读队列，非已确认核心文献',reading_depth_at_selection=a['reading_depth_at_selection'],current_reading_depth=a['current_reading_depth'],reading_basis=a['current_reading_depth'],current_reading_role=a['current_reading_role'],source_conditions_audit='reading_queue35/source_conditions_audit.json#'+item['id'],recommendation_status='优先阅读队列；不代表四项条件均已满足或核心文献终选',processing_state='完整缓存摘要已读；全文核实未开始',next_action=a['followup'])
  dump(POOL/file,q);csvout(POOL/(file[:-5]+'_utf8.csv'),q)
 # Current reading overlay is separate from the immutable initial decisions.
 catalog_rows=list(cat.values())
 for item in catalog_rows:
  if item['id'] in audit and item.get('screening_stage')!='全文核实' and not item.get('fulltext_verified_this_round'):
   a=audit[item['id']];item.update(reading_level_at_queue_selection=a['reading_depth_at_selection'],reading_level=a['current_reading_depth'],current_reading_depth=a['current_reading_depth'],queue35_source_audit='reading_queue35/source_conditions_audit.json#'+item['id'])
 dump(ROOT/'catalog.json',catalog_rows);csvout(ROOT/'literature_master_utf8.csv',catalog_rows)
 sem=json.loads((POOL/'semantic_screening_records.json').read_text())
 for item in sem:
  if item['id'] in audit:item.update(current_reading_depth=audit[item['id']]['current_reading_depth'],current_reading_source='reading_queue35/source_conditions_audit.json#'+item['id'])
 dump(POOL/'semantic_screening_records.json',sem);csvout(POOL/'semantic_screening_records_utf8.csv',sem)
 ris_path=POOL/'new_focus_set.ris'
 ris_path.write_text(ris_path.read_text().replace('N1  - 全池暂定候选；题名/摘要关键句初筛不等于全文核实；优先补读全文（暂定）','N1  - 优先阅读队列，非已确认核心；完整缓存摘要已补读，全文及独立验证未核实'),encoding='utf-8')
 initial20={r['id'] for r in json.loads((POOL/'baseline_screening.json').read_text())};order={r['id']:r['order'] for r in json.loads((POOL/'review_order_v2.json').read_text())};other=[r for r in json.loads((POOL/'provisional_priorities.json').read_text()) if r['id'] not in audit and r['id'] not in initial20];other.sort(key=lambda r:order[r['id']])
 for i,r in enumerate(other,1):r.update(continuation_order=i,queue_status='其他范围相关候选，继续核实；未入35不等于排除',next_action='补读完整摘要及必要全文方法/结果；可随证据进入优先阅读队列')
 dump(OUT/'other_in_scope_continuation_queue.json',other);csvout(OUT/'other_in_scope_continuation_queue_utf8.csv',other)
 summary={'queue_records':35,'initial_categories':dict(collections.Counter(r['initial_screening_category_preserved'] for r in rows)),'new_complete_cached_abstracts_read':35,'new_fulltext_checks':0,'new_independence_verifications':0,'citation_counts_available':sum(r['citation_count'] is not None for r in rows),'citation_count_zero':sum(r['citation_count']==0 for r in rows),'citation_count_missing':sum(r['citation_count'] is None for r in rows),'JIF2024_available':sum(r['JIF2024'] is not None for r in rows),'journal_metric_unverified':sum(r['JIF2024'] is None and '预印本' not in r['journal_impact_condition'] for r in rows),'preprints_JIF_not_applicable':sum('预印本' in r['journal_impact_condition'] for r in rows),'strict_same_claim_independent_replication_confirmed':0,'all_four_confirmed':0,'outside35_in_scope_new_candidates_active':len(other),'all_pool_pending_confirmation_retained':len(json.loads((POOL/'pending_confirmation_queue.json').read_text())),'all_pool_background_retained':len(json.loads((POOL/'background_trace_queue.json').read_text())),'old20_preserved':True,'no_scope_decision_rewritten':True,'citation_or_JIF_requests_this_turn':0,'source_queue_commit':SOURCE_COMMIT};dump(OUT/'summary.json',summary);dump(OUT/'policy.json',[{'dimension':a,'rule':b} for a,b in RULES])
 fields=['order','id','title','year','journal','doi','source_databases','source_query_ids','old_library_ids','initial_screening_batch','initial_screening_category_preserved','initial_category_meaning','initial_selection_reason','selection_reason_refined','current_reading_role','reading_depth_at_selection','current_reading_depth','microbial_objects','AD_relationship_assessment','AD_relationship_basis_and_boundary','abstract_reported_potential_evidence','unresolved_limits','citation_count','citation_source','citation_query_date','citations_per_year','high_citation_condition','JIF2024','JIF_metric_year','JCR2024_categories','JIF_source','JIF_source_snapshot','journal_impact_condition','independent_validation_condition','existing_validation_relations','bounded_comparison_pointers_not_verified','all_four_status','queue_status','cached_abstract_path','cached_abstract_locator','source_record_provenance']
 wb=load_workbook(ROOT/'oxygen_response_review_cumulative.xlsx')
 sheets=[('35篇来源与条件',rows,fields),('35篇评价口径',[{'dimension':a,'rule':b} for a,b in RULES],None),('其他候选续核',other,None),('下一阶段新候选',json.loads((POOL/'next_stage_queue.json').read_text()),None),('重点全文候选',json.loads((POOL/'focus_fulltext_queue.json').read_text()),None),('全池暂定优先',json.loads((POOL/'provisional_priorities.json').read_text()),None),('全池筛选',sem,None)]
 for name,data,fs in sheets:
  if name in wb:idx=wb.sheetnames.index(name);del wb[name];ws=wb.create_sheet(name,idx)
  else:ws=wb.create_sheet(name)
  fs=fs or list(dict.fromkeys(k for r in data for k in r));ws.append(fs)
  for r in data:
   vals=[]
   for key in fs:
    v=r.get(key);v=json.dumps(v,ensure_ascii=False) if isinstance(v,(list,dict)) else v
    if isinstance(v,str):v=v[:32700];v="'"+v if v.startswith(('=','+','@')) else v
    vals.append(v)
   ws.append(vals)
  ws.freeze_panes='C2';ws.auto_filter.ref=ws.dimensions
  for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='244766')
  for col in ws.columns:ws.column_dimensions[col[0].column_letter].width=45
 master=wb['文献总表'];headers={c.value:c.column for c in master[1]}
 for colname in ['reading_level_at_queue_selection','current_reading_depth','queue35_source_audit']:
  if colname not in headers:headers[colname]=master.max_column+1;master.cell(1,headers[colname],colname)
 for row in master.iter_rows(min_row=2):
  rid=row[headers['id']-1].value
  if rid in audit:
   for k in ['reading_level','reading_level_at_queue_selection','current_reading_depth','queue35_source_audit']:master.cell(row[0].row,headers[k],cat[rid][k])
 wb.save(ROOT/'oxygen_response_review_cumulative.xlsx');wb.close()
 progress=json.loads((POOL/'progress.json').read_text());progress.update(queue35_source_audit='reading_queue35/source_conditions_audit.json',queue35_complete_cached_abstracts_read=35,queue35_confirmed_core_records=0,outside35_in_scope_new_candidates_active=len(other));dump(POOL/'progress.json',progress)
 # Human-readable document: concise comparison plus a complete per-record account.
 md=['# 35条优先阅读队列：来源、理由和条件核对','', '这35条不是已经筛选确认的核心文献。当前不能确认其中任何一条同时满足高被引、高期刊影响、AD对象高度相关和具体机制独立验证四项条件。缺失/未核实不等于不符合范围。','', '## 形成来源与评价口径','']
 for a,b in RULES:md+=['- **'+a+'**：'+b]
 md+=['','## 逐篇总览','','原初筛类别均为“核心候选”（原始标签，仅指范围/机制潜力）；入队时为题名/关键句层面，本轮35条完整缓存摘要均已读取，全文核实均未完成。高被引未采用统一阈值；引用、年均引用及JIF只是可核信息。空JIF不能按期刊印象补造。','', '|顺序|ID／简题|来源|当前阅读角色／选择理由|引用数／年均|2024 JIF|AD关联|独立验证条件|','|---:|---|---|---|---:|---:|---|---|']
 for r in rows:
  src='/'.join('原AD库' if x=='Original AD library' else x for x in r['source_databases']);metric=('缺失' if r['citation_count'] is None else str(r['citation_count']))+' / '+('缺失' if r['citations_per_year'] is None else str(r['citations_per_year']));j='预印本不适用' if '预印本' in r['journal_impact_condition'] else (str(r['JIF2024']) if r['JIF2024'] is not None else '未核实');md.append(f"|{r['order']}|{r['id']}<br>{r['short_title_zh']}|{src}|{r['current_reading_role']}；{r['selection_reason_refined']}|{metric}|{j}|{r['AD_relationship_assessment']}|{r['independent_validation_condition']}|")
 md+=['','## 逐篇完整记录','']
 for r in rows:
  md += [f"### {r['order']}. {r['id']} — {r['title']}",'',f"- 年份/期刊/DOI：{r['year']}；{r['journal']}；https://doi.org/{r['doi']}",f"- 来源及原库关联：{r['source_databases']}；查询ID {r['source_query_ids']}；原库ID {r['old_library_ids']}；初筛批次 {r['initial_screening_batch']}。逐来源原始文件及PMID在JSON/Excel的source_record_provenance。",f"- 初筛类别：{r['initial_screening_category_preserved']}；规范解释：{r['initial_category_meaning']}。当前边界：{r['current_scope_caveat']}。",f"- 原入队理由：{r['initial_selection_reason']}。本轮更明确的阅读理由：{r['selection_reason_refined']}。",f"- 阅读深度：入队时{r['reading_depth_at_selection']}；当前{r['current_reading_depth']}。缓存定位：`{r['cached_abstract_path']}` 的 `{r['cached_abstract_locator']}`。",f"- 微生物/AD条件：{r['microbial_objects']}；{r['AD_relationship_assessment']}；{r['AD_relationship_basis_and_boundary']}。",f"- 机制证据仍仅由摘要报告：{r['abstract_reported_potential_evidence']}。需核实：{r['unresolved_limits']}。",f"- 高被引条件：{r['citation_count'] if r['citation_count'] is not None else '缺失'}次，年均{r['citations_per_year'] if r['citations_per_year'] is not None else '缺失'}；OpenAlex，查询{r['citation_query_date'] or '未取得'}，原始缓存`{r['citation_raw_path'] or '无匹配'}`；{r['high_citation_condition']}。",f"- 期刊影响力条件：JIF2024={r['JIF2024'] if r['JIF2024'] is not None else '空'}；JCR类别/分区={json.dumps(r['JCR2024_categories'],ensure_ascii=False)}；{r['journal_impact_condition']}。来源：{r['JIF_source'] or '未核实'}；快照：`{r['JIF_source_snapshot'] or '无'}`。",f"- 独立验证条件：{r['independent_validation_condition']}；严格同对象/同主张独立重复未确认。"]
  for x in r['existing_validation_relations']:md.append(f"  - 复用旧表 {x['claim_id']}：{x['focal_doi']} ↔ {x['related_doi']}；{x['relation']}；具体主张：{x['specific_claim']}；限制：{x['limitations']}。")
  for x in r['bounded_comparison_pointers_not_verified']:md.append(f"  - 配对阅读指针：{x['doi']}；{x['status']}。")
  md += [f"- 四项合并结论：{r['all_four_status']}。队列身份：{r['queue_status']}。",'']
 md+=['## 下载与继续工作','','- [同一累计Excel](../oxygen_response_review_cumulative.xlsx)：新增“35篇来源与条件”“35篇评价口径”“其他候选续核”，原20评价保留。','- [35篇CSV](source_conditions_audit_utf8.csv)／[JSON](source_conditions_audit.json)；[原35队列快照](baseline_queue35.json)；[本轮摘要阅读日志](abstract_reading_log.json)。','- [其余312条范围相关候选CSV](other_in_scope_continuation_queue_utf8.csv)／[JSON](other_in_scope_continuation_queue.json)。它们仍待核实，35不是截断线；其他条目可进入队列。累计479待确认、336背景及原20各自状态继续保留。','- 原[35条RIS](../all_pool/new_focus_set.ris)仍是阅读队列导入文件，不是终选核心集。','- 四项条件的缺口按具体主张补证；没有重新检索或重复查询原20指标，没有新的全文核实或独立验证断言。','']
 (OUT/'README.md').write_text('\n'.join(md),encoding='utf-8');print(json.dumps(summary,ensure_ascii=False))
if __name__=='__main__':refresh()
