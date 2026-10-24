#!/usr/bin/env python3
"""Render whole-pool triage without changing the original 20 expert evaluations.
The annotations are conservative text-derived routing, not validated mechanisms.
"""
import collections,datetime as dt,json,re,time
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.styles import Font,PatternFill
from build_catalog import dump,csvout,normtitle
ROOT=Path(__file__).resolve().parents[1];OUT=ROOT/'all_pool'
def hit(p,t):return bool(re.search(p,t,re.I))
# Explicit assistant selection after reading all packets; not impact-metric selection.
FOCUS={
'M1 损伤与修复':['OR-3A00C82F8A2B','OR-E0944369CA6F','OR-F4B8A98806CB','OR-E05033DBC651','OR-FA5B37263F6F','OR-8C75CC78BFD2','OR-5D303857796F'],
'M2 古菌及厌氧菌解毒':['OR-3D3C740308F5','OR-454A09DABEFE','R09148','OR-C63E5057C2A6','OR-82F6FC6E3A3E','OR-A47D88A5BD92','OR-7018E028D4E6','OR-C78553A0F4D0','OR-49242898AE4A'],
'M3 代谢切换与条件':['OR-81B738B20350','OR-124B294CA364','OR-CEDAA95DD383','OR-3B1D70862C75','OR-F6D3E9F8E669','R13274','OR-52D6CC670C6A'],
'M4 群落互作':['R11596','R13833','R10086','OR-01782ECD28E3'],
'M5 空间保护':['OR-E9124F54EA76','OR-8E58C5BBDB53'],
'M6 恢复及反例':['R07494','OR-15A62A4556E3','OR-071CAFC18C32','OR-A8CEEF69FB67','OR-E551A6ADF474','OR-631665AE20A2']}
POLICY=[
('阅读依据','新1516条为助手逐项读取题名及所选摘要关键句（片段可截断），不是完整摘要/全文审阅；缺摘要为题名初筛。待确认数量由进度动态给出；不称所有摘要信息不足，下一步可先补读完整摘要。'),
('计算预筛','computational_prescreen.json仅路由1516条，不计完成；只有独立decisions逐项判断才计题名/关键句初筛完成。'),
('科学排序','原20评价保留；新增重点全文核验集经助手按具体对象/靶点和跨主题覆盖选择。其他候选按AD工艺/功能关联和潜在干预线索分层；自动注释仅辅助，不确认属种相同、因果或可迁移。'),
('指标辅助','同一科学层内以OpenAlex年均引用、2024 JIF作辅助，不将范围排除/待确认提升为核心。缺失为null/空单元格，查询确认的0才为0。新近文献不因累计引用低而排除。'),
('引用年限','使用OpenAlex出版日；缺日用已知年份1月1日并注明。年限=(2026-10-24-出版日)/365.2425，最小1年，避免新文短时放大；不是学科/同年归一化。'),
('JIF口径','2024指标年/2025发布版已存机构索引缓存；逐期刊或ISSN统一匹配，不是本轮Clarivate账户导出。缺失不补造，不使用其他分区代替JCR。'),
('推荐边界','核心候选是范围/机制潜力初筛标签，不是已证实的核心机制文献；理论/综述/模型背景单独队列，不混入原始研究优先数量。'),
('直接O2与ROS','NO/H2O2/金属/药物ROS分开标记，背景不冒充直接分子氧进入证据。'),
('独立验证','新增候选全部未核实；保留原20既有验证关联。重复题名、预印本/期刊/学位/结构数据不得视作独立实验。'),
('外推限制','16S变化非功能、基因非活性、网络非互作、导电材料非DIET、恢复或驯化非遗传适应；自然体系需考虑水分/温度/矿物/根际混杂。'),
('续跑','未开始队列为空不代表待确认已解决。重点全文队列、动态待确认队列和背景追溯队列独立保存；继续时不重检索、不过度追指标、不复核未改原件。'),
('时长','metadata_stage.json的48.25秒为完整元数据阶段计时。checkpoints时间仅写入区间，render_timing仅导出耗时；不作为完整人工语义筛选/全文阅读耗时。')]
def annotations(r):
 t=r['title']+' '+r.get('abstract','');parts=[]
 ad=0
 if hit(r'anaerobic digest|digester|methanogenic sludge|anaerobic gran',t):ad=3;parts.append('AD/厌氧污泥工艺直接相关；具体菌种和净过程机制仍待核')
 if hit(r'methanogen|methanosarc|methanobrev|methanotherm|methanobacter|methanolobus',t):ad=max(ad,2);parts.append('产甲烷功能类群可比；未据同门/同科认定相同AD物种')
 if hit(r'syntroph|syntrophic acetate|fatty acid.*degrad|LCFA',t):ad=max(ad,2);parts.append('脂肪酸/乙酸互营过程可比；H2/甲酸或DIET须干预实证')
 if hit(r'acetogen|homoacet',t):ad=max(ad,2);parts.append('产乙酸功能可比；底物/路径和伙伴差异待核')
 if hit(r'sulfate.reduc|sulphate.reduc|Desulf|sulfide oxid|sulfur cycle',t):ad=max(ad,2);parts.append('AD硫循环潜在可比；硫形态及氧位置差异需核')
 if hit(r'ferment|hydrolys|Bacteroides|Clostridi|Lactobac|Lactiplanti|Phocaeicola',t):ad=max(ad,1);parts.append('发酵/水解或低电位代谢靶点可类比；不等同该菌就是AD优势功能菌')
 if hit(r'anammox|Kuenenia',t):parts.append('Anammox氮代谢不是AD产甲烷；仅氧敏性/空间保护可比较')
 if hit(r'FNR|iron.sul|Fe.S|ferredoxin|hydrogenase|flavodiiron|rubrerythrin|superoxide reductase',t):ad=max(ad,1);parts.append('氧敏靶点/防御系统线索可比；AD同源性与体内作用未核实')
 if hit(r'gut|intestinal|rumen|host|infection',t):parts.append('非AD宿主体系需分开宿主介导ROS、营养与直接O2效应')
 if hit(r'sediment|soil|peat|wetland|paddy',t):parts.append('自然体系扩散、水分、温度、矿物及根际输入存在外推边界')
 if not parts:parts=['与AD相同属种/功能或靶点的对应未核实；不得按泛厌氧归为高度相关']
 organisms=[]
 for genus in ['Methanosarcina','Methanobrevibacter','Methanothermobacter','Methanobacterium','Methanosaeta','Methanothrix','Methanolobus','Syntrophomonas','Syntrophobacter','Tepidanaerobacter','Clostridium','Clostridioides','Bacteroides','Phocaeicola','Desulfovibrio','Desulfobulbus','Fusobacterium','Filifactor','Pyrococcus','Thermococcus','Pseudomonas','Escherichia','Shewanella','Geobacter','Faecalibacterium','Eubacterium','Lactobacillus','Lactiplantibacillus','Akkermansia','Bifidobacterium','Dehalococcoides','Dehalogenimonas','Wolinella','Paracoccus','Treponema','Prevotella','Giardia','Entamoeba','Spironucleus','Staphylococcus','Salmonella','Klebsiella','Anaerobutyricum']:
  match=re.search(r'\b'+genus+r'(?:\s+[a-z][a-z-]+)?',t)
  if match:organisms.append(match.group(0))
 systems=[]
 for label,pat in [('AD/厌氧污泥',r'anaerobic digest|digester|anaerobic sludge'),('发酵培养',r'ferment|butanol|acidogen'),('纯/共培养线索',r'pure cultur|co.cultur|cocultur|mutant'),('沉积物/土壤',r'sediment|soil|peat|wetland'),('肠道/瘤胃',r'gut|intestinal|rumen'),('离体酶学',r'purified|recombinant|enzyme|enzymatic')]:
  if hit(pat,t):systems.append(label)
 scenarios=[]
 for label,pat in [('急性/短时线索',r'shock|pulse|acute|short.term|transient'),('持续微氧线索',r'microaer|micro.aerat|microoxic|oxygen.limit|low oxygen'),('间歇/重复线索',r'intermittent|cyclic|repeated|fluctuat|oscillat'),('去氧后恢复线索',r'recover|reversib|resilien')]:
  if hit(pat,t):scenarios.append(label)
 return {'id':r['id'],'system_text_hints':systems,'microbial_names_text_mentions':organisms,'AD_relation_text_basis':parts,'AD_relevance_stratum_potential_only':ad,'exposure_text_hints_not_verified_design':scenarios,'annotation_basis':'文本辅助注释，须结合逐项语义理由；不是同物种或实验设计核实'}
def ris(path,rows,cat):
 lines=[]
 for r in rows:
  c=cat[r['id']];typ=str(c.get('document_type','')).lower();rtype='THES' if 'thesis' in typ else ('CONF' if 'conference' in typ else ('RPRT' if 'preprint' in typ or 'posted-content' in typ else ('CHAP' if 'chapter' in typ else 'JOUR')));lines+=['TY  - '+rtype,'ID  - '+r['id'],'TI  - '+c['title']]
  for tag,k in [('PY','year'),('JO','journal'),('DO','doi'),('UR','url')]:
   if c.get(k):lines.append(tag+'  - '+str(c[k]).replace('\n',' '))
  authors=c.get('authors') or []
  if isinstance(authors,str):authors=[authors]
  for a in authors:lines.append('AU  - '+str(a))
  lines+=['N1  - 全池暂定候选；题名/摘要关键句初筛不等于全文核实；'+str(r.get('tier','')),'ER  -','']
 path.write_text('\n'.join(lines),encoding='utf-8')
def run():
 started=time.monotonic();cat=json.loads((ROOT/'catalog.json').read_text());byid={r['id']:r for r in cat};meta={r['id']:r for r in json.loads((OUT/'metadata_enrichment.json').read_text())};screen=json.loads((ROOT/'screening.json').read_text());sb={r['id']:r for r in screen};old=json.loads((ROOT/'reading_priority/evaluations.json').read_text());oldrec={r['文献ID']:r for r in old if r['正式精读推荐']};ann={r['id']:annotations(r) for r in cat};focusmap={rid:theme for theme,ids in FOCUS.items() for rid in ids};allrows=json.loads((OUT/'scope_candidates_before_priority_split.json').read_text());background=[];priority=[]
 for r in allrows:
  a=ann[r['id']];m=meta[r['id']];r.update(a);r['old_library_ids']=byid[r['id']].get('old_library_ids',[]);r['old_AD_classification_preserved']=byid[r['id']].get('old_ad_classification');r['old_manual_classification_preserved']=byid[r['id']].get('old_human_classification');r['publication_type']=byid[r['id']].get('document_type');r['early_foundational_pre2012']=str(r.get('year','')).isdigit() and int(r['year'])<2012
  pub=m.get('publication_date_for_age');method='OpenAlex出版日'
  if not pub and str(r.get('year','')).isdigit():pub=str(r['year'])+'-01-01';method='仅年份，按1月1日近似'
  try:age=max(1,(dt.date(2026,10,24)-dt.date.fromisoformat(pub)).days/365.2425)
  except Exception:age=None
  r['citation_age_years_min1']=round(age,3) if age is not None else None;r['citation_age_basis']=method if age else '缺日期';r['citations_per_year']=round(m['citation_count']/age,3) if age is not None and m['citation_count'] is not None else None
  r['independent_validation']='见原20既有评价，不新增判断' if r['id'] in oldrec else '未核实';r['fulltext_focus_theme']=focusmap.get(r['id']);r['potential_direct_intervention_hint']=any(any(k in e for k in ['遗传','酶学','共培养','微剖面']) for e in r['potential_evidence']);r['AD_relation']='；'.join(a['AD_relation_text_basis'])
  if sb[r['id']]['screening_decision']=='机制背景' and r['id'] not in oldrec:background.append(r);continue
  if r['id'] in oldrec:r['tier']='原20精读推荐（原评价保留）';r['prior_detailed_layer']=oldrec[r['id']]['推荐主层级']
  elif r['id'] in focusmap:r['tier']='优先补读全文（暂定）'
  elif a['AD_relevance_stratum_potential_only']>=2:r['tier']='AD过程/功能相关候选（暂定）'
  else:r['tier']='可迁移机制/边界候选（暂定）'
  r['scope_label_warning']='核心候选仅初筛标签，不代表机制已直接验证';priority.append(r)
 def key(r):
  annual=r['citations_per_year'];jif=r['jif2024']
  return (0 if r['id'] in oldrec or r['id'] in focusmap else 1,-r['AD_relevance_stratum_potential_only'],-int(r['potential_direct_intervention_hint']),annual is None,-annual if annual is not None else 0,jif is None,-jif if jif is not None else 0,r['id'])
 priority.sort(key=key);focus=[r for r in priority if r['id'] in oldrec or r['id'] in focusmap];pending=[{**r,**ann[r['id']],'next_action':'先补读完整摘要；仍无法判断再读必要全文；不预设全文不存在'} for r in screen if r['screening_decision']=='待全文确认' or r['screening_stage']=='待确认'];screen_enriched=[{**r,**ann[r['id']]} if r['id'] not in oldrec else {**r,'previous20_detailed_evaluation':'reading_priority/evaluations.json'} for r in screen]
 for item in screen_enriched:
  if item.get('batch') in ['P001','P002']:item['abstract_fragment_provenance']='缓存摘要前450字定位，未保存当时显示片段；不作逐字已读证据引文'
  elif item.get('batch','').startswith('P'):item['abstract_fragment_provenance']='对应判断包保存的所示摘要关键句，可截断；非全文引文'
 for name,rows in [('assessment_annotations',list(ann.values())),('provisional_priorities',priority),('background_trace_queue',background),('focus_fulltext_queue',focus),('pending_confirmation_queue',pending),('semantic_screening_records',screen_enriched)]:dump(OUT/(name+'.json'),rows);csvout(OUT/(name+'_utf8.csv'),rows)
 newfocus=[];bins=collections.defaultdict(list)
 for r in focus:
  if r['id'] not in oldrec:bins[r['fulltext_focus_theme']].append(r)
 while any(bins.values()):
  for theme in sorted(bins):
   if bins[theme]:
    r=dict(bins[theme].pop(0));r.update(next_stage_order=len(newfocus)+1,next_action='补读完整摘要及必要全文方法/结果；有原件先复用；新增独立验证仍未核实',processing_state='未开始全文核验');newfocus.append(r)
 dump(OUT/'next_stage_queue.json',newfocus);csvout(OUT/'next_stage_queue_utf8.csv',newfocus)
 ris(OUT/'new_focus_set.ris',newfocus,byid)
 # Stable queues with no semantic duplicates removed silently.
 titles=collections.defaultdict(list)
 for r in cat:titles[normtitle(r['title'])].append(r['id'])
 duplicate_groups=[{'ids':ids,'status':'规范化同题版本/身份待核；不能视为独立验证','titles':[byid[i]['title'] for i in ids],'dois':[byid[i].get('doi') for i in ids]} for ids in titles.values() if len(ids)>1];dump(OUT/'possible_version_groups.json',duplicate_groups)
 dump(OUT/'evaluation_policy.json',[{'topic':a,'rule':b} for a,b in POLICY]);ris(OUT/'provisional_priority_set.ris',priority,byid);ris(OUT/'focus_fulltext_set.ris',focus,byid)
 progress=json.loads((OUT/'progress.json').read_text());progress.update(provisional_priority_articles=len(priority),background_trace_records=len(background),focus_fulltext_records=len(focus),new_focus_next_stage_records=len(newfocus),pending_confirmation_records=len(pending),provisional_priority_count_unit='书目记录；未合并可能版本；不等于独立实验数',screening_basis='题名/摘要关键句初筛；不是全摘要/全文阅读',journal_metric_available=sum(m.get('jif2024') is not None for m in meta.values()),missing_abstracts=sum(not m['abstract_enriched'] for m in meta.values()));dump(OUT/'progress.json',progress)
 wb=load_workbook(ROOT/'oxygen_response_review_cumulative.xlsx')
 for name,rows in [('全池筛选',screen_enriched),('全池暂定优先',priority),('背景追溯队列',background),('待确认队列',pending),('重点全文候选',focus),('下一阶段新候选',newfocus),('全池评价口径',[{'topic':a,'rule':b} for a,b in POLICY])]:
  if name in wb:idx=wb.sheetnames.index(name);del wb[name];ws=wb.create_sheet(name,idx)
  else:ws=wb.create_sheet(name)
  fields=list(dict.fromkeys(k for r in rows for k in r)) or ['id','title','status'];ws.append(fields)
  for r in rows:
   values=[]
   for k in fields:
    v=r.get(k);v=json.dumps(v,ensure_ascii=False) if isinstance(v,(dict,list)) else v
    if isinstance(v,str):v=v[:32700];v="'"+v if v.startswith(('=','+','@')) else v
    values.append(v)
   ws.append(values)
  ws.freeze_panes='C2';ws.auto_filter.ref=ws.dimensions
  for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='244766')
  for col in ws.columns:ws.column_dimensions[col[0].column_letter].width=45
 if not json.loads((ROOT/'remaining_queue.json').read_text()):
  idx=wb.sheetnames.index('续跑队列');del wb['续跑队列'];ws=wb.create_sheet('续跑队列',idx);ws.append(['order','id','title','screening_stage','next_action']);(ROOT/'remaining_queue_utf8.csv').write_text('order,id,title,screening_stage,next_action\n',encoding='utf-8-sig')
 wb.save(ROOT/'oxygen_response_review_cumulative.xlsx');wb.close()
 summary=json.loads((ROOT/'summary.json').read_text());summary.update(all_pool_provisional_priority=len(priority),all_pool_pending_confirmation=len(pending),all_pool_focus_fulltext=len(focus),next_stage='重点候选全文及待确认记录补读完整摘要；不重检索');dump(ROOT/'summary.json',summary)
 with (OUT/'finalize_timing.jsonl').open('a') as f:f.write(json.dumps({'at_utc':dt.datetime.now(dt.timezone.utc).isoformat(),'measured_export_seconds':round(time.monotonic()-started,3)})+'\n')
 print(json.dumps(progress,ensure_ascii=False))
if __name__=='__main__':run()
