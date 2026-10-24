#!/usr/bin/env python3
"""Verify the priority additions independently from the immutable GitHub commit."""
import argparse,concurrent.futures,datetime as dt,hashlib,json,os,subprocess
from pathlib import Path
from urllib.parse import quote,urlencode
from urllib.request import Request,urlopen
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parents[1];REPO=ROOT.parents[2];OUT=ROOT/'reading_priority'
API='https://api.github.com/repos/Colliam-Vermilion/FirstCodex/'
def get(path,raw=False):
 h={'User-Agent':'OxygenReadingPriorityAudit/1.0','Accept':'application/vnd.github.raw+json' if raw else 'application/vnd.github+json'}
 if os.environ.get('GH_TOKEN'):h['Authorization']='Bearer '+os.environ['GH_TOKEN']
 with urlopen(Request(API+path,headers=h),timeout=55) as r:return r.read()
def run(commit):
 base=json.loads((OUT/'baseline.json').read_text())['source_commit'];j=json.loads(get('git/commits/'+commit));tree=json.loads(get('git/trees/'+j['tree']['sha']+'?recursive=1'));assert not tree.get('truncated')
 remote={r['path']:r['sha'] for r in tree['tree'] if r['type']=='blob'}
 changed=subprocess.check_output(['git','diff','--name-only',base,commit],cwd=REPO,text=True).splitlines();objects=[]
 for rel in changed:
  f=REPO/rel;assert f.exists();b=f.read_bytes();oid=hashlib.sha1(b'blob '+str(len(b)).encode()+b'\0'+b).hexdigest();assert remote[rel]==oid,rel
  objects.append({'path':rel,'bytes':len(b),'sha256':hashlib.sha256(b).hexdigest(),'remote_git_blob_matches':True})
 files=['oxygen_response_review_cumulative.xlsx','reading_priority/README.md','reading_priority/evaluations.json','reading_priority/priority_reading_utf8.csv','reading_priority/priority_reading_set.ris','reading_priority/mechanism_validation_links.json','reading_priority/mechanism_validation_links_utf8.csv','reading_priority/evaluation_policy.json','reading_priority/evaluation_policy_utf8.csv','reading_priority/progress.json','reading_priority/validation.json','reading_priority/raw/openalex_citations_20.json','reading_priority/raw/pubmed_selected_and_support.xml']
 prefix=ROOT.relative_to(REPO).as_posix()+'/';dest=Path('/tmp/reading-priority-remote')/commit;dest.mkdir(parents=True,exist_ok=True)
 def read_one(name):
  path=prefix+name;b=get('contents/'+quote(path,safe='/')+'?'+urlencode({'ref':commit}),True)
  if b.startswith(b'{'):
   try:j=json.loads(b)
   except Exception:j={}
   if j.get('encoding')=='base64' and 'content' in j:
    import base64;b=base64.b64decode(j['content'])
  assert hashlib.sha256(b).digest()==hashlib.sha256((ROOT/name).read_bytes()).digest(),name
  f=dest/name;f.parent.mkdir(parents=True,exist_ok=True);f.write_bytes(b)
  return {'path':path,'bytes':len(b),'sha256':hashlib.sha256(b).hexdigest(),'independent_remote_byte_readback':True}
 with concurrent.futures.ThreadPoolExecutor(max_workers=2) as pool:readbacks=list(pool.map(read_one,files))
 wb=load_workbook(dest/'oxygen_response_review_cumulative.xlsx',read_only=True);sheets={ws.title:ws.max_row-1 for ws in wb};wb.close()
 assert sheets['优先精读']==20 and sheets['机制验证关联']==26 and sheets['评价口径']==13
 evals=json.loads((dest/'reading_priority/evaluations.json').read_text());assert len(evals)==20;selected={r['文献ID'] for r in evals if r['正式精读推荐']}
 import re
 assert set(re.findall(r'^ID  - (.+)$',(dest/'reading_priority/priority_reading_set.ris').read_text(),re.M))==selected
 result={'payload_commit':commit,'branch':'results/ad-oxygen-literature-20261006','verified_at_Asia_Shanghai':dt.datetime.now(dt.timezone(dt.timedelta(hours=8))).isoformat(),'changed_files_remote_git_object_matches':len(objects),'remote_git_objects':objects,'independently_downloaded_key_files':len(readbacks),'readbacks':readbacks,'workbook_opened_from_downloaded_bytes':True,'workbook_sheets':sheets,'remote_evaluations':len(evals),'remote_RIS_recommended_ids':len(selected),'verification':'passed','scope':'All changed files matched immutable GitHub tree IDs; 13 listed key payloads additionally downloaded, SHA256 checked and workbook/RIS/JSON reopened. Local /tmp is verification workspace, not delivery location.'}
 (OUT/'remote_readback.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8');print(json.dumps({k:v for k,v in result.items() if k not in ['remote_git_objects','readbacks']},ensure_ascii=False))
if __name__=='__main__':
 ap=argparse.ArgumentParser();ap.add_argument('--commit',required=True);run(ap.parse_args().commit)
