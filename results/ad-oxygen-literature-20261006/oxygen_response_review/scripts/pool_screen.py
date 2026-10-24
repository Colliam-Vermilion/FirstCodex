#!/usr/bin/env python3
"""Transparent whole-pool prescreen + separately recorded assistant reviews.

Rules only route records. They do NOT count as semantic screening completion.
Explicit per-record judgments are supplied by the assistant after reading packets.
"""
import argparse,collections,datetime as dt,hashlib,json,math,re,sys,time
from pathlib import Path
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Alignment,Font,PatternFill
ROOT=Path(__file__).resolve().parents[1];OUT=ROOT/'all_pool';TZ=dt.timezone(dt.timedelta(hours=8));sys.path.insert(0,str(ROOT/'scripts'))
from build_catalog import dump,csvout,normtitle
def now():return dt.datetime.now(TZ).isoformat()
def hit(p,t):return bool(re.search(p,t,re.I))
EXPO=r'oxygen (?:exposure|shock|stress|toleran|intrusion|pulse|consum|reduct|sensitiv)|aerotoler|reoxygen|micro.?aerat|micro.?aerob|microoxic|redox (?:fluctuat|oscillat)|low o2|o2 tension|air.adapt|aeration'
ANA=r'anaerob|methanogen|methanosarc|methanotherm|methanobrev|methanosaeta|syntroph|acetogen|clostridi|bacteroid|faecalibacter|desulf|ferment|sediment|peat|rumen|gut |intestinal|waterlogged|paddy'
MECH=r'mutant|delet|knockout|complement|enzyme|protein|metabol|transcript|genom|microbial|communit|physiolog|isotope|microelectrode|microprofile|gradient|rubrerythrin|flavodiiron|superoxide|catalase|peroxidase|ferrodoxin|ferredoxin|hydrogenase|electron'
NONMIC=r'cardiorespir|cardio.pulmonary|myocardial|athlete|exercise (?:capacity|test|train)|rowers|skiers|basketball|biventricular|lymphocyte proliferation|human skeletal|submergence.*(?:seedling|rice)|plant abiotic|soybean seedlings|caspase.3.*prawn|fish.*hypoxia|oxygen tolerance of marine fauna'
def classify_hint(r,m):
 text=r['title']+' '+m.get('abstract_enriched','');title=r['title'];micro=[]
 for pat,label in [(r'Methanosarcina\s+\w+','Methanosarcina'),(r'Methanobacter\w*','Methanobacterium/相关'),(r'Methanothermobacter\s+\w+','Methanothermobacter'),(r'Methanobrevibacter\s+\w+','Methanobrevibacter'),(r'Methanosaeta\s+\w+','Methanosaeta'),(r'Syntroph\w+(?:\s+\w+)?','互营类群'),(r'Clostridi\w+(?:\s+\w+)?','Clostridia相关'),(r'Bacteroides\s+\w+','Bacteroides'),(r'Desulf\w+(?:\s+\w+)?','硫酸盐还原相关'),(r'Faecalibacterium\s*\w*','Faecalibacterium'),(r'Akkermansia\s+\w+','Akkermansia'),(r'Kuenenia|anammox','Anammox')]:
  match=re.search(pat,text,re.I)
  if match:micro.append(match.group(0))
 themes=[]
 for label,pat in [('M1 损伤',r'inactiv|iron.sul|Fe.S|damage|DNA repair|protein repair|pyruvate|glycyl radical'),('M2 解毒',r'superoxide|rubrerythrin|flavodiiron|catalase|peroxidase|detox|scaveng|aerotoler|oxygen toler'),('M3 代谢',r'metabol|ferment|electron flow|respirat|hydrogenase|carbon flux|methanogen'),('M4 互作',r'syntroph|co.cultur|cross.feed|facultative|mutual|hydrogen.transfer|formate'),('M5 空间',r'microprofile|microelectrode|oxygen gradient|biofilm|granul|aggregat|filament|diffusion'),('M6 恢复',r'recover|adapt|resilien|transient|repeated|pulse|shock'),('M7 条件',r'dose|frequency|intermittent|redox|oxygen concentration|aeration rate')]:
  if hit(pat,text):themes.append(label)
 ev=[]
 for label,pat in [('遗传干预线索',r'mutant|deletion|knockout|complement|site.directed'),('直接酶学线索',r'purified|recombinant|enzymatic activity|enzyme activity|kinetic|kcat'),('共培养干预线索',r'co.cultur|cocultur|synthetic consortium'),('微剖面/同位素线索',r'microelectrode|microprofile|isotop|13C|14C'),('生理/代谢实测线索',r'growth|survival|methane production|fermentation product'),('组学/群落线索',r'transcript|proteom|metagenom|16S|community composition')]:
  if hit(pat,text):ev.append(label)
 ad=3 if hit(r'anaerobic digest|methanogenic sludge|anaerobic gran|syntroph|methanosarc|methanosaeta|methanothermobacter|acetogen',text) else 2 if hit(r'methanogen|ferredoxin|flavodiiron|superoxide reductase|rubrerythrin|hydrogenase|ferment|sulfate.reduc',text) else 1
 score=8*int(hit(EXPO,title))+4*int(hit(EXPO,text))+4*int(hit(ANA,text))+3*int(hit(MECH,text))+ad+3*int(any('遗传' in x or '酶学' in x for x in ev))
 if hit(NONMIC,title):bucket='明确非目标对象线索';score=-30
 elif not m.get('abstract_enriched'):bucket='缺摘要，须核实';score-=2
 elif hit(EXPO,text) and hit(ANA,text) and hit(MECH,text):bucket='氧响应机制候选线索'
 elif hit(EXPO,text) and hit(ANA,text):bucket='有氧过程但机制待核'
 else:bucket='范围/机制不明，待语义判断'
 return {'id':r['id'],'title':title,'routing_only':True,'routing_bucket':bucket,'scientific_hint_score':score,'microbial_objects_hint':micro,'AD_relation_hint_level':ad,'AD_relation_warning':'词线索仅排队，不能仅凭同门/同科/厌氧判高度相关','mechanism_themes_hint':themes,'evidence_type_hint':ev,'abstract_available':bool(m.get('abstract_enriched')),'semantic_screening_complete':False}
def init():
 start=time.monotonic();cat=json.loads((OUT/'baseline_catalog.json').read_text());meta={m['id']:m for m in json.loads((OUT/'metadata_enrichment.json').read_text())};prior={r['id'] for r in json.loads((OUT/'baseline_screening.json').read_text())}
 routed=[classify_hint(r,meta[r['id']]) for r in cat if r['id'] not in prior];dump(OUT/'computational_prescreen.json',routed)
 if not (OUT/'review_order_v2.json').exists():
  bins=collections.defaultdict(list)
  for r in routed:bins[(r['mechanism_themes_hint'] or ['M0 未分配'])[0]].append(r)
  for rows in bins.values():rows.sort(key=lambda r:(-r['scientific_hint_score'],r['id']))
  order=[]
  while any(bins.values()):
   for key in sorted(bins):
    if bins[key]:x=bins[key].pop(0);order.append({'order':len(order)+1,'id':x['id'],'theme_bin':key})
  dump(OUT/'review_order_v2.json',order)
 # Cache JIF/JCR ONCE per distinct journal/ISSN. No per-paper network lookup.
 sources=json.loads((ROOT/'reading_priority/journal_metric_source_register.json').read_text());aliases={'thefebsjournal':'febsjournal','theismejournal':'ismejournal','microbiologyreadingengland':'microbiologysgm'};jcache={}
 for r in cat:
  m=meta[r['id']];key=m.get('issn_l') or normtitle(m.get('journal_source_name') or r.get('journal','')) or 'unknown';names=[normtitle(r.get('journal','')),normtitle(m.get('journal_source_name') or '')];match=next((sources.get(aliases.get(n,n)) for n in names if sources.get(aliases.get(n,n))),None)
  if key not in jcache:jcache[key]={'journal_key':key,'journal_names':[],'issn_l':m.get('issn_l'),'jif2024':match['jif'] if match else None,'jcr2024_categories':match['categories'] if match else [],'source':match['source_url'] if match else None,'source_snapshot':match['source_snapshot'] if match else None,'status':'cached 2024 institutional JCR index snapshot; not Clarivate account export' if match else '2024 value not verified; blank; screening not blocked'}
  jcache[key]['journal_names']=sorted(set(jcache[key]['journal_names']+[r.get('journal','')]))
  m['journal_metric_key']=key;m['jif2024']=jcache[key]['jif2024'];m['jcr2024_categories']=jcache[key]['jcr2024_categories'];m['journal_metric_status']=jcache[key]['status']
 dump(OUT/'journal_metric_cache.json',jcache);dump(OUT/'metadata_enrichment.json',list(meta.values()))
 dump(OUT/'prescreen_stage.json',{'stage':'computational_routing_not_semantic_completion','date_Asia_Shanghai':now(),'routed':len(routed),'buckets':dict(collections.Counter(r['routing_bucket'] for r in routed)),'unique_journal_ISSN_cache_keys':len(jcache),'per_paper_JIF_requests':0,'measured_wall_seconds':round(time.monotonic()-start,3)})
 print('routing',len(routed),'journals',len(jcache))
def packet(offset,size,mode):
 cat={r['id']:r for r in json.loads((OUT/'baseline_catalog.json').read_text())};meta={r['id']:r for r in json.loads((OUT/'metadata_enrichment.json').read_text())};done={p.stem for p in (OUT/'decisions').glob('*.json')} if (OUT/'decisions').exists() else set();order=[r for r in json.loads((OUT/'review_order_v2.json').read_text()) if r['id'] not in done]
 if mode=='clear':order=sorted(order,key=lambda o:(not hit(NONMIC,cat[o['id']]['title']),o['order']))
 if mode=='molecular':
  order.sort(key=lambda o:-(8*int(hit(r'mutant|knockout|deletion|complement|enzyme|hydrogenase|rubrerythrin|flavodiiron|superoxide|PFOR|pyruvate|iron.sul|Fe.S|catalase|peroxidase',cat[o['id']]['title']))+3*int(hit(EXPO,cat[o['id']]['title']))))
 selected=order[offset:offset+size];path=OUT/'packets'/f'{mode}_{len(done):04}_{offset:04}_{size}.json'
 for i,o in enumerate(selected,1):
  r=cat[o['id']];ab=meta[r['id']]['abstract_enriched'];sent=re.split(r'(?<=[.!?])\s+',ab);salient=sorted(enumerate(sent),key=lambda z:-(3*int(hit(EXPO,z[1]))+2*int(hit(r'mutant|result|show|found|oxygen|O2',z[1]))))[:2];snippet=' '.join(s for _,s in sorted(salient))
  o['shown_excerpt']=snippet[:350] if mode!='clear' else snippet[:140]
  print(f'{i}|{r["id"]}|{r["title"]}|{o["shown_excerpt"]}')
 dump(path,selected)
 print('PACKET',str(path.relative_to(ROOT)))
def apply(path):
 packet_data=json.loads(Path(path).read_text());byid={r['id']:r for r in json.loads((OUT/'baseline_catalog.json').read_text())};meta={r['id']:r for r in json.loads((OUT/'metadata_enrichment.json').read_text())};routing={r['id']:r for r in json.loads((OUT/'computational_prescreen.json').read_text())};start=now();count=0
 # Input judgments explicitly supplied after packet reading. Not generated by
 # a keyword classifier. No claims of full abstract/full text if excerpt-only.
 for item in packet_data:
  rid=item['id'];dest=OUT/'decisions'/f'{rid}.json'
  if dest.exists():continue
  r=byid[rid];h=routing[rid];decision=item['decision'];reason=item['reason'];ab=meta[rid]['abstract_enriched']
  obj={'id':rid,'batch':item['batch'],'batch_order':h.get('scientific_hint_score',0),'screening_stage':'排除' if decision=='排除' else ('待确认' if decision=='待全文确认' else '摘要初筛'),'screening_decision':decision,'screening_reason':reason,'reading_level':'题名及摘要关键句语义审阅（未阅读全文）' if ab else '题名语义审阅；缺摘要','review_method':'助手逐项语义审阅；区别于计算辅助预筛','title':r['title'],'doi':r.get('doi',''),'year':r.get('year',''),'journal':r.get('journal',''),'research_system':item.get('objects') or h['microbial_objects_hint'] or '摘要未明确识别；不编造物种','AD_relationship':item.get('ad_relation') or ('潜在同功能/特定靶点可比，须精读确认；不以分类层级单独判高度相关' if h['AD_relation_hint_level']>=2 else '与AD对象对应未核实，仅机制/方法类比候选'),'AD_hint_level_not_verified':h['AD_relation_hint_level'],'mechanism_themes':h['mechanism_themes_hint'],'potential_evidence_type':h['evidence_type_hint'],'evidence_not_confirmed':'题名/摘要线索，不代表因果机制已证实或已独立验证','source_abstract_available':bool(ab),'abstract_evidence_excerpt':item.get('excerpt') or item.get('shown_excerpt') or (ab[:450] if ab else r['title']),'source_provenance':r.get('provenance',[]),'processed_Asia_Shanghai':now(),'fulltext_verified_this_round':False,'independent_validation':'未核实；本轮不逐篇无限追溯'}
  dump(dest,obj);count+=1
 checkpoint={'batch':packet_data[0]['batch'] if packet_data else '', 'new_records':count,'total_new_semantic_reviews':len(list((OUT/'decisions').glob('*.json'))),'started_Asia_Shanghai':start,'finished_Asia_Shanghai':now()};dump(OUT/'checkpoints'/f'{checkpoint["batch"]}.json',checkpoint);print(json.dumps(checkpoint,ensure_ascii=False))
 render()
def render():
 began=time.monotonic();base=json.loads((OUT/'baseline_catalog.json').read_text());prior=json.loads((OUT/'baseline_screening.json').read_text());new=[json.loads(p.read_text()) for p in sorted((OUT/'decisions').glob('*.json'))] if (OUT/'decisions').exists() else [];done={r['id']:r for r in prior+new};meta={r['id']:r for r in json.loads((OUT/'metadata_enrichment.json').read_text())};routed={r['id']:r for r in json.loads((OUT/'computational_prescreen.json').read_text())};newids={r['id'] for r in new};priority=[];prior_eval={r['文献ID']:r for r in json.loads((ROOT/'reading_priority/evaluations.json').read_text())}
 for r in base:
  m=meta[r['id']];r.update({'citation_count':m['citation_count'],'citation_source':m['citation_source'],'citation_query_date':m['citation_query_date'],'jif2024':m['jif2024'],'jcr2024_categories':m['jcr2024_categories'],'journal_metric_status':m['journal_metric_status'],'abstract':m['abstract_enriched'] or r.get('abstract','')})
  if r['id'] in newids:
   s=done[r['id']];r.update({k:v for k,v in s.items() if k not in ['source_provenance','abstract_evidence_excerpt']})
  r['semantic_screening_complete']=r['id'] in done
  if r['id'] in done:
   s=done[r['id']];dec=s['screening_decision'];h=routed.get(r['id']) or classify_hint(r,m)
   if (dec in ['核心候选','机制背景'] or dec.startswith('符合本轮') or dec.startswith('机制原始')) and (r['id'] not in prior_eval or prior_eval[r['id']]['正式精读推荐']):
    groups=h['mechanism_themes_hint'];e=h['evidence_type_hint'];ad=h['AD_relation_hint_level'];tier='AD机制优先候选（暂定）' if ad>=3 and e else ('基础机制候选（暂定）' if any('遗传' in x or '酶学' in x for x in e) else '补充/边界候选（暂定）')
    if dec=='机制背景':tier='理论/综述/模型背景（暂定）'
    priority.append({'id':r['id'],'title':r['title'],'year':r.get('year'),'doi':r.get('doi'),'journal':r.get('journal'),'tier':tier,'mechanism_themes':groups,'objects':s.get('research_system'),'AD_relation':s.get('AD_relationship') or s.get('evidence',{}).get('ad_relevance_and_extrapolation_boundary'),'potential_evidence':e,'screening_reason':s['screening_reason'],'citation_count':m['citation_count'],'citation_query_date':m['citation_query_date'],'jif2024':m['jif2024'],'jcr2024_categories':m['jcr2024_categories'],'reading_basis':s['reading_level'],'recommendation_status':'全池题名/摘要暂定推荐；未新增全文核实/独立验证','previous20_detailed_evaluation_retained':r['id'] not in newids,'scientific_sort_score':h['scientific_hint_score']})
 priority.sort(key=lambda r:(-r['scientific_sort_score'],r['id']))
 # Scientific scores define strata; citation/JIF are only displayed auxiliary
 # data and never turn a scope exclusion/unknown into an included candidate.
 order=json.loads((OUT/'review_order_v2.json').read_text());remaining=[{**o,'title':next(r['title'] for r in base if r['id']==o['id']),'screening_stage':'未开始'} for o in order if o['id'] not in done]
 dump(ROOT/'catalog.json',base);dump(ROOT/'screening.json',prior+new);dump(ROOT/'remaining_queue.json',remaining);dump(OUT/'semantic_screening_records.json',prior+new);dump(OUT/'provisional_priorities.json',priority);csvout(OUT/'provisional_priorities_utf8.csv',priority);csvout(OUT/'semantic_screening_utf8.csv',prior+new);csvout(ROOT/'literature_master_utf8.csv',base);csvout(ROOT/'remaining_queue_utf8.csv',remaining)
 progress={'pool':len(base),'prior20_retained':len(prior),'new_semantic_screened':len(new),'cumulative_semantic_screened':len(done),'remaining_semantic_unscreened':len(remaining),'computational_prescreen_only_not_counted_as_complete':len(routed),'provisional_priority_articles':len(priority),'new_decisions':dict(collections.Counter(r['screening_decision'] for r in new)),'next_id':remaining[0]['id'] if remaining else None,'citation_records':sum(m['citation_count'] is not None for m in meta.values()),'new_fulltext_verifications':0,'new_independent_validation_claims':0,'updated_Asia_Shanghai':now()};dump(OUT/'progress.json',progress)
 old=json.loads((OUT/'baseline_summary.json').read_text());old.update(cumulative_semantic_screened=len(done),remaining_unscreened_records=len(remaining),next_id=progress['next_id'],all_pool_progress='all_pool/progress.json');dump(ROOT/'summary.json',old)
 # Preserve original six sheets and all three reading-evaluation sheets, then
 # update the master/queue and add whole-pool fields in separate cumulative tabs.
 wb=load_workbook(ROOT/'oxygen_response_review_cumulative.xlsx')
 sheets=[('文献总表',base,['id','title','year','journal','doi','pmid','screening_stage','screening_decision','screening_reason','reading_level','citation_count','citation_query_date','jif2024','journal_metric_status','semantic_screening_complete','old_ad_classification','old_human_classification','abstract']),('续跑队列',remaining,None),('全池筛选',prior+new,None),('全池暂定优先',priority,None),('全池指标',list(meta.values()),None)]
 for name,rows,fields in sheets:
  if name in wb:
   ix=wb.sheetnames.index(name);del wb[name];ws=wb.create_sheet(name,ix)
  else:ws=wb.create_sheet(name)
  fields=fields or list(dict.fromkeys(k for r in rows for k in r));ws.append(fields)
  for r in rows:
   values=[]
   for k in fields:
    v=r.get(k);v=json.dumps(v,ensure_ascii=False) if isinstance(v,(dict,list)) else v
    if isinstance(v,str):v=v[:32700];v="'"+v if v.startswith(('=','+','@')) else v
    values.append(v)
   ws.append(values)
  ws.freeze_panes='C2';ws.auto_filter.ref=ws.dimensions
  for c in ws[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='244766')
  for col in ws.columns:ws.column_dimensions[col[0].column_letter].width=50 if any(w in str(col[0].value) for w in ['title','reason','abstract']) else 24
 wb.save(ROOT/'oxygen_response_review_cumulative.xlsx')
 # Batch check deliberately excludes unchanged originals/hashes.
 wb=load_workbook(ROOT/'oxygen_response_review_cumulative.xlsx',read_only=True);assert wb['文献总表'].max_row-1==len(base);assert wb['全池筛选'].max_row-1==len(done);assert wb['全池暂定优先'].max_row-1==len(priority);wb.close();assert len({r['id'] for r in base})==len(base);assert len(done)+len(remaining)==len(base)
 with (OUT/'render_timing.jsonl').open('a') as f:f.write(json.dumps({'at_Asia_Shanghai':now(),'records':len(done),'measured_render_seconds':round(time.monotonic()-began,3)})+'\n')
 print(json.dumps(progress,ensure_ascii=False))
if __name__=='__main__':
 ap=argparse.ArgumentParser();sp=ap.add_subparsers(dest='command');sp.add_parser('init');p=sp.add_parser('packet');p.add_argument('--offset',type=int,default=0);p.add_argument('--size',type=int,default=50);p.add_argument('--mode',default='ranked');a=sp.add_parser('apply');a.add_argument('path');sp.add_parser('render');x=ap.parse_args()
 if x.command=='init':init()
 elif x.command=='packet':packet(x.offset,x.size,x.mode)
 elif x.command=='apply':apply(x.path)
 else:render()
